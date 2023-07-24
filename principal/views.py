# -*- coding: utf-8 -*-
from django.utils import timezone
from datetime import datetime, timedelta
from django.contrib.auth.decorators import login_required, permission_required, user_passes_test
from django.contrib.auth.models import User
from django.contrib.auth import login as auth_login, logout, authenticate
#from htmlmin.decorators import minified_response
# from django.core import serializers
from django.core.mail import send_mail, EmailMultiAlternatives, EmailMessage
# from django.core.serializers import serialize
# from django.core.serializers.json import DjangoJSONEncoder
from django.http import HttpResponse, HttpResponseRedirect, JsonResponse, Http404,FileResponse
from django.template.loader import render_to_string, get_template
from django.shortcuts import render,redirect
from django.urls import reverse
#from django.core.urlresolvers import //reverse esta ya no en la version 1.11.17
# from django.utils.decorators import method_decorator
# from django.utils.encoding import force_text
# from django.template import RequestContext
from principal.models import *
import os
from django.conf import settings

#---------NUEVO-------#
from django.db.models import Q
from django.core.files.images import get_image_dimensions
from django.contrib.auth.hashers import make_password
from django.contrib.auth.models import User
from django.db import transaction,connections
#import cStringIO as StringIO // ya no es valido en python 3.X 
from io import StringIO, BytesIO
from xhtml2pdf import pisa
from cgi import *
# import re
import json as simplejson
##para generar barcode
from base64 import b64encode
from reportlab.lib import units
from reportlab.graphics import renderPM
from reportlab.graphics.barcode import createBarcodeDrawing
from reportlab.lib.pagesizes import letter
from reportlab.lib.units import mm
from reportlab.pdfgen import canvas
from reportlab.graphics import renderPDF
from reportlab.graphics.barcode.qr import QrCodeWidget
from reportlab.graphics.shapes import Drawing
import io
import tempfile
#from barcode import Code128   //Revisar esta importacion
##barcode##
from django.db.models import Sum
#reload(sys)
#sys.setdefaultencoding('utf-8')
# ##MENSAJE DE TEXTO
from twilio.rest import Client
from django.db import transaction
import pandas as pd
from datetime import datetime, date
import calendar
from urllib.parse import unquote
import json
import numpy as np
from openpyxl import Workbook
#from htmlmin.decorators import minified_response
###############################################################################
def get_barcode(value, width, barWidth = 0.05 * units.inch, fontSize = 30, humanReadable = True):
	barcode = createBarcodeDrawing('Code128', value = value, barWidth = barWidth, fontSize = fontSize, humanReadable = humanReadable)
	drawing_width = width
	barcode_scale = drawing_width / barcode.width
	drawing_height = barcode.height * barcode_scale
	drawing = Drawing(drawing_width, drawing_height)
	drawing.scale(barcode_scale, barcode_scale)
	drawing.add(barcode, name='barcode')
	return drawing

def generate_qr_code(value, width, height):
    qr_code = QrCodeWidget(value)
    bounds = qr_code.getBounds()
    qr_code_width = bounds[2] - bounds[0]
    qr_code_height = bounds[3] - bounds[1]
    drawing = Drawing(width, height)
    drawing.add(qr_code)
    drawing.scale(width/qr_code_width, height/qr_code_height)
    buffer = io.BytesIO()
    renderPM.drawToFile(drawing, buffer, fmt='PNG')
    return b64encode(buffer.getvalue())

def inicio_cliente(request):
	return render(request,'inicio_cliente.html')

def buscar_envio_cliente(request):
	return render(request,'busqueda_cliente.html')

def resultado_busqueda(request):
	if request.is_ajax():
		try:
			codigo = request.GET['codigo']
			try:
				resultado_list = HistorialEnvio.objects.filter(codigo_envio__codigo=codigo).order_by('-pk')
				
				ctx = {'resultado':resultado_list}
				return render(request,'resultado_busqueda.html',ctx)
			except Exception as e:
				print (e,'error aqui es')
				raise e
				return render(request,'resultado_busqueda.html')
		except Exception as e:
			print (e)
			return render(request,'resultado_busqueda.html')

@login_required
def buscar_envio(request):
	empleado = Empleado.objects.get(usuario=request.user)
	empresa = EmpresaEmpleado.objects.get(empleado = empleado)
	hoy = timezone.now() 
	fecha_hoy = hoy.strftime("%Y-%m-%d")
	fecha_inicio = hoy - timedelta(days=180)
	envios = Envio.objects.filter(empresa=empresa.empresa, fecha_envio__range=[fecha_inicio,fecha_hoy])
	er = Revendedor.objects.count()
	if er >0:
		r = Revendedor.objects.filter(usuario=request.user)
		es_revendedor = False
		if r.count() >= 1:
			es_revendedor = True
			usuarios_estafeta = User.objects.filter(empleado_usuario__nombres_empleado__in=["ESTAFETA USA", "ESTAFETA","ESTAFETA_CENTRAL"]).values_list("id", flat=True)
			envios = Envio.objects.filter(revendedor=True, usuario_registro__in=usuarios_estafeta)
	else:
		revendedor = ''

	dic_envios = []
	for e in envios:
		envio = {}
		envio['pk'] = e.pk
		envio['codigo'] = e.codigo
		envio['quien_envia'] = e.quien_envia
		envio['pais_destino'] = e.pais_destino.nombre
		envio['departamento_destino'] = e.departamento_destino.nombre
		envio['direccion'] = e.direccion_registrar
		try:
			envio['estado'] = HistorialEnvio.objects.filter(codigo_envio=e.pk).latest('estado')
		except Exception as exe:
			envio['estado'] = e.pk 
		dic_envios.append(envio)
	ctx = {'es_revendedor':es_revendedor,'envios':dic_envios,'empleado':empleado}
	return render(request,'buscar_envio.html',ctx)

@login_required
def inicio_admin(request):
	empleado = Empleado.objects.get(usuario=request.user)
	empresa = EmpresaEmpleado.objects.get(empleado = empleado)
	estado_envio = EstadoEnvio.objects.filter(empresa = empresa.empresa)
	dic_datos_seguimiento = []
	for e in estado_envio:
		estado = {}
		estado['pk'] = e.pk
		estado['estado'] = e.estado
		estado['cantidad'] = Envio.objects.filter(empresa = e.empresa, estado_envio = e.pk).count()
		#estado['estado'] = HistorialEnvio.objects.filter(codigo_envio=e.pk).latest('estado')
		dic_datos_seguimiento.append(estado)

	ctx = {'estado_envio':estado_envio, 'dic_datos_seguimiento':dic_datos_seguimiento,'empleado':empleado}
	return render(request,'inicio_admin.html',ctx)

@login_required
def ver_paquetes(request,id_estado):
	empleado = Empleado.objects.get(usuario=request.user)
	empresa = EmpresaEmpleado.objects.get(empleado = empleado)
	estado_final = EstadoEnvio.objects.filter(empresa=empresa.empresa).last()
	estado = EstadoEnvio.objects.get(pk = id_estado)
	hoy = timezone.now() 
	fecha_hoy = hoy.strftime("%Y-%m-%d")
	fecha_inicio = hoy - timedelta(days=180)
	if int(id_estado) != int(estado_final.pk): 
		estado_sigui = int(id_estado)+1
		estado = EstadoEnvio.objects.get(pk=id_estado)
		estado_siguiente = EstadoEnvio.objects.get(pk=estado_sigui)
		
		if(int(id_estado) == 7):
			envios = Envio.objects.filter(estado_envio=estado, fecha_envio__range=[fecha_inicio,fecha_hoy]) 
		else:
			envios = Envio.objects.filter(estado_envio=estado)
		#envios = SeguimientoEnvio.objects.filter(estado=id_estado)
		volumen = sum(de.tipo_caja.tipo_caja.espacio_cubico for en in envios for de in DetalleEnvio.objects.filter(envio_id=en.pk))
		f_volumen = "{:.2f}".format(volumen)
		
		ctx = {'envios':envios,'estado':estado,'estado_siguiente':estado_siguiente,'estado_final':estado_final, 'volumen':f_volumen}
		return render(request,'ver_paquetes.html',ctx)
	else:
		estado = EstadoEnvio.objects.get(pk=id_estado)
		#envios = SeguimientoEnvio.objects.filter(estado=estado_final)
		if(int(id_estado) == 7):
			envios = Envio.objects.filter(estado_envio=estado, fecha_envio__range=[fecha_inicio,fecha_hoy])
		else:
			envios = Envio.objects.filter(estado_envio=estado)
		ctx = {'envios':envios,'estado':estado,'estado_final':estado_final}
		return render(request,'ver_paquetes.html',ctx) 

@login_required
def trasladar_post(request):
	if request.method == 'POST':
		envios = request.POST.getlist('mover')
		estado = request.POST.get('estado_siguiente')
		# print(envios)
		try:
			with transaction.atomic():
				for e in envios:
					envio_cliente = Envio.objects.get(pk=e)
					detalle_envio_qs = DetalleEnvio.objects.filter(envio=envio_cliente)
					estado_envio_q = EstadoEnvio.objects.get(pk=estado)
					envio_cliente.estado_envio =estado_envio_q
					
					for detalle_envio in detalle_envio_qs:
						detalle_envio.estado_hija = estado_envio_q
						detalle_envio.save()

					envio_cliente.save()

					empresa = Envio.objects.get(pk=e)
					estado_final = EstadoEnvio.objects.filter(empresa=empresa.empresa).last()
					# print(empresa.guia_revendedor)
					### UPDATE EEHN
					envio = SeguimientoEnvio.objects.filter(codigo_envio=Envio.objects.get(pk=e)).update(estado=estado)
					historial = HistorialEnvio.objects.create(codigo_envio=Envio.objects.get(pk=e),estado=EstadoEnvio.objects.get(pk=estado),usuario_registro=request.user)
					# print('no va')
					###UPDATE KRAKEN
					#OBTENER ENVIO EN kraken_cargo
					es_kraken = False

					# kraken_envio = SistemaEmpresaenvio.objects.using('kraken_cargo').filter(codigo = empresa.guia_revendedor).count()
					kraken_envio = 0
					# print('no va2 ', kraken_envio)
					if kraken_envio >= 1:
						es_kraken = True
					
					if es_kraken:
						estado_kraken = ''
						pk = 0
						if estado_envio_q.pk == 1:
							estado_kraken = 'BODEGA EEUU'
							pk = 8
						elif estado_envio_q.pk == 2:
							estado_kraken = 'PUERTO EEUU'
							pk = 9
						elif estado_envio_q.pk == 3:
							estado_kraken = 'TRANSITO MARITIMO'
							pk = 10
						elif estado_envio_q.pk == 4:
							estado_kraken = 'PUERTO HONDURAS'
							pk = 11
						elif estado_envio_q.pk == 5:
							estado_kraken = 'BODEGA HONDURAS'
							pk = 12
						elif estado_envio_q.pk == 6:
							estado_kraken = 'EN TRANSITO PARA ENTREGA'
							pk = 13
						elif estado_envio_q.pk == 7:
							estado_kraken = 'ENTREGADO'
							pk = 14
						estadokraken = SistemaEmpresaestadoenvio.objects.using('kraken_cargo').get(pk=pk)
						seguimiento_kraken_cargo = SistemaSeguimientoenvio.objects.db_manager('kraken_cargo').filter(codigo_envio=SistemaEmpresaenvio.objects.using('kraken_cargo').get(codigo =  envio_cliente.guia_revendedor)).update(estado=SistemaEmpresaestadoenvio.objects.using('kraken_cargo').get(pk=pk))
						historial_kraken_cargo = SistemaHistorialenvio.objects.db_manager('kraken_cargo').create(codigo_envio=SistemaEmpresaenvio.objects.using('kraken_cargo').get(codigo =  envio_cliente.guia_revendedor),estado=SistemaEmpresaestadoenvio.objects.using('kraken_cargo').get(pk=pk),usuario_registro=AuthUser.objects.using('kraken_cargo').get(pk=14),fechahora = timezone.now())
						envio_kraken = SistemaEmpresaenvio.objects.using('kraken_cargo').get(codigo =  envio_cliente.guia_revendedor)
						hoy = timezone.now()
					if envio_cliente.credito:
						try:
							if int(estado) == 4:
								account_sid = 'AC17aec886df904438ae784475a29acd60'
								auth_token = 'd36a6c11614c87a982ff21406c2e8bb3'
								twilio_client= Client(account_sid, auth_token)
								message = twilio_client.messages.create(
								body='ESTE ES UN MENSAJE GENERADO AUTOMATICAMENTE FAVOR NO RESPONDER. ' + '\n' + 'Saludos ' + envio_cliente.quien_envia.nombre_completo + '\n\n' +envio_cliente.empresa.nombre_empresa+' le informa que su envio '+envio_cliente.codigo + ' para ' + envio_cliente.quien_recibe.nombre_completo+ ', se encuentra en el Puerto de Honduras esperando el desembarque. Para evitar atrasos con su entrega favor realizar el pago pendiente de $'+ str(envio_cliente.saldo_pendiente) + '. Para más información contáctarse al número ' + envio_cliente.empresa.telefono_empresa ,
								from_='+13473345592',
								to='+1' + envio_cliente.quien_envia.celular
								)
								mensaje = EmpresaMensaje.objects.create(tipo_mensaje=TipoMensajes.objects.get(pk=2),envio=envio_cliente,texto = message.body)
								#print(message.sid)
						except Exception as e:
							print('No mensaje llegando a puerto destino error en numero: ', e)
							
					try:
						if int(estado) == int(estado_final.pk):
							account_sid = 'AC17aec886df904438ae784475a29acd60'
							auth_token = 'd36a6c11614c87a982ff21406c2e8bb3'
							twilio_client= Client(account_sid, auth_token)
							message = twilio_client.messages.create(
							body='ESTE ES UN MENSAJE GENERADO AUTOMATICAMENTE FAVOR NO RESPONDER. ' + '\n' +'Saludos ' + envio_cliente.quien_envia.nombre_completo + '\n\n' +envio_cliente.empresa.nombre_empresa+' le informa que su envio '+envio_cliente.codigo + ' para ' + envio_cliente.quien_recibe.nombre_completo+ ', ha sido entregado satisfactoriamente.' + '\n' + 'Muchas gracias por confiar en nuestros servicios esperamos pronto seguirle brindando el mejor servicio.',
							from_='+13473345592',
							# agregar concicionantes por pais
							to='+1' + envio_cliente.quien_envia.celular
							)
							#print(message.sid)
							mensaje = EmpresaMensaje.objects.create(tipo_mensaje=TipoMensajes.objects.get(pk=1),envio=envio_cliente,texto = message.body)
					except Exception as e:
							print('No mensaje estado final error en numero: ', e)
				return HttpResponseRedirect(reverse('ver_paquetes',kwargs={'id_estado':int(estado)-1}))
		except Exception as e:
			#print e, 'error'
			print("error --> ", e)
			transaction.rollback()
			return HttpResponseRedirect(reverse('ver_paquetes',kwargs={'id_estado':int(estado)-1}))
	elif request.method == 'GET':
		print(datos_envio)
		return datos_envio

@login_required
def validar_empresa(request):
	query_empresa,errores,ret_data,retorno = {},{},{},{}
	if request.POST.get('codigo_empresa') != '':
		ret_data.update({'codigo_empresa':request.POST.get('codigo_empresa').upper()})
		query_empresa.update({'codigo_empresa':request.POST.get('codigo_empresa').upper()})
	else:
		errores['codigo_empresa'] = u'Falta ingresar el codigo de la empresa para mostrarlo en los recibos'
	
	if request.POST.get('nombre_empresa') != '':
		ret_data.update({'nombre_empresa':request.POST.get('nombre_empresa').upper()})
		query_empresa.update({'nombre_empresa':request.POST.get('nombre_empresa').upper()})
	else:
		errores['nombre_empresa'] = u'Falta ingresar el nombre de la empresa'

	if  request.FILES.get('logo_empresa') != None:
		ret_data.update({'logo_empresa':request.FILES.get('logo_empresa')})
		query_empresa.update({'logo_empresa':request.FILES.get('logo_empresa')})
	else:
		errores['logo_empresa'] = u'Falta ingresar el logo de la empresa'

	if request.POST.get('direccion_empresa') != '':
		ret_data.update({'direccion_empresa':request.POST.get('direccion_empresa')})
		query_empresa.update({'direccion_empresa':request.POST.get('direccion_empresa')})
	else:
		errores['direccion_empresa'] = u'Falta ingresar la direccion de la empresa'

	if request.POST.get('telefono_empresa') != '':
		ret_data.update({'telefono_empresa':request.POST.get('telefono_empresa')})
		query_empresa.update({'telefono_empresa':request.POST.get('telefono_empresa')})
	else:
		errores['telefono_empresa'] = u'Falta ingresar el telefono de la empresa empresa'
	
	if request.POST.get('correo_empressa') != '':
		ret_data.update({'correo_empressa':request.POST.get('correo_empressa')})
		query_empresa.update({'correo_empressa':request.POST.get('correo_empressa')})
	else:
		errores['correo_empressa'] = u'Falta ingresar el correo de la empresa'

	retorno['query_empresa'] = query_empresa
	retorno['errores'] = errores
	retorno['ret_data'] = ret_data
	return retorno

@login_required
def empresa(request):
	empresas = Empresa.objects.all()
	print(empresas)
	if request.method == 'POST':
		ret_data,errores,query_empresa,recibo,ingreso_correcto = {},{},{},{},{}
		recibo = validar_empresa(request)
		ret_data.update(recibo['ret_data'])
		errores.update(recibo['errores'])
		query_empresa.update(recibo['query_empresa'])		
		if not errores:
			print('entra sin errores')
			try:
				#print('entra try')
				query_empresa.update({'usuario_registro':request.user})
				query_empresa.update({'celular_empresa':request.POST.get('celular_empresa')})
				empresa = Empresa(**query_empresa)
				empresa.save()
			except Exception as e:
				errores['extra'] = e
				transaction.rollback()
				ctx = {'empresas':empresas,'ret_data':ret_data,'errores':errores}
				return render(request,'empresa.html',ctx)
			else:
				transaction.commit()
				ingreso_correcto['mensaje'] = u"Datos guardados correctamente."
				ctx = {'ingreso_correcto':ingreso_correcto,'empresas':empresas}
				return HttpResponseRedirect(reverse('empresa'))
		else:
			#print errores,'erroresdentro de errores'
			ctx = {'ret_data':ret_data,'errores':errores,'empresas':empresas}
			return render(request,'empresa.html',ctx)
	elif request.method == 'GET':
		ctx = {'empresas':empresas}
		return render(request,'empresa.html', ctx)

@login_required
def validar_empleado(request,rol_empleado):
	query_empleado,errores,ret_data,retorno = {},{},{},{}
	if request.POST.get('nombres_empleado') != '':
		ret_data.update({'nombres_empleado':request.POST.get('nombres_empleado').upper()})
		query_empleado.update({'nombres_empleado':request.POST.get('nombres_empleado').upper()})
	else:
		errores['nombres_empleado'] = u'Falta ingresar el nombre del empleado'

	if request.POST.get('apellidos_empleado') != '':
		ret_data.update({'apellidos_empleado':request.POST.get('apellidos_empleado')})
		query_empleado.update({'apellidos_empleado':request.POST.get('apellidos_empleado')})
	else:
		errores['apellidos_empleado'] = u'Falta ingresar el telefono del empleado'

	if request.POST.get('telefono_empleado') != '':
		ret_data.update({'telefono_empleado':request.POST.get('telefono_empleado')})
		query_empleado.update({'telefono_empleado':request.POST.get('telefono_empleado')})
	else:
		errores['telefono_empleado'] = u'Falta ingresar el telefono del empleado'

	if request.POST.get('correo_empleado') != '':
		ret_data.update({'correo_empleado':request.POST.get('correo_empleado')})
		query_empleado.update({'correo_empleado':request.POST.get('correo_empleado')})
	else:
		errores['correo_empleado'] = u'Falta ingresar el correo del empleado'
	
	if request.POST.get('rol_empleado', '') and request.POST.get('rol_empleado', '') != '0':
		ret_data['id_rol_empleado'] = int(request.POST.get('rol_empleado',''))
		rol_empleado_query = rol_empleado.get(pk=request.POST.get('rol_empleado',''))
		query_empleado['rol_empleado'] = rol_empleado_query
	else:
		errores['rol_empleado'] = u'Falta ingresar el rol del empleado'

	retorno['query_empleado'] = query_empleado
	retorno['errores'] = errores
	retorno['ret_data'] = ret_data
	return retorno

@login_required
def empleados(request):
	usuario = User.objects.get(username=request.user)
	empleado = Empleado.objects.get(usuario = usuario)
	empresa = EmpresaEmpleado.objects.get(empleado = empleado)
	empresas = Empresa.objects.all()
	rol_empleado_actual = EmpresaEmpleado.objects.get(empleado = empleado)
	if request.user.is_superuser:
		empleados = EmpresaEmpleado.objects.all()
	else:
		empleados = EmpresaEmpleado.objects.filter(empresa = empresa.empresa)
	if request.user.is_superuser:
		rol_empleado = Rol.objects.all()
	else:
		#rol_empleado = Rol.objects.all()
		rol_empleado = Rol.objects.exclude(pk=rol_empleado_actual.empleado.rol_empleado.pk)
	
	if request.method == 'POST':
		ret_data,errores,query_empleado,recibo,ingreso_correcto = {},{},{},{},{}
		recibo = validar_empleado(request,rol_empleado)
		ret_data.update(recibo['ret_data'])
		errores.update(recibo['errores'])
		query_empleado.update(recibo['query_empleado'])	
		query_empleado.update({'usuario_registro':request.user})	
		if not errores:
			try:
				random_pass = User.objects.make_random_password(length=8, allowed_chars='abcdefghjkmnpqrstuvwxyz01234567889')
				contrasena = make_password(random_pass)
				nombre = request.POST.get('nombres_empleado')
				apellido = request.POST.get('apellidos_empleado')
				correo = request.POST.get('correo_empleado')
				nombre_primera = nombre[0].lower()
				apellidos = apellido.split()
				usuario_name = nombre_primera + str(apellidos[0]).lower()
				nusuario = User.objects.create(
					password = contrasena,
					last_login = datetime.now(),
					is_superuser = False,
					username = usuario_name,
					first_name = nombre,
					last_name = apellido,
					email = correo,
					is_staff = False,
					is_active = True,
					date_joined = datetime.now()
					)
				query_empleado.update({'usuario':nusuario})
				empleado = Empleado(**query_empleado)
				empleado.save()
				#print('guardo empleado')
				if request.user.is_superuser:
					#print('entro crear empresaempleado superuser')
					empresa_empleado = EmpresaEmpleado.objects.create(
						empresa = Empresa.objects.get(pk = request.POST.get('empresa_empleado')),
						empleado = empleado,
						usuario_registro = request.user
					)
					#print('creo empresaempleado superuser')
				else:
					#print('entro crear empresaempleado')
					empresa_empleado = EmpresaEmpleado.objects.create(
						empresa = empresa,
						empleado = empleado,
						usuario_registro = request.user
					)
					#print('creo empresaempleado')
			except Exception as e:
				errores['extra'] = e
				#print(errores)
				# borrar_usuario = User.objects.filter(pk=nusuario.pk)
				# #print(borrar_usuario, 'usuario creado')
				transaction.rollback()
				ctx = {'empleados':empleados,'empresas':empresas,'rol_empleado':rol_empleado,'rol_empleado_actual':rol_empleado_actual,'ret_data':ret_data,'errores':errores}
				return render(request,'empleados.html',ctx)
			else:
				transaction.commit()
				name = 'Bienvenido al sistema de ' + empresa.empresa.nombre_empresa
				subject = 'Nota de registro'
				to_email = nusuario.email
				mensaje = 'Estimado Sr. '+empleado.nombres_empleado.upper() +' '+ empleado.apellidos_empleado.upper() +' las credenciales para ingreso del sistema son las siguientes: '+'\n'+'Usuario: '+ nusuario.username +'\n'+'Contraseña: '+random_pass
				send_mail(subject, mensaje,to_email, [empleado.correo_empleado])
				##MENSAJE DE TEXTO
				account_sid = 'AC17aec886df904438ae784475a29acd60'
				auth_token = 'd36a6c11614c87a982ff21406c2e8bb3'
				twilio_client= Client(account_sid, auth_token)
				message = twilio_client.messages.create(
					body='Bienvenido al sistema de ' + empresa.empresa.nombre_empresa + '\n' + 'Estimado Sr./Sra. '+empleado.nombres_empleado.upper() +' '+ empleado.apellidos_empleado.upper()  + ' las credenciales para ingreso del sistema son las siguientes: '+'\n'+'Usuario: '+ nusuario.username +'\n'+'Contraseña: '+random_pass,
					from_='+13473345592',
					to='+1' + empleado.telefono_empleado
                          )
				mensaje = EmpresaMensaje.objects.create(tipo_mensaje=TipoMensajes.objects.get(pk=1),texto = message.body)
				ingreso_correcto['mensaje'] = u"Datos guardados correctamente."
				ctx = {'ingreso_correcto':ingreso_correcto,'empleados':empleados,'empresa':empresa,'rol_empleado':rol_empleado}
				return render(request,'empleados.html',ctx)
                #return HttpResponseRedirect(reverse('empleados')+"?ok")
		else:
			#print errores,'erroresdentro de errores'
			ctx = {'ret_data':ret_data,'errores':errores,'empleados':empleados,'empresas':empresas,'rol_empleado':rol_empleado,'rol_empleado_actual':rol_empleado_actual}
			return render(request,'empleados.html',ctx)
	elif request.method == 'GET':
		ctx = {'empleados':empleados,'empresas':empresas,'rol_empleado':rol_empleado,'rol_empleado_actual':rol_empleado_actual}
		return render(request,'empleados.html', ctx)

@login_required
def desactivar_usuario(request,id):
	usuario = User.objects.filter(pk = id).update(is_active = False)
	return HttpResponseRedirect(reverse('empleados'))

@login_required
def activar_usuario(request,id):
	usuario = User.objects.filter(pk = id).update(is_active = True)
	return HttpResponseRedirect(reverse('empleados'))

@login_required
def validar_cajas(request):
	query_caja,errores,ret_data,retorno = {},{},{},{}
	if request.POST.get('medida_caja') != '':
		ret_data.update({'medida_caja':request.POST.get('medida_caja').upper()})
		query_caja.update({'medida_caja':request.POST.get('medida_caja').upper()})
	else:
		errores['medida_caja'] = u'Falta ingresar el tamaño de la caja'
	
	if request.POST.get('compra') != '':
		ret_data.update({'compra':request.POST.get('compra').upper()})
		query_caja.update({'compra':request.POST.get('compra').upper()})
	else:
		errores['compra'] = u'Falta ingresar la cantidad comprada'
	
	if request.POST.get('valor') != '':
		ret_data.update({'valor':request.POST.get('valor').upper()})
		query_caja.update({'valor':request.POST.get('valor').upper()})
	else:
		errores['valor'] = u'Falta ingresar el valor de la compra'

	if request.POST.get('recibo_no') != '':
		ret_data.update({'recibo_no':request.POST.get('recibo_no')})
		query_caja.update({'recibo_no':request.POST.get('recibo_no')})
	else:
		errores['recibo_no'] = u'Falta ingresar el numero del recibo'

	if  request.FILES.get('recibo') != None:
		ret_data.update({'recibo':request.FILES.get('recibo')})
		query_caja.update({'recibo':request.FILES.get('recibo')})
	else:
		errores['recibo'] = u'Falta ingresar el logo de la empresa'
	
	if request.POST.get('fecha') != '':
		ret_data.update({'fecha':request.POST.get('fecha').upper()})
		query_caja.update({'fecha':request.POST.get('fecha').upper()})
	else:
		errores['fecha'] = u'Falta ingresar la fecha de compra'

	retorno['query_caja'] = query_caja
	retorno['errores'] = errores
	retorno['ret_data'] = ret_data
	return retorno

@login_required
def cajas(request):
	es_sadmin = Empleado.objects.get(usuario=request.user)
	empresa = EmpresaEmpleado.objects.get(empleado = es_sadmin)
	
	if request.user.is_superuser == True:
		empresas = Empresa.objects.all()
		cajas = EmpresaControlCaja.objects.all()
	else:
		if es_sadmin.rol_empleado.pk == 1:
			cajas = EmpresaControlCaja.objects.filter(empresa=empresa)
	
	if request.method == 'POST':
		ret_data,errores,query_caja,recibo,ingreso_correcto = {},{},{},{},{}
		recibo = validar_cajas(request)
		ret_data.update(recibo['ret_data'])
		errores.update(recibo['errores'])
		query_caja.update(recibo['query_caja'])
		if not errores:
			try:
				if request.user.is_superuser == True or es_sadmin.rol_empleado.pk == 1:
					query_caja.update({'empresa':Empresa.objects.get(pk=request.POST.get('empresa'))})
				else:
					query_caja.update({'empresa':empresa})

				query_caja.update({'existencia':request.POST.get('compra')})
				query_caja.update({'usuario_registro':request.user})
				caja = EmpresaControlCaja(**query_caja)
				caja.save()
			except Exception as e:
				errores['extra'] = e
				transaction.rollback()
				ctx = {'empresas':empresas,'cajas':cajas,'es_sadmin':es_sadmin,'ret_data':ret_data,'errores':errores}
				return render(request,'cajas.html',ctx)
			else:
				transaction.commit()
				ingreso_correcto['mensaje'] = u"Datos guardados correctamente."
				ctx = {'ingreso_correcto':ingreso_correcto,'empresas':empresas,'cajas':cajas,'es_sadmin':es_sadmin}
				return HttpResponseRedirect(reverse('cajas'))
		else:
			ctx = {'ret_data':ret_data,'errores':errores,'empresas':empresas,'cajas':cajas,'es_sadmin':es_sadmin}
			return render(request,'cajas.html',ctx)
	elif request.method == 'GET':
		ctx = {'empresas':empresas,'cajas':cajas,'es_sadmin':es_sadmin}
		return render(request,'cajas.html', ctx)

@login_required
def tipo_gasto(request):
	es_sadmin = Empleado.objects.get(usuario=request.user)
	empresa = EmpresaEmpleado.objects.get(empleado = es_sadmin)
	
	if request.user.is_superuser == True:
		#print('entro a superuser')
		empresas = Empresa.objects.all()
		gastos = EmpresaGasto.objects.all()
		tipo_gasto = EmpresaTipoGasto.objects.all()
	else:
		if es_sadmin.rol_empleado.pk == 1:
			#print('entro a rol empleado')
			gastos = EmpresaGasto.objects.filter(empresa=empresa)
			tipo_gasto = EmpresaTipoGasto.filter(empresa=empresa)

	if request.method == 'POST':
		ret_data,errores_tipo_gasto,ingreso_correcto_tipo_gasto = {},{},{}

		if not errores_tipo_gasto:
			#print('entra sin errores_tipo_gasto')
			try:
				#print('entra try')
				tipo = request.POST.get('tipo_gasto').upper()
				if request.user.is_superuser == True or es_sadmin.rol_empleado.pk == 1:
					eingresar = Empresa.objects.get(pk=request.POST.get('empresa'))
					tipo_gastos = EmpresaTipoGasto.objects.create(
						tipo_gasto = tipo,
						empresa = eingresar)
				else:
					tipo_gastos = EmpresaTipoGasto.objects.create(
						tipo_gasto = tipo,
						empresa = empresa)
				#print('guardo')
			except Exception as e:
				#print(e, 'errores_tipo_gasto extras')
				errores_tipo_gasto['extra'] = e
				transaction.rollback()
				ctx = {'empresas':empresas,'cajas':cajas,'es_sadmin':es_sadmin,'ret_data':ret_data,'errores_tipo_gasto':errores_tipo_gasto}
				return render(request,'gastos.html',ctx)
			else:
				transaction.commit()
				ingreso_correcto_tipo_gasto['mensaje'] = u"Datos guardados correctamente."
				ctx = {'ingreso_correcto_tipo_gasto':ingreso_correcto_tipo_gasto,'empresas':empresas,'cajas':cajas,'es_sadmin':es_sadmin}
				return HttpResponseRedirect(reverse('gastos')+"?okt")
		else:
			#print (errores_tipo_gasto,'errores_tipo_gastodentro de errores_tipo_gasto')
			ctx = {'ret_data':ret_data,'errores_tipo_gasto':errores_tipo_gasto,'empresas':empresas,'cajas':cajas,'es_sadmin':es_sadmin}
			return render(request,'gastos.html',ctx)
	elif request.method == 'GET':
		ctx = {'empresas':empresas,'cajas':cajas,'es_sadmin':es_sadmin}
		return render(request,'gastos.html', ctx)

@login_required
def validar_gastos(request,tipo_gasto):
	query_gastos,errores,ret_data,retorno = {},{},{},{}
	if request.POST.get('tipo_gasto_ingresar', '') and request.POST.get('tipo_gasto_ingresar', '') != '0':
		ret_data['id_tipo_gasto'] = int(request.POST.get('tipo_gasto_ingresar',''))
		tipo_gasto_query = tipo_gasto.get(pk=request.POST.get('tipo_gasto_ingresar',''))
		query_gastos['tipo_gasto'] = tipo_gasto_query
	else:
		errores['tipo_gasto'] = u'Falta ingresar el tipo de gasto'
	
	if request.POST.get('recibo') != '':
		ret_data.update({'recibo':request.POST.get('recibo').upper()})
		query_gastos.update({'recibo':request.POST.get('recibo').upper()})
	else:
		errores['recibo'] = u'Falta ingresar el numero de recibo'
	
	if request.POST.get('valor') != '':
		ret_data.update({'valor':request.POST.get('valor').upper()})
		query_gastos.update({'valor':request.POST.get('valor').upper()})
	else:
		errores['valor'] = u'Falta ingresar el valor de la compra'

	if request.POST.get('descripcion') != '':
		ret_data.update({'descripcion':request.POST.get('descripcion').upper()})
		query_gastos.update({'descripcion':request.POST.get('descripcion')})
	else:
		errores['descripcion'] = u'Falta ingresar la descripcion del gasto'

	if  request.FILES.get('pdf') != None:
		ret_data.update({'pdf':request.FILES.get('pdf')})
		query_gastos.update({'pdf':request.FILES.get('pdf')})
	else:
		errores['pdf'] = u'Falta ingresar el recibo'
	
	if request.POST.get('fecha') != '':
		ret_data.update({'fecha':request.POST.get('fecha').upper()})
		query_gastos.update({'fecha':request.POST.get('fecha').upper()})
	else:
		errores['fecha'] = u'Falta ingresar la fecha de compra'

	retorno['query_gastos'] = query_gastos
	retorno['errores'] = errores
	retorno['ret_data'] = ret_data
	return retorno

@login_required
def gastos(request):
	es_sadmin = Empleado.objects.get(usuario=request.user)
	empresa = EmpresaEmpleado.objects.get(empleado = es_sadmin)
	
	if request.user.is_superuser == True:
		#print('entro a superuser')
		empresas = Empresa.objects.all()
		gastos = EmpresaGasto.objects.all()
		tipo_gasto = EmpresaTipoGasto.objects.all()
	else:
		if es_sadmin.rol_empleado.pk == 1:
			#print('entro a rol empleado')
			gastos = EmpresaGasto.objects.filter(empresa=empresa)
			tipo_gasto = EmpresaTipoGasto.filter(empresa=empresa)

	if request.method == 'POST':
		ret_data,errores,query_gastos,recibo,ingreso_correcto = {},{},{},{},{}
		recibo = validar_gastos(request, tipo_gasto)
		ret_data.update(recibo['ret_data'])
		errores.update(recibo['errores'])
		query_gastos.update(recibo['query_gastos'])
		#print request.POST.get('tipo_gasto_ingresar'), 'tipo_gasto'
		if not errores:
			#print('entra sin errores')
			try:
				#print('entra try')
				gasto_empresa =  EmpresaTipoGasto.objects.get(pk=request.POST.get('tipo_gasto_ingresar'))
				query_gastos.update({'empresa':gasto_empresa.empresa})
				query_gastos.update({'usuario_registro':request.user})

				gasto = EmpresaGasto(**query_gastos)
				gasto.save()
				#print('guardo')
			except Exception as e:
				#print(e, 'errores extras')
				errores['extra'] = e
				transaction.rollback()
				ctx = {'empresas':empresas, 'es_sadmin':es_sadmin,'ret_data':ret_data,'errores':errores,'tipo_gasto':tipo_gasto,'gastos':gastos}
				return render(request,'gastos.html',ctx)
			else:
				transaction.commit()
				ingreso_correcto['mensaje'] = u"Datos guardados correctamente."
				ctx = {'ingreso_correcto':ingreso_correcto,'empresas':empresas, 'es_sadmin':es_sadmin,'tipo_gasto':tipo_gasto,'gastos':gastos}
				return HttpResponseRedirect(reverse('gastos')+"?ok")
		else:
			#print errores,'erroresdentro de errores'
			ctx = {'ret_data':ret_data,'errores':errores,'empresas':empresas, 'es_sadmin':es_sadmin,'tipo_gasto':tipo_gasto,'gastos':gastos}
			return render(request,'gastos.html',ctx)
	elif request.method == 'GET':
		ctx = {'empresas':empresas, 'es_sadmin':es_sadmin,'tipo_gasto':tipo_gasto,'gastos':gastos}
		return render(request,'gastos.html', ctx)

@login_required
def ver_recibo_gastos(request, id):
    gasto = EmpresaGasto.objects.get(pk = id)
    filename = gasto.pdf.name.split('/')[-1]
    separar = filename.split('.')
    if separar[1] == 'pdf':
        filepath = os.path.join(settings.MEDIA_ROOT, "archivos_gastos", filename)
        return FileResponse(open(filepath, 'rb'), content_type='application/pdf')
    elif separar[1].lower() == 'jpg' or separar[1].lower() == 'jpeg' or separar[1].lower() == 'png':
        width, height = get_image_dimensions(gasto.pdf)
        return render(request, 'imagen_gasto.html', {'gasto': gasto,'anchura': width, 'altura': height})
    else:
        response = HttpResponse(caja.pdf, content_type='text/plain')
        response['Content-Disposition'] = 'attachment; filename=%s' % filename
        return response

@login_required
def ver_recibos_cajas(request, id):
    caja = EmpresaControlCaja.objects.get(pk = id)
    filename = caja.recibo.name.split('/')[-1]
    separar = filename.split('.')
    if separar[1] == 'pdf':
        filepath = os.path.join(settings.MEDIA_ROOT, "archivos_compra_cajas", filename)
        return FileResponse(open(filepath, 'rb'), content_type='application/pdf')
    elif separar[1].lower() == 'jpg' or separar[1].lower() == 'jpeg' or separar[1].lower() == 'png':
        width, height = get_image_dimensions(caja.recibo)
        return render(request, 'imagen_caja.html', {'caja': caja,'anchura': width, 'altura': height})
    else:
        response = HttpResponse(caja.recibo, content_type='text/plain')
        response['Content-Disposition'] = 'attachment; filename=%s' % filename
        return response

#########
@login_required
def validar_cajas_detalle(request,caja):
	query_caja,errores,ret_data,retorno = {},{},{},{}
	if request.POST.get('caja', '') and request.POST.get('caja', '') != '0':
		ret_data['id_caja'] = int(request.POST.get('caja',''))
		caja_query = caja.get(pk=request.POST.get('caja',''))
		query_caja['caja'] = caja_query
	else:
		errores['rol_empleado'] = u'Falta ingresar el rol del empleado'

	if request.POST.get('cantidad') != '':
		ret_data.update({'cantidad':request.POST.get('cantidad').upper()})
		query_caja.update({'cantidad':request.POST.get('cantidad').upper()})
	else:
		errores['cantidad'] = u'Falta ingresar la cantidad comprada'
	
	if request.POST.get('valor') != '':
		ret_data.update({'valor':request.POST.get('valor').upper()})
		query_caja.update({'valor':request.POST.get('valor').upper()})
	else:
		errores['valor'] = u'Falta ingresar el valor de la compra'

	if request.POST.get('recibo_no_detalle') != '':
		ret_data.update({'recibo_no_detalle':request.POST.get('recibo_no_detalle')})
		query_caja.update({'recibo_no':request.POST.get('recibo_no_detalle')})
	else:
		errores['recibo_no_detalle'] = u'Falta ingresar el numero del recibo'

	if  request.FILES.get('recibo_detalle') != None:
		ret_data.update({'recibo':request.FILES.get('recibo_detalle')})
		query_caja.update({'recibo':request.FILES.get('recibo_detalle')})
	else:
		errores['recibo_detalle'] = u'Falta ingresar el recibo de la compra'
	
	if request.POST.get('fecha_detalle') != '':
		ret_data.update({'fecha':request.POST.get('fecha_detalle').upper()})
		query_caja.update({'fecha':request.POST.get('fecha_detalle').upper()})
	else:
		errores['fecha_detalle'] = u'Falta ingresar la fecha de compra'

	retorno['query_caja'] = query_caja
	retorno['errores'] = errores
	retorno['ret_data'] = ret_data
	return retorno

@login_required
def detalle_cajas(request):
	es_sadmin = Empleado.objects.get(usuario=request.user)
	empresa = EmpresaEmpleado.objects.get(empleado = es_sadmin)
	caja =  EmpresaControlCaja.objects.all()
	if request.user.is_superuser == True:
		empresas = Empresa.objects.all()
		cajas = EmpresaControlCaja.objects.all()
	else:
		if es_sadmin.rol_empleado.pk == 1:
			cajas = EmpresaControlCaja.objects.filter(empresa=empresa)
	
	if request.method == 'POST':
		ret_data,errores,query_caja,recibo,ingreso_correcto = {},{},{},{},{}
		recibo = validar_cajas_detalle(request,caja)
		ret_data.update(recibo['ret_data'])
		errores.update(recibo['errores'])
		query_caja.update(recibo['query_caja'])
		if not errores:
			#print('entra sin errores')
			try:
				query_caja.update({'usuario_registro':request.user})
				query_caja.update({'accion':False})
				query_caja.update({'fecha':request.POST.get('fecha_detalle')})
				#print request.POST.get('fecha_detalle'), 'fecha'
				caja = EmpresaDetalleCaja(**query_caja)
				caja.save()
				#print('guardo')
			except Exception as e:
				#print(e, 'errores extras')
				errores['extra'] = e
				transaction.rollback()
				ctx = {'empresas':empresas,'cajas':cajas,'es_sadmin':es_sadmin,'ret_data':ret_data,'errores':errores}
				return render(request,'cajas.html',ctx)
			else:
				transaction.commit()
				existencia = EmpresaControlCaja.objects.get(pk = caja.caja.pk)
				#print existencia.existencia, 'existencia'
				control_caja = EmpresaControlCaja.objects.filter(pk = caja.caja.pk).update(existencia = int(existencia.existencia) + int(caja.cantidad))
				#print 'actualizo control caja'
				ingreso_correcto['mensaje'] = u"Datos guardados correctamente."
				ctx = {'ingreso_correcto':ingreso_correcto,'empresas':empresas,'cajas':cajas,'es_sadmin':es_sadmin}
				return HttpResponseRedirect(reverse('cajas'))
		else:
			#print errores,'erroresdentro de errores'
			ctx = {'ret_data':ret_data,'errores':errores,'empresas':empresas,'cajas':cajas,'es_sadmin':es_sadmin}
			return render(request,'cajas.html',ctx)
	elif request.method == 'GET':
		ctx = {'empresas':empresas,'cajas':cajas,'es_sadmin':es_sadmin}
		return render(request,'cajas.html', ctx)

@login_required
def detalle_compra_cajas(request, id):
	compras = EmpresaDetalleCaja.objects.filter(caja=id)
	caja = EmpresaControlCaja.objects.get(pk=id)
	ctx = {'compras':compras,'caja':caja}
	return render(request,'detalle_compras_cajas.html', ctx)
	
@login_required
def ver_detalle_recibos_cajas(request, id):
    caja = EmpresaDetalleCaja.objects.get(pk = id)
    filename = caja.recibo.name.split('/')[-1]
    separar = filename.split('.')
    if separar[1] == 'pdf':
        filepath = os.path.join(settings.MEDIA_ROOT, "archivos_compra_cajas", filename)
        return FileResponse(open(filepath, 'rb'), content_type='application/pdf')
    elif separar[1].lower() == 'jpg' or separar[1].lower() == 'jpeg' or separar[1].lower() == 'png':
        width, height = get_image_dimensions(caja.recibo)
        return render(request, 'imagen_caja.html', {'caja': caja,'anchura': width, 'altura': height})
    else:
        response = HttpResponse(caja.recibo, content_type='text/plain')
        response['Content-Disposition'] = 'attachment; filename=%s' % filename
        return response

@login_required
def validad_actividades(request,quien_envia):
	query_actividades,errores,ret_data,retorno = {},{},{},{}
	if request.POST.get('cliente', '') and request.POST.get('cliente', '') != '0':
		ret_data['id_cliente'] = int(request.POST.get('cliente',''))
		quien_envia_query = quien_envia.get(pk=request.POST.get('cliente',''))
		query_actividades['cliente'] = quien_envia_query
	else:
		errores['cliente'] = u'Falta ingresar el cliente'
	
	if request.POST.get('fecha') != '':
		ret_data.update({'fecha':request.POST.get('fecha').upper()})
		query_actividades.update({'fecha':request.POST.get('fecha')})
	else:
		errores['fecha'] = u'Falta ingresar la fecha'
	
	if request.POST.get('actividad') != '':
		ret_data.update({'actividad':request.POST.get('actividad').upper()})
		query_actividades.update({'actividad':request.POST.get('actividad')})
	else:
		errores['actividad'] = u'Falta ingresar la actividad'

	if request.POST.get('hora') != '':
		ret_data.update({'hora':request.POST.get('hora')})
		query_actividades.update({'hora':request.POST.get('hora')})
	else:
		errores['hora'] = u'Falta ingresar el numero del recibo'

	if request.POST.get('descripcion_pedido') != '':
		ret_data.update({'descripcion_pedido':request.POST.get('descripcion_pedido').upper()})
		query_actividades.update({'descripcion_pedido':request.POST.get('descripcion_pedido').upper()})
	else:
		errores['descripcion_pedido'] = u'Falta ingresar la descripcion del pedido'

	retorno['query_actividades'] = query_actividades
	retorno['errores'] = errores
	retorno['ret_data'] = ret_data
	return retorno

@login_required
def actividades(request):
	es_sadmin = Empleado.objects.get(usuario=request.user)
	empresa = EmpresaEmpleado.objects.get(empleado = es_sadmin)
	hoy = timezone.now() 
	fecha_hoy = hoy.strftime("%Y-%m-%d")
	fecha_inicio = hoy - timedelta(days=180)
	# print(fecha_hoy)
	cajas = EmpresaControlCaja.objects.filter(empresa=empresa.empresa)
	if request.user.is_superuser == True:
		empresas = Empresa.objects.all()
		cajas = EmpresaControlCaja.objects.all()
		actividades = EmpresaActividades.objects.filter(estado=False,fecha__range=[fecha_inicio,fecha_hoy]).order_by('-fecha')
		actividades_terminadas = EmpresaActividades.objects.filter(estado=True,fecha__range=[fecha_inicio,fecha_hoy]).order_by('-fecha')
		quien_envia = Cliente.objects.all()
	else:
		if es_sadmin.rol_empleado.pk == 1:
			cajas = EmpresaControlCaja.objects.filter(empresa=empresa)
			actividades = EmpresaActividades.objects.filter(estado=False,empresa=empresa,fecha__range=[fecha_inicio,fecha_hoy]).order_by('-fecha')
			actividades_terminadas = EmpresaActividades.objects.filter(estado=True,empresa=empresa,fecha__range=[fecha_inicio,fecha_hoy]).order_by('-fecha')
			quien_envia = Cliente.objects.filter(empresa= empresa)

	if request.method == 'POST':
		ret_data,errores,query_actividades,recibo,ingreso_correcto = {},{},{},{},{}
		recibo = validad_actividades(request,quien_envia)
		ret_data.update(recibo['ret_data'])
		errores.update(recibo['errores'])
		query_actividades.update(recibo['query_actividades'])
		if not errores:
			#print('entra sin errores')
			try:
				fecha = request.POST.get('fecha')
				#print fecha, 'fecha obtenida'
				fecha_ = fecha.split('/')
				fechafinal = fecha_[2] + '-' + fecha_[0] + '-' + fecha_[1]
				#print fechafinal, 'fecha convertirda'
				query_actividades.update({'fecha':fechafinal})
				query_actividades.update({'empresa':empresa.empresa})
				query_actividades.update({'usuario_registro':request.user})
				caja = EmpresaActividades(**query_actividades)
				caja.save()
				#print('guardo')
			except Exception as e:
				print('errores: ', e)
				errores['extra'] = e
				transaction.rollback()
				ctx = {'empresas':empresas,'cajas':cajas,'es_sadmin':es_sadmin,'ret_data':ret_data,'errores':errores,'quien_envia':quien_envia,'actividades':actividades,'cajas':cajas,'actividades_terminadas':actividades_terminadas}
				return render(request,'actividades.html',ctx)
			else:
				transaction.commit()
				ingreso_correcto['mensaje'] = u"Datos guardados correctamente."
				ctx = {'ingreso_correcto':ingreso_correcto,'empresas':empresas,'cajas':cajas,'es_sadmin':es_sadmin,'quien_envia':quien_envia,'actividades':actividades,'cajas':cajas,'actividades_terminadas':actividades_terminadas}
				return HttpResponseRedirect(reverse('actividades'))
		else:
			#print errores,'erroresdentro de errores'
			ctx = {'ret_data':ret_data,'errores':errores,'empresas':empresas,'cajas':cajas,'es_sadmin':es_sadmin,'quien_envia':quien_envia,'actividades':actividades,'actividades_terminadas':actividades_terminadas}
			return render(request,'actividades.html',ctx)
	elif request.method == 'GET':
		ctx = {'empresas':empresas,'cajas':cajas,'es_sadmin':es_sadmin,'quien_envia':quien_envia,'actividades':actividades,'actividades_terminadas':actividades_terminadas}
		return render(request,'actividades.html', ctx)

@login_required
def finalizar_actividad(request):
	recibo_entrega = ''
	if request.method == 'POST':
		error = False
		ret_data,errores,ingreso_correcto = {},{},{}
		cajas_post = request.POST.getlist('cajas')
		actividad = request.POST.get('actividad')
		deposito = request.POST.get('deposito_pedido')
		descripcion_entrega = ''
		texto = ''
		if len(cajas_post) == 0:
			errores.update({'cajas_post':'Debe seleccionar al menos una caja'})
			ret_data.update({'cajas_envio': request.POST.getlist('cajas')})
			error = True
			
		
		if not error:
			try:
				with transaction.atomic():
					for p in cajas_post:
						separar = p.split("|")
						#print separar[0],'separar0'
						#print separar[1],'separar1'
						caja = EmpresaControlCaja.objects.get(pk = separar[0])
						if int(separar[1]) > 1:
							texto = 'CAJAS'
						else:
							texto = 'CAJA'
						if descripcion_entrega == '':
							descripcion_entrega += separar[1] + " " + texto + " de " + caja.medida_caja
						else:
							descripcion_entrega += ", "+separar[1] + " " + texto + " de " + caja.medida_caja 
					
					cliente = EmpresaActividades.objects.get(pk = actividad)
					recibo_no = ReciboCaja.objects.filter(empresa=cliente.empresa).count()
					recibo_no += 1
					codigo_final = cliente.empresa.codigo_empresa+ '00' + str(recibo_no)
					new_recibo_caja = ReciboCaja.objects.create(
						empresa = cliente.empresa,
						recibo_no=codigo_final,
						cliente=cliente.cliente,
						descripcion_entrega=descripcion_entrega,
						valor_caja=deposito,
						usuario_registro=request.user)
					nactividad = EmpresaActividades.objects.filter(pk=actividad).update(estado = True, deposito =deposito, fecha_realizo = datetime.now(), descripcion_entrega = descripcion_entrega, recibo_entrega= new_recibo_caja.pk)
					recibo_entrega = new_recibo_caja.pk
			except Exception as e:
				#print 'hay errores',e
				errores['extra'] = e
				print("Error --> | ", str(e))
				transaction.rollback()
				ctx = {'ret_data':ret_data,'cajas_post':cajas_post,'error':error}
				return HttpResponseRedirect(reverse('actividades'))
			else:
				transaction.commit()
				empresa = ''
				for p in cajas_post:
					separar = p.split("|")
					#print separar[0],'separar0 commit'
					#print separar[1],'separar1 commit'
					caja = EmpresaControlCaja.objects.get(pk = separar[0])
					empresa= caja
					update_caja = EmpresaDetalleCaja.objects.create(caja = caja,cantidad = separar[1],recibo_caja = ReciboCaja.objects.get(pk=recibo_entrega),usuario_registro = request.user)
					#print 'creo empresa detalle caja al entregar una caja'
					update_control_caja = EmpresaControlCaja.objects.filter(pk=separar[0]).update(existencia = caja.existencia - int(separar[1]))
				ingreso_correcto['mensaje'] = u"Datos guardados correctamente."
				ctx = {'ingreso_correcto':ingreso_correcto,}
				return HttpResponseRedirect(reverse('actividades')+"?ok" +"&envio="+ str(actividad))
		else:
			ctx = {'ret_data':ret_data,'errores':errores, 'cajas_post':cajas_post,'error':error}
			return HttpResponseRedirect(reverse('actividades'))
	else:
		return HttpResponseRedirect(reverse('actividades'))

def recibo_caja_pdf(request, id):
	recibo_actividad = EmpresaActividades.objects.get(id=id) 
	caja = ReciboCaja.objects.get(pk = recibo_actividad.recibo_entrega)
	fecha = caja.fecha
	barcode = get_barcode(value = id, width = 600)
	codigo = b64encode(renderPM.drawToString(barcode, fmt = 'PNG'))
	return generar_pdf('recibo_caja_pdf.html',{'pagesize':'A4','orientation':'landscape','caja':caja,'codigo':codigo, 'fecha':fecha})

@login_required
def terminar_tarea(request):
	hoy = timezone.now().date()
	tarea = request.GET['id']

	nactividad = EmpresaActividades.objects.filter(pk=tarea).update(estado = True,fecha_realizo = datetime.now())

@login_required
def reactivar_tarea(request):
	if request.is_ajax():
		hoy = datetime.now().date()
		tarea = request.GET['id']
		try:
			act = AgendaActividades.objects.filter(pk=tarea).update(estado=False, deposito = 0.00, fecha_realizo=hoy)
		except Exception as e:
			data = {'valor':1}
			return HttpResponse(simplejson.dumps(data), content_type='application/json')
		else:
			data = {'valor':1}
			return HttpResponse(simplejson.dumps(data), content_type='application/json')
	else:
		data = {'valor':1}
		return HttpResponse(simplejson.dumps(data), content_type='application/json')

@login_required
def validar_clientes(request):
	query_clientes,errores,ret_data,retorno = {},{},{},{}
	if request.POST.get('nombre_completo') != '':
		ret_data.update({'nombre_completo':request.POST.get('nombre_completo')})
		query_clientes.update({'nombre_completo':request.POST.get('nombre_completo')})
	else:
		errores['nombre_completo'] = u'Falta ingresar el nombre del cliente'

	if request.POST.get('celular') != '':
		ret_data.update({'celular':request.POST.get('celular')})
		query_clientes.update({'celular':request.POST.get('celular')})
	else:
		errores['direccion'] = u'Falta ingresar la direccion'

	if request.POST.get('direccion') != '':
		ret_data.update({'direccion':request.POST.get('direccion')})
		query_clientes.update({'direccion':request.POST.get('direccion')})
	else:
		errores['direccion'] = u'Falta ingresar la direccion'

	retorno['query_clientes'] = query_clientes
	retorno['errores'] = errores
	retorno['ret_data'] = ret_data

	return retorno

@login_required
def clientes(request):
	empleado = Empleado.objects.get(usuario=request.user)
	empresa = EmpresaEmpleado.objects.get(empleado = empleado)
	clientes = Cliente.objects.filter(empresa=empresa.empresa).order_by('-pk')
	duplicado = ''
	if request.method == 'POST':
		ret_data,errores,query_clientes,recibo,ingreso_correcto = {},{},{},{},{}
		recibo = validar_clientes(request)
		ret_data.update(recibo['ret_data'])
		errores.update(recibo['errores'])
		query_clientes.update(recibo['query_clientes'])
		codigo = Cliente.objects.filter(celular=request.POST.get('celular'))			
		if not errores:
			# if codigo.count() == 1:
			# 	duplicado = 'Ya existe un Cliente con ese numero telefonico'
			# 	#print('error duplicado')
			# 	ctx = {'ret_data':ret_data,'duplicado':duplicado, 'clientes':clientes}
			# 	return render(request, 'clientes.html',ctx)
			try:
				nombre_completo = request.POST.get('nombre_completo')
				celular = request.POST.get('celular')
				direccion = request.POST.get('direccion')
				correo = request.POST.get('correo')
				usuario_registro = request.user

				cliente = Cliente.objects.create(nombre_completo = nombre_completo.upper(),celular = celular,direccion = direccion.upper(),usuario_registro = usuario_registro, empresa = empresa.empresa)				

			except Exception as e:
				#print e,'aqui esta reventando'
				errores['extra'] = e
				transaction.rollback()
				ctx = {'ret_data':ret_data,'errores':errores, 'clientes':clientes,'empleado':empleado}
				return render(request, 'clientes.html',ctx)
			else:
				transaction.commit()
				ingreso_correcto['mensaje'] = u"Datos guardados correctamente."
				ctx = {'ingreso_correcto':ingreso_correcto, 'clientes':clientes,'empleado':empleado}
				return HttpResponseRedirect(reverse('clientes'))
		else:
			#print errores,'erroresdentro de errores'
			ctx = {'ret_data':ret_data,'errores':errores, 'clientes':clientes,'empleado':empleado}
			return render(request,'clientes.html',ctx)
	else:
		ctx = {'clientes':clientes,'empleado':empleado}
		return render(request,'clientes.html', ctx)

@login_required
def recibe(request):
	if request.is_ajax():
		quien_envia = request.POST.get('quien_envia')
		nombre = request.POST.get('nombre').upper()
		celular = request.POST.get('celular')
		direccion = request.POST.get('direccion').upper()
		try:
			quien_recibe = ClienteRecibe.objects.create(
				cliente_envia = Cliente.objects.get(pk=quien_envia),
				nombre_completo = nombre,
				celular = celular,
				direccion = direccion,
				usuario_registro = request.user
			)
		except Exception as e:
			#print e,'error recibe'
			data = {'valor':1}
			return HttpResponse(simplejson.dumps(data), content_type='application/json')
		else:
			data = {'valor':1}
			return HttpResponse(simplejson.dumps(data), content_type='application/json')
	else:
		data = {'valor':2}
		return HttpResponse(simplejson.dumps(data), content_type='application/json')

@login_required
def ver_recibe(request):
	recibe = list(ClienteRecibe.objects.filter(cliente_envia=request.GET.get('quien_envia')).values('pk','nombre_completo','celular','direccion','pais','departamento'))
	return HttpResponse(simplejson.dumps(recibe), content_type='application/javascript')
	connection.close()

@login_required
def datos_recibe(request):
	recibe = list(ClienteRecibe.objects.filter(pk=request.GET.get('cliente_recibe')).values('pk','nombre_completo','celular','direccion','pais','departamento','pais_pk__pk','departamento_pk__pk'))
	return HttpResponse(simplejson.dumps(recibe), content_type='application/javascript')
	connection.close()

@login_required
def datos_caja_edit(request):
	#print request.GET.get('pk_caja')
	caja = request.GET.get('pk_caja').split('|')
	pk_caja_pais = caja[3]
	
	caja = list(CajaPais.objects.filter(pk=pk_caja_pais).values('pk','precio','tipo_caja__descripcion','tipo_caja__alto','tipo_caja__ancho','tipo_caja__largo'))
	#print caja
	return HttpResponse(simplejson.dumps(caja), content_type='application/javascript')
	connection.close()

#MODULO REVENDEDOR START
@login_required
def validar_revendedor(request):
	query_revendedor,errores,ret_data,retorno = {},{},{},{}

	if request.POST.get('nombre_completo') != '':
		ret_data.update({'nombre_completo':request.POST.get('nombre_completo')})
		query_revendedor.update({'nombre_completo':request.POST.get('nombre_completo')})
	else:
		errores['nombre_completo'] = u'Falta ingresar el nombre del revendedor'
	
	if request.POST.get('telefono') != '':
		ret_data.update({'telefono':request.POST.get('telefono')})
		query_revendedor.update({'telefono':request.POST.get('telefono')})
	else:
		errores['telefono'] = u'Falta ingresar el telefono del revendedor'
	
	if request.POST.get('nombre_empresa') != '':
		ret_data.update({'nombre_empresa':request.POST.get('nombre_empresa')})
		query_revendedor.update({'nombre_empresa':request.POST.get('nombre_empresa')})
	else:
		errores['nombre_empresa'] = u'Falta ingresar el nombre de la empresa'
	
	if request.POST.get('direccion') != '':
		ret_data.update({'direccion':request.POST.get('direccion')})
		query_revendedor.update({'direccion':request.POST.get('direccion')})
	else:
		errores['direccion'] = u'Falta ingresar la direccion'
	
	if request.POST.get('telefono_empresa') != '':
		ret_data.update({'telefono_empresa':request.POST.get('telefono_empresa')})
		query_revendedor.update({'telefono_empresa':request.POST.get('telefono_empresa')})
	else:
		errores['telefono_empresa'] = u'Falta ingresar el telefono de la empresa'
	
	if request.POST.get('celular_empresa') != '':
		ret_data.update({'celular_empresa':request.POST.get('celular_empresa')})
		query_revendedor.update({'celular_empresa':request.POST.get('celular_empresa')})
	else:
		errores['celular_empresa'] = u'Falta ingresar el celular de la empresa'
	
	if request.POST.get('correo') != '':
		ret_data.update({'correo':request.POST.get('correo')})
		query_revendedor.update({'correo':request.POST.get('correo')})
	else:
		errores['correo'] = u'Falta ingresar el correo de la empresa'

	retorno['query_revendedor'] = query_revendedor
	retorno['errores'] = errores
	retorno['ret_data'] = ret_data
	return retorno

@login_required
def revendedor(request):
	empleado = Empleado.objects.get(usuario=request.user)
	empresa = EmpresaEmpleado.objects.get(empleado = empleado)
	revendedores = Revendedor.objects.all()
	nusuario_pk = ''
	#print(empresa.empresa.pk,'pk empresa')
	if request.method == 'POST':
		ret_data,errores,query_revendedor,recibo,ingreso_correcto = {},{},{},{},{}
		recibo = validar_revendedor(request)
		ret_data.update(recibo['ret_data'])
		errores.update(recibo['errores'])

		if request.FILES.get('logo') != None:
			query_revendedor.update({'logo':request.FILES.get('logo')})
		
		if not errores:
			#print('entra sin errores')
			logo = request.FILES.get('logo')
			#print(logo, 'logo')
			try:
				random_pass = User.objects.make_random_password(length=8, allowed_chars='abcdefghjkmnpqrstuvwxyz01234567889')
				contrasena = make_password(random_pass)
				revendedor = request.POST.get('nombre_completo').lower()
				empresa_revendedor = request.POST.get('nombre_empresa')
				correo = request.POST.get('correo')
				nombre_primera = revendedor[0].lower()
				nombre_empresa = empresa_revendedor.split(' ')
				usuario_name = nombre_primera + nombre_empresa[0]
				nusuario = User.objects.create(
					password = contrasena,
					last_login = datetime.now(),
					is_superuser = False,
					username = usuario_name,
					first_name = revendedor,
					last_name = '',
					email = correo,
					is_staff = False,
					is_active = True,
					date_joined = datetime.now()
					)
				nusuario_pk = nusuario.pk
				new_revendedor = Revendedor.objects.create(
					nombre_completo = request.POST.get('nombre_completo').upper(),
					telefono = request.POST.get('telefono'),
					nombre_empresa = request.POST.get('nombre_empresa').upper(),
					direccion = request.POST.get('direccion').upper(),
					telefono_empresa = request.POST.get('telefono_empresa'),
					celular_empresa = request.POST.get('celular_empresa'),
					correo = request.POST.get('correo'),
					logo=logo,
					usuario = nusuario,
					empresa = Empresa.objects.get(pk=empresa.empresa.pk),
					usuario_registro=request.user)
				new_cliente = Cliente.objects.create(
					nombre_completo = request.POST.get('nombre_empresa').upper(),
					celular = request.POST.get('telefono_empresa'),
					direccion = request.POST.get('direccion').upper(),
					correo = request.POST.get('correo'),
					usuario_registro = request.user,
					revendedor = True,
					revenedor_creo = new_revendedor)
				#print('paso recibo')
			except Exception as e:
				#print(e, 'error')
				errores['extra'] = e
				transaction.rollback()
				
				ctx = {'revendedores':revendedores,'ret_data':ret_data,'errores':errores}
				return render(request,'revendedor.html',ctx)
			else:	
				transaction.commit()
				name = 'Bienvenido al sistema de ' + empresa.empresa.nombre_empresa
				subject = 'Nota de registro'
				to_email = nusuario.email
				mensaje = 'Estimado Sr./Sra. '+ new_revendedor.nombre_completo.upper()  + ' las credenciales para ingreso del sistema son las siguientes: '+'\n'+'Usuario: '+ nusuario.username +'\n'+'Contraseña: '+random_pass
				send_mail(subject, mensaje,to_email, [nusuario.email])
				##MENSAJE DE TEXTO
				message = client.messages.create(
					body='Bienvenido al sistema de ' + empresa.empresa.nombre_empresa + '\n' + 'Estimado Sr./Sra. '+new_revendedor.nombre_completo.upper()  + ' las credenciales para ingreso del sistema son las siguientes: '+'\n'+'Usuario: '+ nusuario.username +'\n'+'Contraseña: '+random_pass,
					from_='+13473345592',
					to='+1' + new_revendedor.telefono
                          )
				##print(message.sid)
				#send_message('Bienvenido al sistema de ' + empresa.empresa.nombre_empresa + 'Estimado Sr./Sra. '+new_revendedor.nombre_completo.upper()  + ' las credenciales para ingreso del sistema son las siguientes: '+'\n'+'Usuario: '+ nusuario.username +'\n'+'Contraseña: '+random_pass , '+1' + new_revendedor.telefono)
				ingreso_correcto['mensaje'] = u"Datos guardados correctamente."
				ctx = {'ingreso_correcto':ingreso_correcto,'revendedores':revendedores}
				return HttpResponseRedirect(reverse('revendedor'))
		else:
			borrar = User.objects.get(pk=nusuario_pk).delete()
			#print errores,'erroresdentro de errores'
			ctx = {'ret_data':ret_data,'errores':errores,'revendedores':revendedores}
			return render(request,'revendedor.html',ctx)
	elif request.method == 'GET':
		ctx = {'revendedores':revendedores}
		return render(request,'revendedor.html',ctx)

@login_required
def desactivar_usuario(request,id):
	usuario = User.objects.filter(pk = id).update(is_active = False)
	return HttpResponseRedirect(reverse('empleados'))

@login_required
def activar_usuario(request,id):
	usuario = User.objects.filter(pk = id).update(is_active = True)
	return HttpResponseRedirect(reverse('empleados'))

@login_required
def aprobar_envios(request):
	envios = Envio.objects.filter(revendedor=True,aprobado=False)
	dic_envios = []
	for e in envios:
		envio = {}
		envio['pk'] = e.pk
		envio['codigo'] = str(e.codigo)
		envio['revendedor'] = Revendedor.objects.get(usuario=e.usuario_registro)
		envio['guia_revendedor'] = e.guia_revendedor
		envio['quien_envia'] = e.quien_envia.nombre_completo
		envio['destino'] = e.pais_destino.nombre + " | " + e.departamento_destino.nombre
		dic_envios.append(envio)
	ctx = {'envios':dic_envios}
	return render(request,'aprobar_revendedor.html',ctx)

@login_required
def aprobar_post(request):
	if request.method == 'POST':
		envios = request.POST.getlist('mover')
		try:
			for e in envios:
				#print e,'error arriba'
				envio = Envio.objects.filter(pk=e).update(aprobado=True, usuario_aprobo=request.user)
				#print 'va hacer la redireccion'
			return HttpResponseRedirect(reverse('aprobar_envios')+"?ok")
		except Exception as e:
			print (e, 'error')
	elif request.method == 'GET':
		return datos_envio
####MODULO REVENDEDOR END

@login_required
def departamento(request):
	#print "mando pais"
	departamentos = list(Departamento.objects.filter(pais=request.GET.get('idpais')).values('id','nombre'))
	return HttpResponse(simplejson.dumps(departamentos), content_type='application/javascript')
	connection.close()

@login_required
def municipio(request):
	municipios = list(Municipio.objects.filter(departamento=request.GET.get('iddepartamento')).values('id','nombre'))
	return HttpResponse(simplejson.dumps(municipios), content_type='application/javascript')
	connection.close()

@login_required
def validad_envio(request,quien_envia,quien_recibe,pais,departamento):
	query_envio,errores,ret_data,retorno = {},{},{},{}
	if request.POST.get('quien_envia', '') and request.POST.get('quien_envia', '') != '0':
		ret_data['id_quien_envia'] = int(request.POST.get('quien_envia',''))
		quien_envia_query = quien_envia.get(pk=request.POST.get('quien_envia',''))
		query_envio['quien_envia'] = quien_envia_query
	else:
		errores['quien_envia'] = u'Falta ingresar la persona que realiza el envío'
	
	if request.POST.get('quien_recibe', '') and request.POST.get('quien_recibe', '') != '0':
		ret_data['id_quien_recibe'] = int(request.POST.get('quien_recibe',''))
		quien_recibe_query = quien_recibe.get(pk=request.POST.get('quien_recibe',''))
		query_envio['quien_recibe'] = quien_recibe_query
	else:
		errores['quien_recibe'] = u'Falta ingresar la persona que realiza el envío'
	if request.POST.get('pais', '') and request.POST.get('pais', '') != '0':
		ret_data['id_pais'] = int(request.POST.get('pais',''))
		pais_query = pais.get(pk=request.POST.get('pais',''))
		query_envio['pais_destino'] = pais_query
	else:
		errores['pais'] = u'Elija una opción válida (País)'

	if request.POST.get('departamento', '') and request.POST.get('departamento', '') != '0':
		ret_data['id_departamento'] = int(request.POST.get('departamento',''))
		departamento_query = departamento.get(pk=request.POST.get('departamento',''))
		query_envio['departamento_destino'] = departamento_query
	else:
		errores['departamento'] = u'Elija una opción válida (Departamento)'
	
	if request.POST.get('direccion_registrar') != '':
		ret_data.update({'direccion_registrar':request.POST.get('direccion_registrar').upper()})
		query_envio.update({'direccion_registrar':request.POST.get('direccion_registrar').upper()})
	else:
		errores['direccion_registrar'] = u'Falta ingresar la dirección de envío'

	if request.POST.get('valor_envio') != '' or request.POST.get('valor_envio', '') != '0' :
		ret_data.update({'valor_envio':request.POST.get('valor_envio')})
		query_envio.update({'valor_envio':request.POST.get('valor_envio')})
	else:
		errores['valor_envio'] = u'Falta ingresar precio de envío'

	if request.POST.get('celular_registrar') != '':
		ret_data.update({'celular_registrar':request.POST.get('celular_registrar')})
		query_envio.update({'celular_registrar':request.POST.get('celular_registrar')})
	else:
		errores['celular_registrar'] = u'Falta ingresar los teléfonos del consignatario'

	if request.POST.get('pago_recibido') != '' or request.POST.get('pago_recibido', '') != '0' :
		ret_data.update({'pago_recibido':request.POST.get('pago_recibido')})
		query_envio.update({'pago_recibido':request.POST.get('pago_recibido')})
	else:
		errores['pago_recibido'] = u'Falta ingresar valor del pago'

	retorno['query_envio'] = query_envio
	retorno['errores'] = errores
	retorno['ret_data'] = ret_data
	return retorno

@login_required
def registrar_envio(request):
	empleado = Empleado.objects.get(usuario=request.user)
	empresa = EmpresaEmpleado.objects.get(empleado = empleado)
	r = Revendedor.objects.filter(usuario=request.user)
	pais = Pais.objects.all()
	ver_clientes = Cliente.objects.all()
	departamento = Departamento.objects.all()
	quien_recibe = ClienteRecibe.objects.all()
	es_revendedor = False
	correlativo = Envio.objects.filter(empresa=empresa.empresa)
	#print correlativo.count(), 'correlativo'
	tipo_contenido = TipoContenido.objects.all()
	tipo_envio = TipoEnvio.objects.all()
	if r.count() >= 1:
		revendedor_creo = Revendedor.objects.get(usuario = request.user)
		es_revendedor = True
		revendedor_creo = Revendedor.objects.get(usuario = request.user)
		quien_envia = Cliente.objects.filter(empresa=empresa.empresa,revendedor = True,revenedor_creo=revendedor_creo.pk)
	else:
		quien_envia = Cliente.objects.filter(empresa=empresa.empresa)

	#imagen = request.FILES.get('imagen')
	if request.method == 'POST':
		valor_adicional = 0
		valor_emplasticado = 0
		valor_seguro = 0
		codigo_inicial = 2006
		correlativo = Envio.objects.filter(empresa=empresa.empresa)
		codigo_envio = int(codigo_inicial) + int(correlativo.count())
		codigo_empresa = empresa.empresa.codigo_empresa
		codigo_final = codigo_empresa + '00' + str(codigo_envio)
		ret_data,errores,query_envio,recibo,ingreso_correcto = {},{},{},{},{}
		recibo = validad_envio(request,quien_envia,quien_recibe,pais,departamento)
		ret_data.update(recibo['ret_data'])
		errores.update(recibo['errores'])
		query_envio.update({'empresa':empresa.empresa})
		query_envio.update({'codigo':codigo_final})
		query_envio.update(recibo['query_envio'])
		if request.POST.get('valor_adicional') != '':
			valor_adicional = request.POST.get('valor_adicional')
			ret_data.update({'precioadicional':valor_adicional})
			query_envio.update({'valor_adicional':valor_adicional})

		if request.POST.get('valor_emplasticado') != '':
			valor_emplasticado = request.POST.get('valor_emplasticado')
			ret_data.update({'valor_emplasticado':valor_emplasticado})
			query_envio.update({'valor_emplasticado':valor_emplasticado})
		
		if request.POST.get('valor_seguro') != '':
			valor_seguro = request.POST.get('valor_seguro')
			ret_data.update({'valor_seguro':valor_seguro})
			query_envio.update({'valor_seguro':valor_seguro})

		if request.POST.get('valor_envio') == '':
			errores.update({'valor_envio':'DEBE REALIZAR EL CALCULO DEL ENVIO'})
		else:
			if float(request.POST.get('valor_envio')) <=0:
				errores.update({'valor_envio':'DEBE REALIZAR EL CALCULO DEL ENVIO O AGREGAR CAJAS'})
		if not errores:
			valor = float(request.POST.get('valor_envio'))
			total = valor + float(valor_adicional) + float(valor_emplasticado) + float(valor_seguro)
			pago_recibido = float(request.POST.get('pago_recibido'))
			
			if pago_recibido < total:
				credito = True
				saldo_pendiente = (total - pago_recibido)
			else:
				credito = False
				saldo_pendiente = 0.00

			if es_revendedor:
				if request.POST.get('guia_revendedor') != '':
					ret_data.update({'guia_revendedor':request.POST.get('guia_revendedor')})
					query_envio.update({'guia_revendedor':request.POST.get('guia_revendedor').upper()})
					query_envio.update({'revendedor':True})
					query_envio.update({'aprobado':False})
				else:
					errores['guia_revendedor'] = u'Falta ingresar la guia de revendedor'

			query_envio.update({'guia_revendedor':request.POST.get('guia_revendedor').upper()})
			if request.POST.get('tipo_contenido', '') and request.POST.get('tipo_contenido', '') != '0':
				ret_data['id_tipo_contenido'] = int(request.POST.get('tipo_contenido',''))
				tipo_contenido_query = tipo_contenido.get(pk=request.POST.get('tipo_contenido',''))
				query_envio['tipo_contenido'] = tipo_contenido_query
			else:
				tipo_contenido_query = TipoContenido.objects.get(pk=1)
				query_envio['tipo_contenido'] = tipo_contenido_query
			
			if request.POST.get('tipo_envio', '') and request.POST.get('tipo_envio', '') != '0':
				ret_data['id_tipo_envio'] = int(request.POST.get('tipo_envio',''))
				tipo_envio_query = tipo_envio.get(pk=request.POST.get('tipo_envio',''))
				query_envio['tipo_envio'] = tipo_envio_query
			else:
				tipo_envio_query = TipoEnvio.objects.get(pk=1)
				query_envio['tipo_envio'] = tipo_envio_query


			query_envio.update({'total':total,'usuario_registro':request.user,
								'usuario_aprobo':request.user,
								'credito':credito,
								'valor_emplasticado':valor_emplasticado,
								'saldo_pendiente':saldo_pendiente,
								'comentario':request.POST.get('comentario').upper(),
								'estado_envio':EstadoEnvio.objects.get(pk=1)})
			try:
				with transaction.atomic():
					envio = Envio(**query_envio)
					envio.save()
					acum = 0
					for detalle in request.POST.getlist('cajas'):
						acum+=1
						dato = detalle.split('|')
						pk_caja = dato[0]
						cantidad = dato[1]
						pk = dato[2]
						caja_pais = CajaPais.objects.get(pk=int(pk))
						precio = caja_pais.precio
						#codigo_detalle = codigo_final + str(acum)
						#total = float(precio) * int(cantidad)
						
						acum_prueba = 0
						for x in range(0, int(cantidad)):
							acum_prueba += 1
							codigo_detalle = codigo_final + str(caja_pais.pk) +str(acum_prueba)
							total = float(precio) * int(1)
							try:
								detalle = DetalleEnvio.objects.create(envio=envio,
																	tipo_caja=caja_pais,
																	precio=precio,
																	cantidad=1,
																	codigo_orden=acum_prueba,
																	codigo=codigo_detalle,
																	total=total)
							except Exception as e:
								#print 'ERROR EN DETALLE', e
								return 0
						seguimiento = SeguimientoEnvio.objects.create(codigo_envio=Envio.objects.get(pk=envio.pk),estado=EstadoEnvio.objects.get(pk=1),empresa=empresa.empresa,usuario_registro=request.user)
						historial = HistorialEnvio.objects.create(codigo_envio=Envio.objects.get(pk=envio.pk),estado=EstadoEnvio.objects.get(pk=1),usuario_registro=request.user)
			except Exception as e:
				#print (e,'errores')
				errores['extra'] = e
				transaction.rollback()
				ctx = {'es_revendedor':es_revendedor,'pais':pais,'ret_data':ret_data,'errores':errores}
				return render(request,'registrar_envio.html',ctx)
			else:
				transaction.commit()
				#opais= Pais.objects.get(pk=request.POST.get('pais'))
				#npais = str(opais.pk) + '|' + str(opais.nombre)
				#odepto = Departamento.objects.get(pk=request.POST.get('departamento'))
				#ndepto = str(odepto.pk) + '|' + str(odepto.nombre)
				#cliente_recibe_update = ClienteRecibe.objects.filter(pk = envio.quien_recibe.pk).update(celular=request.POST.get('celular_registrar'), direccion=request.POST.get('direccion_registrar').upper(),pais=npais,departamento=ndepto)
				
				ingreso_correcto['mensaje'] = u"Datos guardados correctamente."
				ctx = {'es_revendedor':es_revendedor,'ingreso_correcto':ingreso_correcto, 'envio':envio,'tipo_contenido':tipo_contenido,'tipo_envio':tipo_envio}
				#return HttpResponseRedirect(reverse('registrar_envio')+"?ok" +"&envio="+ str(envio.pk))
				return redirect(reverse('envio_pdf', kwargs={ 'id': envio.pk }))
		else:
			#print errores,'erroresdentro de errores'
			ctx = {'es_revendedor':es_revendedor,'pais':pais,'ret_data':ret_data,'errores':errores,'quien_envia':quien_envia,'tipo_contenido':tipo_contenido,'tipo_envio':tipo_envio}
			return render(request,'registrar_envio.html',ctx)
	elif request.method == 'GET':
		ctx = {'es_revendedor':es_revendedor,
				'pais':pais,
				'quien_envia':quien_envia,
				'quien_recibe':quien_recibe,
				'tipo_contenido':tipo_contenido,
				'tipo_envio':tipo_envio,
				'empleado':empleado}
		return render(request,'registrar_envio.html',ctx)

#ejemplo de vista ver en pdf
def envio_pdf(request, id):
	envio = Envio.objects.get(pk = id)
	detalle = DetalleEnvio.objects.filter(envio=envio).order_by('pk')
	abonos = PagosCredito.objects.filter(envio=envio).aggregate(abonos_envios=Sum('pago'))
	abonos_totales = 0
	if abonos['abonos_envios'] == None:
		abonos_totales = 0
	else:
		abonos_totales = float(abonos['abonos_envios'])
	saldo = round(envio.total,2) - (float(envio.pago_recibido)+float(abonos_totales))
	detalle_guia = []
	for d in detalle:
		lista = {}
		lista['descripcion'] = d.tipo_caja.tipo_caja.descripcion
		lista['cantidad'] = d.cantidad
		barcode = get_barcode(value = d.codigo, width = 600)
		codigo = b64encode(renderPM.drawToString(barcode, fmt = 'PNG'))	
		lista['codigo'] = codigo
		detalle_guia.append(lista)
	filename = envio.empresa.logo_empresa.name.split('/')[-1]
	separar = filename.split('.')
	if separar[1] == 'pdf':
		nada = ''
	elif separar[1].lower() == 'jpg' or separar[1].lower() == 'jpeg' or separar[1].lower() == 'png':
		width, height = get_image_dimensions(envio.empresa.logo_empresa)
	#print width, height
	barcode = get_barcode(value = envio.codigo, width = 600)
	codigo = b64encode(renderPM.drawToString(barcode, fmt = 'PNG'))	
	return generar_pdf('envio_pdf.html',
						{'pagesize':'A4',
						'orientation':'landscape',
						'envio':envio,
						'codigo':codigo,
						'anchura': width, 
						'altura': height,
						'detalle':detalle,
						'detalle_guia':detalle_guia,
						'abonos':abonos_totales,
						'saldo':saldo}) 

###########GENERAR PDF##############
def generar_pdf(template_src, context_dict={}):
	template = get_template(template_src)
	#context = Context(context_dict)
	html = template.render(context_dict)
	result = BytesIO()

	pdf = pisa.pisaDocument(BytesIO(html.encode("utf-8")), result)
	if not pdf.err:
		return HttpResponse(result.getvalue(), content_type='application/pdf')
	return HttpResponse('Error al generar el PDF<pre>%s</pre>'% escape(html))

def login(request):
	mensaje = ''
	if request.user.is_authenticated:
		print
		return HttpResponseRedirect(reverse('inicio_admin'))

	if request.method == 'POST':
		username = request.POST.get('usuario')
		password = request.POST.get('pass')
		user = authenticate(username=username, password=password)
		if user is not None:
			if user.is_active:
				auth_login(request, user)
				#if user.is_superuser:
				return redirect('inicio_admin') #aqui se debe crear la vista de la parte administrativa d               
				#else:
				#    return redirect('index_cuadrilla') #esto es para los operardoes
			else:
				mensaje = 'Usuario inactivo'
		else:
			mensaje = 'Nombre de usuario o contraseña no válido.'

	return render(request, 'login.html', {'mensaje': mensaje})


def cerrar_sesion(request): 
	logout(request)
	return HttpResponseRedirect(reverse('login'))

@login_required
def editarpass(request):
	user = User.objects.get(pk=request.user.id)
	mensaje = ''
	if request.method == 'POST':
		user.password = make_password(request.POST.get('password'))
		contrasena1 = request.POST.get('password')
		contrasena2 = request.POST.get('password2')
		if contrasena1 == contrasena2:
			contrasena =  User.objects.filter(pk=request.user.id).update(password= make_password(contrasena1))
			logout(request)
			return HttpResponseRedirect(reverse('login'))
		else:
			mensaje = 'Las contraseñas no coinciden'
			ctx = {'mensaje':mensaje}
			return render(request,'cambiar_pass.html', ctx)
	else:
		return render(request, 'cambiar_pass.html')

##ultimo
@login_required
def envios_credito(request):
	credito = Envio.objects.filter(credito=True,cancelado=False).order_by('-pk')
	saldo = 0
	dic_creditos = []
	for c in credito:
		#print c.quien_envia.nombre_completo
		creditos = {}
		creditos['pk'] = c.pk
		creditos['codigo'] = c.codigo
		creditos['cliente'] = c.quien_envia.nombre_completo
		creditos['pais'] = c.pais_destino
		creditos['depto'] = c.departamento_destino
		creditos['fecha'] = c.fecha_cierre
		estado = HistorialEnvio.objects.filter(codigo_envio = c.pk).values('estado')
		lon = estado.count()
		creditos['estado']  = lon
		creditos['embarque']  = c.descripcion_embarque
		creditos['comentario']  = c.comentario
		creditos['celular'] = c.quien_envia.celular
		creditos['recibe'] = c.quien_recibe
		creditos['saldo'] = c.saldo_pendiente
		saldo += c.saldo_pendiente
		dic_creditos.append(creditos)
	#print dic_creditos
	ctx = {'creditos':dic_creditos,'saldo':saldo}
	return render(request, 'envios_credito.html', ctx)

@login_required
def tickets_envio(request):
	empleado = Empleado.objects.get(usuario=request.user)
	empresa = EmpresaEmpleado.objects.get(empleado = empleado)

	revendedor = ''
	tickets = ''
	er = Revendedor.objects.count()
	es_revendedor = False
	if er > 0:
		r = Revendedor.objects.filter(usuario=request.user)
		es_revendedor = False
		if r.count() >= 1:
			es_revendedor = True
			tickets = Envio.objects.filter(revendedor=True, usuario_registro=request.user,ticket=False)
		else:
			tickets = Envio.objects.filter(empresa=empresa.empresa,ticket=False, aprobado = True)
	else:
		revendedor = ''
		tickets = Envio.objects.filter(empresa=empresa.empresa,ticket=False, aprobado = True)

	if empleado.rol_empleado == 1 and request.user.is_superuser:
		tickets = Envio.objects.all()

	ctx = {'tickets':tickets,'empleado':empleado}
	return render(request, 'tickets.html', ctx)

@login_required
def imprimir_ticket (request, id = None):
	envio = Envio.objects.get(pk= id)
	if envio.revendedor:
		revendedor = Revendedor.objects.get(usuario=envio.usuario_registro)
	else:
		revendedor = ''
	# barcode = get_barcode(value = envio.codigo, width = 600)
	# codigo = b64encode(renderPM.drawToString(barcode, fmt = 'PNG'))
	empleado = Empleado.objects.get(usuario=request.user)
	cantTicket = 0
	detalle = DetalleEnvio.objects.filter(envio=envio)
	dic = []
	for d in detalle:
		lista = {}
		lista['guia_padre'] = d.envio.codigo
		lista['guia'] = d.codigo
		lista['producto'] = d.tipo_caja.tipo_caja.descripcion
		lista['destinatario'] = d.envio.quien_recibe.nombre_completo
		lista['departamento'] = d.envio.departamento_destino.nombre
		lista['direccion'] = d.envio.direccion_registrar
		lista['pais'] = d.envio.pais_destino.nombre
		lista['telefono'] = d.envio.quien_recibe.celular
		barcode = get_barcode(value = d.codigo, width = 600)
		codigo = b64encode(renderPM.drawToString(barcode, fmt = 'PNG'))
		lista['codigo'] = codigo
		#qr
		qr_code = generate_qr_code(d.codigo, width=180, height=180)
		lista['qr_code'] = qr_code.decode()
		
		if d.envio.comentario:
			if d.envio.comentario.strip():
				partes = d.envio.comentario.split('|')
				guia_estafeta = None
				valor_declarado = None
				comentario = None

				for parte in partes:
					if 'Guia Estafeta:' in parte:
						guia_estafeta = parte.split(':')[1].strip()
					elif 'Valor declarado' in parte:
						valor_declarado = parte.split(':')[1].strip()
					elif parte.strip() != ' ':
						comentario = parte.strip()

				lista['comentario'] = ''

				if guia_estafeta:
					lista['comentario'] += f" - Guia Estafeta: {guia_estafeta}"

				if valor_declarado:
					lista['comentario'] += f" - Valor declarado: {valor_declarado}"

				if comentario:
					lista['comentario'] += f" - Comentario: {comentario}"
		#qr
		dic.append(lista)
		cantTicket += 1
	# ctx = {'envio':envio,'codigo':codigo,'revendedor':revendedor,'lista':dic,'cantT':cantTicket}
	# return render(request,'imprimir_ticket.html',ctx)
	if revendedor:
		return generar_pdf('imprimir_ticket_rv.html',
						{
						'envio':envio,'codigo':codigo,'revendedor':revendedor,'lista':dic,'cantT':cantTicket,'empleado':empleado,
						}
		)
	else:
		return generar_pdf('imprimir_ticket.html',
							{'pagesize':'A4',
							'orientation':'landscape',
							'envio':envio,'codigo':codigo,'revendedor':revendedor,'lista':dic,'cantT':cantTicket
							}
		)

@login_required
def cerrar_ticket(request):
	if request.is_ajax():
		id = request.GET['id']
		try:
			envio = Envio.objects.filter(pk=id).update(ticket=True)
		except Exception as e:
			data = {'valor':1}
			return HttpResponse(simplejson.dumps(data), content_type='application/json')
		else:
			data = {'valor':1}
			return HttpResponse(simplejson.dumps(data), content_type='application/json')
	else:
		data = {'valor':2}
		return HttpResponse(simplejson.dumps(data), content_type='application/json')

@login_required
def pago_envio(request,id=None):
	pagos = PagosCredito.objects.filter(envio = id)
	envio = Envio.objects.get(pk = id)
	ubicacion = UbicacionEmpleado.objects.all()
	if request.method == 'POST':
		pago = request.POST.get('pago')
		pk = request.POST.get('envio')
		tp = request.POST.get('tipo_pago')
		user = request.user
		ubicacion = UbicacionEmpleado.objects.get(pk=int(request.POST.get('ubicacion')))
		if int(tp) == 0:
			tipo_pago = False
		else:
			tipo_pago = True

		pendiente = float(envio.saldo_pendiente) - float(pago)
		envio.saldo_pendiente = float(pendiente)
		#envio.save()
		pagos = float(envio.pago_recibido) + float(pago)
		npago = PagosCredito.objects.create(envio=Envio.objects.get(pk=id),pago=pago,saldo= pendiente, tipo_pago = tipo_pago, usuario_registro = user,ubicacion=ubicacion)
		if pendiente == 0:
			envio.cancelado=True
			#envios = Envio.objects.filter(pk=envio.pk).update(saldo_pendiente = pendiente,pago_recibido = pagos, credito=False)
		envio.save()
		#return envios_credito(request)
		return HttpResponseRedirect(reverse('pago_envio',kwargs={ 'id': id }))	
	elif request.method == 'GET':
		ctx = {'pagos':pagos, 'envio':envio,'ubicaciones':ubicacion}
		return render(request,'pago_envio.html',ctx)

@login_required
def imprimir_credito(request):
	credito = Envio.objects.filter(credito=True)#.values('pk','codigo','fecha_recoleccion','fecha_envio','valor_envio','valor_adicional','quien_envia','cajas','total','departamento_destino','pais_destino','pago_recibido')
	saldo = 0
	estado = HistorialEnvio.objects.filter(codigo_envio = 25).values('estado')
	lon = estado.count()
	#print lon, 'numero'
	#print estado[lon-1], 'estado'
	dic_creditos = []
	for c in credito:
		#print c.quien_envia.nombre_completo
		creditos = {}
		creditos['pk'] = c.pk,
		creditos['codigo'] = c.codigo,
		creditos['cliente'] = c.quien_envia.nombre_completo,
		creditos['pais'] = c.pais_destino,
		creditos['depto'] = c.departamento_destino,
		estado = HistorialEnvio.objects.filter(codigo_envio = c.pk).values('estado')
		lon = estado.count()
		creditos['estado']  = lon,
		creditos['embarque']  = c.cajas,
		creditos['comentario']  = c.comentario,
		creditos['celular'] = c.quien_envia.celular,
		creditos['recibe'] = c.quien_recibe,
		creditos['saldo'] = c.saldo_pendiente
		saldo += c.saldo_pendiente
		dic_creditos.append(creditos)
	#print dic_creditos
	ctx = {'creditos':dic_creditos,'saldo':saldo}
	return render(request, 'imprimir_creditos.html', ctx)

@login_required
def cierre(request):
	fecha = timezone.now()
	envios = Envio.objects.filter(cierre= False, aprobado = True)
	cajas = ReciboCaja.objects.filter(cierre = False)
	vehiculos = ReciboVehiculos.objects.filter(cierre = False)
	contenedores = ReciboContenedor.objects.filter(cierre = False)
	pagos = PagosCredito.objects.filter(cierre = False)

	saldo = 0
	saldo_cajas = 0
	saldo_vehiculos = 0
	saldo_contenedores = 0
	saldo_pagos = 0
	saldo_pagos_efectivo = 0
	dic_datos = []
	dic_cajas = []
	dic_vehiculos = []
	dic_contenedores = []
	dic_pagos = []

	for e in envios:
		datos = {}
		datos['pk'] = e.pk,
		datos['codigo'] = e.codigo
		datos['cliente'] = e.quien_envia.nombre_completo
		datos['pais'] = e.pais_destino
		datos['depto'] = e.departamento_destino
		datos['embarque']  = e.descripcion_embarque
		datos['comentario']  = e.comentario
		datos['celular'] = e.quien_envia.celular
		datos['fecha'] = e.fecha_recoleccion.strftime("%H:%M"),
		datos['saldo'] = e.saldo_pendiente
		datos['pago'] = e.pago_recibido
		saldo += e.pago_recibido
		dic_datos.append(datos)
	
	for i in cajas:
		datos = {}
		datos['codigo'] = i.recibo_no
		datos['cliente'] = i.cliente
		datos['size_caja'] = i.size_caja
		datos['valor_caja'] = i.valor_caja
		saldo_cajas += i.valor_caja
		dic_cajas.append(datos)
	
	for i in vehiculos:
		datos = {}
		datos['codigo'] = i.recibo_no
		datos['cliente'] = i.cliente
		datos['vehiculo'] = i.marca_vehiculo
		datos['modelo_vehiculo'] = i.modelo_vehiculo
		datos['valor_vehiculo'] = i.valor_vehiculo
		saldo_vehiculos += i.valor_vehiculo
		dic_vehiculos.append(datos)
	
	for i in contenedores:
		datos = {}
		datos['codigo'] = i.recibo_no
		datos['cliente'] = i.cliente
		datos['pais_destino'] = i.pais_destino
		datos['puerto_destino'] = i.puerto_destino
		datos['tamano_contenedor'] = i.tamano_contenedor
		datos['valor_contenedor'] = i.valor_contenedor
		saldo_contenedores += i.valor_contenedor
		dic_contenedores.append(datos)

	for p in pagos:
		datos = {}
		datos['factura'] = p.envio.quien_envia
		datos['pago'] = p.pago
		datos['tipo'] = p.tipo_pago
		saldo_pagos += p.pago
		if p.tipo_pago == True:
			saldo_pagos_efectivo += p.pago
		dic_pagos.append(datos)


	ctx = {'datos':dic_datos,'dic_cajas':dic_cajas,'dic_vehiculos':dic_vehiculos,'dic_contenedores':dic_contenedores,'saldo':saldo, 'fecha':fecha,'saldo_cajas':saldo_cajas,'saldo_vehiculos':saldo_vehiculos,'saldo_contenedores':saldo_contenedores, 'pagos':dic_pagos, 'saldo_pagos':saldo_pagos, 'saldo_pagos_efectivo':saldo_pagos_efectivo}
	return render(request, 'cierre.html', ctx)

@login_required
def realizar_cierre(request):
	fecha = datetime.now()
	user = request.user
	envios = Envio.objects.filter(cierre= False, aprobado = True)
	
	# cajas = ReciboCaja.objects.filter(cierre=False, usuario_registro = user)
	# vehiculos = ReciboVehiculos.objects.filter(cierre=False,usuario_registro = user)
	# contenedores = ReciboContenedor.objects.filter(cierre=False, usuario_registro = user)
	# #tareas = AgendaActividades.objects.filter(cierre = False, usuario_registro = user)
	# pagos = PagosCredito.objects.filter(cierre = False, usuario_registro = user)
	cajas = ReciboCaja.objects.filter(cierre=False)
	vehiculos = ReciboVehiculos.objects.filter(cierre=False)
	contenedores = ReciboContenedor.objects.filter(cierre=False)
	#tareas = AgendaActividades.objects.filter(cierre = False, usuario_registro = user)
	pagos = PagosCredito.objects.filter(cierre = False)
	for e in envios:
		envio = Envio.objects.filter(pk=e.pk).update(cierre=True, fecha_cierre=fecha.strftime('%Y-%m-%d'))
	for c in cajas:
		caj = ReciboCaja.objects.filter(pk=c.pk).update(cierre=True, fecha_cierre=fecha.strftime('%Y-%m-%d'))
	for v in vehiculos:
		veh = ReciboVehiculos.objects.filter(pk=v.pk).update(cierre=True, fecha_cierre=fecha.strftime('%Y-%m-%d'))
	for co in contenedores:
		con = ReciboContenedor.objects.filter(pk=co.pk).update(cierre=True, fecha_cierre=fecha.strftime('%Y-%m-%d'))
	# for e in tareas:
	# 	envio = AgendaActividades.objects.filter(pk=e.pk).update(cierre=True, fecha_cierre=fecha.strftime('%Y-%m-%d'))

	for p in pagos:
		envio = PagosCredito.objects.filter(pk=p.pk).update(cierre=True, fecha_cierre=fecha.strftime('%Y-%m-%d'))
	return cierre(request)

@login_required
def cierre_mensual(request):
	return render(request, 'cierre_mensual.html')

@login_required
def cierre_mensual_print(request):
	empleado = Empleado.objects.get(usuario=request.user)
	empresa_e = EmpresaEmpleado.objects.get(empleado = empleado)
	empresa = Empresa.objects.get(id=empresa_e.empresa_id)
	
	if request.method == 'POST':
		month = request.POST['mes']
		month = int(month)
		year = request.POST['year']
		year = int(year)
		# date_object = datetime.strptime(date_p, '%Y-%m-%d')
		# month = date_object.month
		# year = date_object.year
		nombre_mes = {1: 'Enero', 2: 'Febrero', 3: 'Marzo', 4: 'Abril', 5: 'Mayo', 6: 'Junio', 7: 'Julio', 8: 'Agosto', 9: 'Septiembre', 10: 'Octubre', 11: 'Noviembre', 12: 'Diciembre'}.get(month, 'Mes inválido')

		total_envios = 0.00
		total_cajas = 0.00
		total_vehiculos = 0.00
		total_contenedores = 0.00
		total = 0.00

		inicio = date(year, month, 1)
		fin = date(year, month, calendar.monthrange(year, month)[1])
		print(fin)
		try:
			envios = Envio.objects.filter(fecha_cierre__range = (inicio,fin), cierre=True, aprobado=True)
			cajas = ReciboCaja.objects.filter(fecha_cierre__range = (inicio,fin))
			vehiculos = ReciboVehiculos.objects.filter(fecha_cierre__range = (inicio,fin))
			contenedores = ReciboContenedor.objects.filter(fecha_cierre__range = (inicio,fin))
		except Exception as e:
			print("Error es-->",str(e))
		else:
			for e in envios:
				total_envios += e.total
			for c in cajas:
				total_cajas += c.valor_caja
			for v in vehiculos:
				total_vehiculos += v.valor_vehiculo
			for co in contenedores:
				total_contenedores += co.valor_contenedor
			# print('envios-->',total_envios)
			# print('cajas-->',total_cajas)
			# print('contenedores-->',total_contenedores)
			total = total_envios + total_cajas + total_vehiculos + total_contenedores
			# print(total)
			return generar_pdf('cierre_mensual_print.html',
						{'pagesize':'A4',
						'orientation':'landscape',
						'empresa':empresa, 'empleado':empleado,'mes':nombre_mes, 'inicio':inicio, 'fin':fin,'envios':envios, 'year':year,
						'cajas':cajas,'vehiculos':vehiculos,'contenedores':contenedores, 'total_envios':total_envios,'total_cajas':total_cajas,
						'total_vehiculos':total_vehiculos,'total_contenedores':total_contenedores,'total':total,
						}
					)

@login_required
def cierre_diario_print(request):
	fecha = datetime.now()
	user = request.user
	empleado = Empleado.objects.get(usuario=user)
	empresa_e = EmpresaEmpleado.objects.get(empleado = empleado)
	empresa = Empresa.objects.get(id=empresa_e.empresa_id)
	print(empresa)
	envios = Envio.objects.filter(cierre= False, aprobado = True)
	cajas = ReciboCaja.objects.filter(cierre = False)
	vehiculos = ReciboVehiculos.objects.filter(cierre = False)
	contenedores = ReciboContenedor.objects.filter(cierre = False)
	pagos = PagosCredito.objects.filter(cierre = False)

	saldo = 0
	saldo_cajas = 0
	saldo_vehiculos = 0
	saldo_contenedores = 0
	saldo_pagos = 0
	saldo_pagos_efectivo = 0
	dic_datos = []
	dic_cajas = []
	dic_vehiculos = []
	dic_contenedores = []
	dic_pagos = []

	for e in envios:
		datos = {}
		datos['pk'] = e.pk
		datos['codigo'] = e.codigo
		datos['cliente'] = e.quien_envia.nombre_completo
		datos['pais'] = e.pais_destino
		datos['depto'] = e.departamento_destino
		datos['embarque']  = e.descripcion_embarque
		datos['comentario']  = e.comentario
		datos['celular'] = e.quien_envia.celular
		datos['fecha'] = e.fecha_recoleccion.strftime("%H:%M")
		datos['saldo'] = e.saldo_pendiente
		datos['pago'] = e.pago_recibido
		saldo += e.pago_recibido
		dic_datos.append(datos)
	
	for i in cajas:
		datos = {}
		datos['codigo'] = i.recibo_no
		datos['cliente'] = i.cliente
		datos['size_caja'] = i.size_caja
		datos['valor_caja'] = i.valor_caja
		saldo_cajas += i.valor_caja
		dic_cajas.append(datos)
	
	for i in vehiculos:
		datos = {}
		datos['codigo'] = i.recibo_no
		datos['cliente'] = i.cliente
		datos['marca_vehiculo'] = i.marca_vehiculo
		datos['modelo_vehiculo'] = i.modelo_vehiculo
		datos['valor_vehiculo'] = i.valor_vehiculo
		saldo_vehiculos += i.valor_vehiculo
		dic_vehiculos.append(datos)
	
	for i in contenedores:
		datos = {}
		datos['codigo'] = i.recibo_no
		datos['cliente'] = i.cliente
		datos['pais_destino'] = i.pais_destino
		datos['puerto_destino'] = i.puerto_destino
		datos['tamano_contenedor'] = i.tamano_contenedor
		datos['valor_contenedor'] = i.valor_contenedor
		saldo_contenedores += i.valor_contenedor
		dic_contenedores.append(datos)

	for p in pagos:
		datos = {}
		datos['codigo'] = p.envio.codigo
		datos['factura'] = p.envio.quien_envia
		datos['pago'] = p.pago
		datos['tipo'] = p.tipo_pago
		saldo_pagos += p.pago
		if p.tipo_pago == True:
			saldo_pagos_efectivo += p.pago
		dic_pagos.append(datos)

	return generar_pdf('cierre_diario_print.html',
						{'pagesize':'A4',
						'orientation':'landscape',
						'datos':dic_datos,'dic_cajas':dic_cajas,'dic_vehiculos':dic_vehiculos,'dic_contenedores':dic_contenedores,
						'saldo':saldo, 'fecha':fecha,'saldo_cajas':saldo_cajas,'saldo_vehiculos':saldo_vehiculos,
						'saldo_contenedores':saldo_contenedores, 'pagos':dic_pagos, 'saldo_pagos':saldo_pagos, 
						'saldo_pagos_efectivo':saldo_pagos_efectivo,'empresa':empresa,'usuario':user,
						}
	)

def validar_contenedor(request,quien_envia):
	query_contenedor,errores,ret_data,retorno = {},{},{},{}
	if request.POST.get('quien_envia', '') and request.POST.get('quien_envia', '') != '0':
		ret_data['id_quien_envia'] = int(request.POST.get('quien_envia',''))
		quien_envia_query = quien_envia.get(pk=request.POST.get('quien_envia',''))
		query_contenedor['quien_envia'] = quien_envia_query
	else:
		if request.POST.get('nuevo_cliente') != '':
			ret_data.update({'nuevo_cliente':request.POST.get('nuevo_cliente')})
			query_contenedor.update({'nuevo_cliente':request.POST.get('nuevo_cliente')})
		else:
			errores['nuevo_cliente'] = u'Falta selecionar o ingresar los datos del cliente'
		
		if request.POST.get('telefono1') != '':
			ret_data.update({'telefono1':request.POST.get('telefono1')})
			query_contenedor.update({'telefono1':request.POST.get('telefono1')})
		else:
			errores['telefono1'] = u'Falta ingresar el telefono del cliente'
				
		if request.POST.get('cliente_direccion') != '':
			ret_data.update({'cliente_direccion':request.POST.get('cliente_direccion')})
			query_contenedor.update({'cliente_direccion':request.POST.get('cliente_direccion')})
		else:
			errores['cliente_direccion'] = u'Falta ingresar la direccion del cliente'

		if request.POST.get('nuevo_cliente') == '' and request.POST.get('telefono1') == '' and request.POST.get('cliente_direccion') == '':
			errores['quien_envia'] = u'Falta selecionar o ingresar los datos del cliente'

	if request.POST.get('telefono2') != '':
		ret_data.update({'telefono2':request.POST.get('telefono2')})
		query_contenedor.update({'telefono2':request.POST.get('telefono2')})
	else:
		errores['telefono2'] = u'Falta ingresar el telefono del cliente'
	
	if request.POST.get('correo') != '':
		ret_data.update({'correo':request.POST.get('correo')})
		query_contenedor.update({'correo':request.POST.get('correo')})
	else:
		errores['correo'] = u'Falta ingresar el correo del cliente'

	if request.POST.get('pasaporte') != '':
		ret_data.update({'pasaporte':request.POST.get('pasaporte')})
		query_contenedor.update({'pasaporte':request.POST.get('pasaporte')})
	else:
		errores['pasaporte'] = u'Falta ingresar el telefono del cliente'

	if request.POST.get('quien_recibe') != '':
		ret_data.update({'quien_recibe':request.POST.get('quien_recibe')})
		query_contenedor.update({'quien_recibe':request.POST.get('quien_recibe')})
	else:
		errores['quien_recibe'] = u'Falta ingresar la persona que realiza el envío'

	if request.POST.get('rnt') != '':
		ret_data.update({'rnt':request.POST.get('rnt')})
		query_contenedor.update({'rnt':request.POST.get('rnt')})
	else:
		errores['rnt'] = u'Falta ingresar el RTN/NIP de quien recibe'

	if request.POST.get('telefono_recibe') != '':
		ret_data.update({'telefono_recibe':request.POST.get('telefono_recibe')})
		query_contenedor.update({'telefono_recibe':request.POST.get('telefono_recibe')})
	else:
		errores['telefono_recibe'] = u'Falta ingresar el telefono de quien recibe'

	if request.POST.get('tamano_contenedor') != '' or request.POST.get('tamano_contenedor', '') != '0' :
		ret_data.update({'tamano_contenedor':request.POST.get('tamano_contenedor')})
		query_contenedor.update({'tamano_contenedor':request.POST.get('tamano_contenedor')})
	else:
		errores['tamano_contenedor'] = u'Falta ingresar el tamano_contenedor del contenedor'

	if request.POST.get('precio') != '' or request.POST.get('precio', '') != '0' :
		ret_data.update({'precio':request.POST.get('precio')})
		query_contenedor.update({'precio':request.POST.get('precio')})
	else:
		errores['precio'] = u'Falta ingresar el precio del contenedor'

	if request.POST.get('pais_destino') != '':
		ret_data.update({'pais_destino':request.POST.get('pais_destino')})
		query_contenedor.update({'pais_destino':request.POST.get('pais_destino')})
	else:
		errores['pais_destino'] = u'Falta ingresar el pais de destino'
	
	if request.POST.get('puerto_destino') != '':
		ret_data.update({'puerto_destino':request.POST.get('puerto_destino')})
		query_contenedor.update({'puerto_destino':request.POST.get('puerto_destino')})
	else:
		errores['puerto_destino'] = u'Falta ingresar el pais de destino'

	retorno['query_contenedor'] = query_contenedor
	retorno['errores'] = errores
	retorno['ret_data'] = ret_data
	return retorno

@login_required
def recibo_contenedor(request):
	quien_envia = Cliente.objects.all()
	contenedores = ReciboContenedor.objects.all()
	if request.method == 'POST':
		ret_data,errores,query_contenedor,recibo,ingreso_correcto = {},{},{},{},{}
		recibo = validar_contenedor(request,quien_envia)
		ret_data.update(recibo['ret_data'])
		errores.update(recibo['errores'])

		if request.FILES.get('pdf') != None:
			query_contenedor.update({'pdf':request.FILES.get('pdf')})
		
		if not errores:
			#print('entra sin errores')
			correo = request.POST.get('correo')
			pasaporte = request.POST.get('pasaporte').upper()
			quien_recibe = request.POST.get('quien_recibe')
			rtn = request.POST.get('rtn').upper()
			telefono_recibe = request.POST.get('telefono_recibe')
			tamano_contenedor = request.POST.get('tamano_contenedor').upper()
			precio = request.POST.get('precio')
			pais_destino = request.POST.get('pais_destino').upper()
			puerto_destino = request.POST.get('puerto_destino').upper()
			pdf = request.FILES.get('pdf')
			telefono2 = request.POST.get('telefono2')
			####NUEVO CLIENTE
			nuevo_cliente = request.POST.get('nuevo_cliente')
			if nuevo_cliente == '':
				cliente = Cliente.objects.get(pk=request.POST.get('quien_envia'))
			else:
				cliente = ''
			telefono1 = request.POST.get('telefono1')
			cliente_direccion = request.POST.get('cliente_direccion').upper()
			recibo_no = ReciboContenedor.objects.count() 
			recibo_no += 1
			codigo_final = 'EECT0' + str(recibo_no)
			try:
				#print('entro try')
				if cliente != '':
					#print('if try')
					new_recibo_contenedor = ReciboContenedor.objects.create(
						recibo_no=codigo_final,
						cliente=cliente,
						num_pasaporte=pasaporte,
						correo=correo,
						telefono_adicional=telefono2,
						pais_destino=pais_destino,
						puerto_destino=puerto_destino,
						nombre_recibe=quien_recibe,
						telefono_recibe=telefono_recibe,
						nip_rtn=rtn,tamano_contenedor=tamano_contenedor,
						valor_contenedor=precio,pdf=pdf,
						usuario_registro=request.user )
				else:
					#print('else try')
					nuevo_cliente = Cliente.objects.create(nombre_completo=nuevo_cliente,celular=telefono1,direccion=cliente_direccion,usuario_registro=request.user)
					#print('paso nuevo cliente')
					ultimo_cliente = Cliente.objects.last()
					#print('paso ultimo cliente')
					new_recibo_contenedor = ReciboContenedor.objects.create(recibo_no=codigo_final,cliente=Cliente.objects.get(pk=ultimo_cliente.pk),num_pasaporte=pasaporte,correo=correo,telefono_adicional=telefono2,pais_destino=pais_destino,puerto_destino=puerto_destino,nombre_recibe=quien_recibe,telefono_recibe=telefono_recibe,nip_rtn=rtn,tamano_contenedor=tamano_contenedor,valor_contenedor=precio,pdf=pdf,usuario_registro=request.user )
					#print('paso recibo')
			except Exception as e:
				#print (e,'error')
				errores['extra'] = e
				transaction.rollback()
				ctx = {'quien_envia':quien_envia,'ret_data':ret_data,'errores':errores, 'contenedores':contenedores}
				return render(request,'recibo_contenedor.html',ctx)
			else:
				transaction.commit()
				ingreso_correcto['mensaje'] = u"Datos guardados correctamente."
				ctx = {'ingreso_correcto':ingreso_correcto,'quien_envia':quien_envia,'contenedores':contenedores}
				return HttpResponseRedirect(reverse('recibo_contenedor_pdf',kwargs={ 'id': new_recibo_contenedor.pk }))
		else:
			#print errores,'erroresdentro de errores'
			ctx = {'ret_data':ret_data,'errores':errores,'quien_envia':quien_envia,'contenedores':contenedores}
			return render(request,'recibo_contenedor.html',ctx)
	elif request.method == 'GET':
		ctx = {'quien_envia':quien_envia,'contenedores':contenedores}
		return render(request,'recibo_contenedor.html',ctx)

def recibo_contenedor_pdf(request, id):
	envio = ReciboContenedor.objects.get(pk = id)
	hoy = datetime.now().date()
	barcode = get_barcode(value = id, width = 600)
	codigo = b64encode(renderPM.drawToString(barcode, fmt = 'PNG'))	
	return generar_pdf('recibo_contenedor_pdf.html',{'pagesize':'A4','orientation':'landscape','envio':envio,'codigo':codigo, 'fecha':hoy})

def ver_pdf_contenedor_pdf(request, id):
	envio = ReciboContenedor.objects.get(pk = id)
	filename = envio.pdf.name.split('/')[-1]
	separar = filename.split('.')
	if separar[1] == 'pdf':
		filepath = os.path.join(settings.MEDIA_ROOT, "archivos_pdf_contenedores", filename)
		return FileResponse(open(filepath, 'rb'), content_type='application/pdf')
	##caso ser imagen
	elif separar[1].lower() == 'jpg' or separar[1].lower() == 'jpeg' or separar[1].lower() == 'png':
		width, height = get_image_dimensions(envio.pdf)
		return render(request, 'imagen_envio.html', {'envio': envio,'opc': 1})
	else:
		###DESCARGAR AL MOMENTO DE SER OTRO ARCHIVO NO PDF
		response = HttpResponse(envio.pdf, content_type='text/plain')
		response['Content-Disposition'] = 'attachment; filename=%s' % filename
		return response

def validar_vehiculo(request,quien_envia):
	query_contenedor,errores,ret_data,retorno = {},{},{},{}
	if request.POST.get('quien_envia', '') and request.POST.get('quien_envia', '') != '0':
		ret_data['id_quien_envia'] = int(request.POST.get('quien_envia',''))
		quien_envia_query = quien_envia.get(pk=request.POST.get('quien_envia',''))
		query_contenedor['quien_envia'] = quien_envia_query
	else:
		if request.POST.get('nuevo_cliente') != '':
			ret_data.update({'nuevo_cliente':request.POST.get('nuevo_cliente')})
			query_contenedor.update({'nuevo_cliente':request.POST.get('nuevo_cliente')})
		else:
			errores['nuevo_cliente'] = u'Falta selecionar o ingresar los datos del cliente'
		
		if request.POST.get('telefono1') != '':
			ret_data.update({'telefono1':request.POST.get('telefono1')})
			query_contenedor.update({'telefono1':request.POST.get('telefono1')})
		else:
			errores['telefono1'] = u'Falta ingresar el telefono del cliente'
				
		if request.POST.get('cliente_direccion') != '':
			ret_data.update({'cliente_direccion':request.POST.get('cliente_direccion')})
			query_contenedor.update({'cliente_direccion':request.POST.get('cliente_direccion')})
		else:
			errores['cliente_direccion'] = u'Falta ingresar la direccion del cliente'

		if request.POST.get('nuevo_cliente') == '' and request.POST.get('telefono1') == '' and request.POST.get('cliente_direccion') == '':
			errores['quien_envia'] = u'Falta selecionar o ingresar los datos del cliente'

	if request.POST.get('telefono2') != '':
		ret_data.update({'telefono2':request.POST.get('telefono2')})
		query_contenedor.update({'telefono2':request.POST.get('telefono2')})
	else:
		errores['telefono2'] = u'Falta ingresar el telefono del cliente'
	
	if request.POST.get('correo') != '':
		ret_data.update({'correo':request.POST.get('correo')})
		query_contenedor.update({'correo':request.POST.get('correo')})
	else:
		errores['correo'] = u'Falta ingresar el correo del cliente'

	if request.POST.get('pasaporte') != '':
		ret_data.update({'pasaporte':request.POST.get('pasaporte')})
		query_contenedor.update({'pasaporte':request.POST.get('pasaporte')})
	else:
		errores['pasaporte'] = u'Falta ingresar el telefono del cliente'

	if request.POST.get('quien_recibe') != '':
		ret_data.update({'quien_recibe':request.POST.get('quien_recibe')})
		query_contenedor.update({'quien_recibe':request.POST.get('quien_recibe')})
	else:
		errores['quien_recibe'] = u'Falta ingresar la persona que realiza el envío'

	if request.POST.get('rnt') != '':
		ret_data.update({'rnt':request.POST.get('rnt')})
		query_contenedor.update({'rnt':request.POST.get('rnt')})
	else:
		errores['rnt'] = u'Falta ingresar el RTN/NIP de quien recibe'

	if request.POST.get('telefono_recibe') != '':
		ret_data.update({'telefono_recibe':request.POST.get('telefono_recibe')})
		query_contenedor.update({'telefono_recibe':request.POST.get('telefono_recibe')})
	else:
		errores['telefono_recibe'] = u'Falta ingresar el telefono de quien recibe'

	if request.POST.get('marca_vehiculo') != '' or request.POST.get('marca_vehiculo', '') != '0' :
		ret_data.update({'marca_vehiculo':request.POST.get('marca_vehiculo')})
		query_contenedor.update({'marca_vehiculo':request.POST.get('marca_vehiculo')})
	else:
		errores['marca_vehiculo'] = u'Falta ingresar la marca del vehiculo'

	if request.POST.get('modelo_vehiculo') != '' or request.POST.get('modelo_vehiculo', '') != '0' :
		ret_data.update({'modelo_vehiculo':request.POST.get('modelo_vehiculo')})
		query_contenedor.update({'modelo_vehiculo':request.POST.get('modelo_vehiculo')})
	else:
		errores['modelo_vehiculo'] = u'Falta ingresar el modelo del vehiculo'
	
	if request.POST.get('vin_vehiculo') != '' or request.POST.get('vin_vehiculo', '') != '0' :
		ret_data.update({'vin_vehiculo':request.POST.get('vin_vehiculo')})
		query_contenedor.update({'vin_vehiculo':request.POST.get('vin_vehiculo')})
	else:
		errores['modelo_vehiculo'] = u'Falta ingresar el vin del vehiculo'
	
	if request.POST.get('valor') != '' or request.POST.get('valor', '') != '0' :
		ret_data.update({'valor':request.POST.get('valor')})
		query_contenedor.update({'valor':request.POST.get('valor')})
	else:
		errores['valor'] = u'Falta ingresar el valor del vehiculo'

	if request.POST.get('pais_destino') != '':
		ret_data.update({'pais_destino':request.POST.get('pais_destino')})
		query_contenedor.update({'pais_destino':request.POST.get('pais_destino')})
	else:
		errores['pais_destino'] = u'Falta ingresar el pais de destino'
	
	if request.POST.get('puerto_destino') != '':
		ret_data.update({'puerto_destino':request.POST.get('puerto_destino')})
		query_contenedor.update({'puerto_destino':request.POST.get('puerto_destino')})
	else:
		errores['puerto_destino'] = u'Falta ingresar el pais de destino'

	retorno['query_contenedor'] = query_contenedor
	retorno['errores'] = errores
	retorno['ret_data'] = ret_data
	return retorno

@login_required
def recibo_vehiculo(request):
	quien_envia = Cliente.objects.all()
	vehiculos = ReciboVehiculos.objects.all()
	if request.method == 'POST':
		ret_data,errores,query_contenedor,recibo,ingreso_correcto = {},{},{},{},{}
		recibo = validar_vehiculo(request,quien_envia)
		ret_data.update(recibo['ret_data'])
		errores.update(recibo['errores'])

		if request.FILES.get('pdf') != None:
			query_contenedor.update({'pdf':request.FILES.get('pdf')})
		
		if not errores:
			#print('entra sin errores')
			correo = request.POST.get('correo')
			pasaporte = request.POST.get('pasaporte').upper()
			quien_recibe = request.POST.get('quien_recibe')
			rtn = request.POST.get('rtn').upper()
			telefono_recibe = request.POST.get('telefono_recibe')
			marca_vehiculo = request.POST.get('marca_vehiculo').upper()
			modelo_vehiculo = request.POST.get('modelo_vehiculo').upper()
			vin_vehiculo = request.POST.get('vin_vehiculo').upper()
			valor_vehiculo = request.POST.get('valor_vehiculo')
			pais_destino = request.POST.get('pais_destino').upper()
			puerto_destino = request.POST.get('puerto_destino').upper()
			pdf = request.FILES.get('pdf')
			#print(pdf, 'pdf')
			telefono2 = request.POST.get('telefono2')
			####NUEVO CLIENTE
			nuevo_cliente = request.POST.get('nuevo_cliente')
			if nuevo_cliente == '':
				cliente = Cliente.objects.get(pk=request.POST.get('quien_envia'))
			else:
				cliente = ''
			telefono1 = request.POST.get('telefono1')
			cliente_direccion = request.POST.get('cliente_direccion').upper()
			recibo_no = ReciboVehiculos.objects.count() 
			recibo_no += 1
			codigo_final = 'EECR0' + str(recibo_no)
			try:
				#print('entro try')
				if cliente != '':
					#print('if try')
					new_recibo_vehiculo = ReciboVehiculos.objects.create(recibo_no=codigo_final,cliente=cliente,num_pasaporte=pasaporte,correo=correo,telefono_adicional=telefono2,pais_destino=pais_destino,puerto_destino=puerto_destino,nombre_recibe=quien_recibe,telefono_recibe=telefono_recibe,nip_rtn=rtn,marca_vehiculo=marca_vehiculo,modelo_vehiculo=modelo_vehiculo,vin_vehiculo=vin_vehiculo,valor_vehiculo=valor_vehiculo,pdf=pdf,usuario_registro=request.user )
				else:
					#print('else try')
					nuevo_cliente = Cliente.objects.create(nombre_completo=nuevo_cliente,celular=telefono1,direccion=cliente_direccion,usuario_registro=request.user)
					#print('paso nuevo cliente')
					ultimo_cliente = Cliente.objects.last()
					#print('paso ultimo cliente')
					new_recibo_vehiculo = ReciboVehiculos.objects.create(recibo_no=codigo_final,cliente=Cliente.objects.get(pk=ultimo_cliente.pk),num_pasaporte=pasaporte,correo=correo,telefono_adicional=telefono2,pais_destino=pais_destino,puerto_destino=puerto_destino,nombre_recibe=quien_recibe,telefono_recibe=telefono_recibe,nip_rtn=rtn,marca_vehiculo=marca_vehiculo,modelo_vehiculo=modelo_vehiculo,vin_vehiculo=vin_vehiculo,valor_vehiculo=valor_vehiculo,pdf=pdf,usuario_registro=request.user )
					#print('paso recibo')
			except Exception as e:
				errores['extra'] = e
				transaction.rollback()
				ctx = {'quien_envia':quien_envia,'vehiculos':vehiculos,'ret_data':ret_data,'errores':errores}
				return render(request,'recibo_vehiculo.html',ctx)
			else:
				transaction.commit()
				ingreso_correcto['mensaje'] = u"Datos guardados correctamente."
				ctx = {'ingreso_correcto':ingreso_correcto,'quien_envia':quien_envia,'vehiculos':vehiculos}
				return HttpResponseRedirect(reverse('recibo_vehiculo_pdf',kwargs={ 'id': new_recibo_vehiculo.pk }))
		else:
			#print errores,'erroresdentro de errores'
			ctx = {'ret_data':ret_data,'errores':errores,'quien_envia':quien_envia,'vehiculos':vehiculos}
			return render(request,'recibo_vehiculo.html',ctx)
	elif request.method == 'GET':
		ctx = {'quien_envia':quien_envia,'vehiculos':vehiculos}
		return render(request,'recibo_vehiculo.html',ctx)

def recibo_vehiculo_pdf(request, id):
	envio = ReciboVehiculos.objects.get(pk = id)
	hoy = datetime.now().date()
	barcode = get_barcode(value = id, width = 600)
	codigo = b64encode(renderPM.drawToString(barcode, fmt = 'PNG'))
	return generar_pdf('recibo_vehiculo_pdf.html',{'pagesize':'A4','orientation':'landscape','envio':envio,'codigo':codigo,'fecha':hoy})

def ver_pdf_vehiculo_pdf(request, id):
	envio = ReciboVehiculos.objects.get(pk = id)
	filename = envio.pdf.name.split('/')[-1]
	separar = filename.split('.')
	if separar[1] == 'pdf':
		filepath = os.path.join(settings.MEDIA_ROOT, "archivos_pdf_vehiculos", filename)
		return FileResponse(open(filepath, 'rb'), content_type='application/pdf')
	##caso ser imagen
	elif separar[1].lower() == 'jpg' or separar[1].lower() == 'jpeg' or separar[1].lower() == 'png':
		width, height = get_image_dimensions(envio.pdf)
		return render(request, 'imagen_envio.html', {'envio': envio,'opc': 1})
	else:
		###DESCARGAR AL MOMENTO DE SER OTRO ARCHIVO NO PDF
		response = HttpResponse(envio.pdf, content_type='text/plain')
		response['Content-Disposition'] = 'attachment; filename=%s' % filename
		return response

@login_required
def top_clientes(request):
	clientes = Cliente.objects.filter(revendedor=False)
	dic_clientes = []

	for e in clientes:
		cliente = {}
		cliente['cliente'] = e.nombre_completo
		cliente['celular'] = e.celular
		cliente['cantidad'] = Envio.objects.filter(quien_envia= e.pk,revendedor=False).count()
		dic_clientes.append(cliente)

	ctx = {'dic_clientes':dic_clientes}
	return render(request,'top_clientes.html', ctx)


###cammbios fran
@login_required
def cajas_tipo(request):
	cajas_pais = list(CajaPais.objects.filter(pais=request.GET.get('idpais')).values('tipo_caja__pk','tipo_caja__descripcion','precio','pk','empresa'))
	#cajas_pais = CajaPais.objects.filter(pais=request.GET.get('idpais'))
	#print cajas_pais
	return JsonResponse(cajas_pais, safe = False)
	connection.close()

@login_required()
def modal_agregar_recibe(request):
	if request.method == 'POST':
		#print "lol es recibe pk",request.POST['envia_pk']

		if int(request.POST['envia_pk']) == 0:
			#print "es 0"
			return JsonResponse({'option': 'error','detalle_error':'DEBES ELEGIR QUIEN ENVIA'})
		else:
			if request.POST['nombre_recibe'] == '' or request.POST['telefono_recibe'] == '' or request.POST['direccion_recibe'] == '':
				return JsonResponse({'option': 'error','detalle_error':'INGRESA TODOS LOS DATOS SOLICITADOS DE QUIEN RECIBE'})
			else:
				envia = Cliente.objects.get(pk=(request.POST['envia_pk']))
				recibe = ClienteRecibe.objects.create(cliente_envia = envia, 
														nombre_completo = request.POST['nombre_recibe'],
														celular = request.POST['telefono_recibe'],
														direccion=request.POST['direccion_recibe'],
														usuario_registro=request.user
													)
				
				option = '''
					<option value="{}" selected>{}</option>
				'''.format(recibe.pk, recibe.nombre_completo+'|'+recibe.celular)
				return JsonResponse({'option': option})
	else:
		#print "no post"
		return render(request, '404.html')


@login_required()
def modal_editar_recibe(request):
	if request.method == 'POST':
		if request.POST['recibe_pk_modal_Edit'] == '':
			return JsonResponse({'option': 'error','detalle_error':'DEBES ELEGIR QUIEN ENVIA'})
		else:
			if request.POST['nombre_recibe_edit'] == '' or request.POST['telefono_recibe_edit'] == '' or request.POST['direccion_recibe_edit'] == '':
				return JsonResponse({'option': 'error','detalle_error':'INGRESA TODOS LOS DATOS SOLICITADOS DE QUIEN RECIBE'})
			else:
				#print request.POST['recibe_pk_modal_Edit'],'recibe_pk_modal_Edit'
				recibe = ClienteRecibe.objects.filter(pk=request.POST['recibe_pk_modal_Edit']).update(
													   nombre_completo = request.POST['nombre_recibe_edit'],
														celular = request.POST['telefono_recibe_edit'],
														direccion=request.POST['direccion_recibe_edit'],
														usuario_registro=request.user
														)
				recibe_ = ClienteRecibe.objects.get(pk=request.POST['recibe_pk_modal_Edit'])
				option = '''
					<option value="{}" selected>{}</option>
				'''.format(recibe_.pk, recibe_.nombre_completo+'|'+recibe_.celular)
				return JsonResponse({'option': option})
	else:
		return render(request, '404.html')

@login_required()
def modal_agregar_tipo_caja(request):
	if request.method == 'POST':
		if int(request.POST['pais_caja']) == 0:
			return JsonResponse({'option': 'error','detalle_error':'DEBES ELEGIR PAIS'})
		else:
			if int(request.POST['pais_caja']) == 0 or request.POST['descripcion_caja'] == '' or request.POST['largo'] == '' or request.POST['ancho'] == '' or request.POST['alto'] == '' or request.POST['pie_cubico'] == '' or request.POST['precio_caja'] == '':
				return JsonResponse({'option': 'error','detalle_error':'INGRESE TODOS LOS DATOS DEL TIPO DE CAJA'})
			else:
				tipo_caja_q=TipoCaja.objects.create(descripcion=request.POST['descripcion_caja'],
												largo=request.POST['largo'],
												ancho=request.POST['ancho'],
												alto=request.POST['alto'],
												espacio_cubico=request.POST['pie_cubico'],
												usuario_registro=request.user)
				caja_pais = CajaPais.objects.create(pais=Pais.objects.get(pk=request.POST['pais_caja']),
													tipo_caja=tipo_caja_q,
													precio=request.POST['precio_caja'],
													usuario_registro=request.user)
				option = '''
					<option value="{}|{}|{}|{}" selected>{} | Precio => {}</option>
				'''.format(caja_pais.tipo_caja.pk, caja_pais.tipo_caja.descripcion,caja_pais.precio,caja_pais.pk,caja_pais.tipo_caja.descripcion,caja_pais.precio,)
				return JsonResponse({'option': option})
	else:
		return render(request, '404.html')


@login_required()
def modal_editar_tipo_caja(request):
	if request.method == 'POST':
		if request.POST['pk_caja_edit'] == '':
			return JsonResponse({'option': 'error','detalle_error':'DEBES SELECCIONA UN PRODUCTO A EDITAR'})
		else:
			if int(request.POST['pk_caja_edit']) == 0 or request.POST['descripcion_caja_edi'] == '' or request.POST['largo_edi'] == '' or request.POST['ancho_edi'] == '' or request.POST['alto_edi'] == '' or request.POST['pie_cubico_edi'] == '' or request.POST['precio_caja_edi'] == '':
				return JsonResponse({'option': 'error','detalle_error':'INGRESE TODOS LOS DATOS DEL TIPO DE CAJA'})
			else:
				caja_pais = CajaPais.objects.get(pk=int(request.POST['pk_caja_edit']))
				caja_pais.precio = float(request.POST['precio_caja_edi'])
				caja_pais.save()
				tipo_caja_q = TipoCaja.objects.get(pk=caja_pais.tipo_caja.pk)
				tipo_caja_q.descripcion = request.POST['descripcion_caja_edi']
				tipo_caja_q.largo = request.POST['largo_edi']
				tipo_caja_q.ancho = request.POST['ancho_edi']
				tipo_caja_q.alto = request.POST['alto_edi']
				tipo_caja_q.espacio_cubico = request.POST['pie_cubico_edi']
				tipo_caja_q.save()
				
				caja_pais = CajaPais.objects.get(pk=int(request.POST['pk_caja_edit']))
				option = '''
					<option value="{}|{}|{}|{}" selected>{} | Precio => {}</option>
				'''.format(caja_pais.tipo_caja.pk, caja_pais.tipo_caja.descripcion,caja_pais.precio,caja_pais.pk,caja_pais.tipo_caja.descripcion,caja_pais.precio,)
				#print option,"aqui fran"
				return JsonResponse({'option': option})
	else:
		return render(request, '404.html')

@login_required()
def cargar_contenedor(request):
	pais = Pais.objects.all()
	contenedores = Contenedor.objects.all().order_by('-pk')
	if request.method == 'POST':
		if request.POST['codigo_original'] == '' or request.POST['codigo_express'] == '' or request.POST['pies_cubicos'] == '':
			return JsonResponse({'option': 'error','detalle_error':'DEBES INGRESAR TODOS LOS DATOS'})
		else:
			if int(request.POST['pies_cubicos']) == 0:
				return JsonResponse({'option': 'error','detalle_error':'DEBES INGRESAR PIES CUBICOS'})
			else:
				contenedor = Contenedor.objects.create(codigo_original=request.POST['codigo_original'],
															codigo_express=request.POST['codigo_express'],
															pais_destino=Pais.objects.get(pk=int(request.POST['pais_destino'])),
															pies_cubicos=request.POST['pies_cubicos'],
													    	estado=EstadoEnvio.objects.get(pk=1),
														)
				historial_contenedor = HistorialContenedor.objects.create(contenedor=contenedor,
																			estado=EstadoEnvio.objects.get(pk=1),
																			comentario='INGRESO A BODEGA DE HOUSTON',
																			usuario_registro=request.user)
				#return HttpResponseRedirect(reverse('ver_contenedor_enviar',kwargs={ 'id': contenedor.pk }))
				return JsonResponse({'option': 'save','url':reverse('ver_contenedor_enviar',kwargs={ 'id': contenedor.pk })})
	else:
		data = {'pais':pais,'contenedores':contenedores}
		return render(request, 'cargar_contenedor.html',data)

#@minified_response
@login_required()
def trasladar_guia(request,id_envio,guia_hija):
	if request.method== 'POST':
		envio = Envio.objects.get(pk=id_envio)
		contenedor = envio.contenedor.pk
		empresa = Empresa.objects.get(pk=envio.empresa.pk)
		#print request.POST['estado_envio_guia']
		if int(request.POST['estado_envio_guia'])==1:
			#print 'opcion 1'
			try:
				estado = EstadoEnvio.objects.get(pk=int(request.POST['estado_envio_guia']))
				#print estado,'estado'
				update = Envio.objects.filter(pk=id_envio).update(contenedor='',estado_envio=estado)
				detalle = DetalleEnvio.objects.filter(codigo=guia_hija).update(fue_enviada=False, estado_hija=estado)
				historial = HistorialEnvio.objects.create(codigo_envio=envio,
														estado=estado,
														usuario_registro=request.user)
				seguimiento = SeguimientoEnvio.objects.create(codigo_envio=envio,estado=estado,empresa=empresa,comentario='REVERSION',usuario_registro=request.user)
				#print 'guardo todo'
			except Exception as e:
				#print 'error todo'
				return 0
		else:
			#print 'opcion 2'
			estado = EstadoEnvio.objects.get(pk=int(request.POST['estado_envio_guia']))
			detalle = DetalleEnvio.objects.filter(codigo=guia_hija).update(fue_enviada=True, estado_hija=estado)
			historial = HistorialEnvio.objects.create(codigo_envio=envio,
													estado=estado,
													usuario_registro=request.user)
			update = Envio.objects.filter(pk=id_envio).update(contenedor=Contenedor.objects.get(pk=contenedor),estado_envio=estado)
			seguimiento = SeguimientoEnvio.objects.create(codigo_envio=envio,estado=estado,empresa=empresa,comentario='REVERSION',usuario_registro=request.user)
		
		#se modifica todo el datelle
		empresa = Envio.objects.get(pk=envio.pk)
		###UPDATE KRAKEN
		#OBTENER ENVIO EN kraken_cargo
		es_kraken = False
		kraken_envio = SistemaEmpresaenvio.objects.using('kraken_cargo').filter(codigo = empresa.guia_revendedor)
		#print kraken_envio.count(),'kraken_envio'
		if kraken_envio.count() >= 1:
			es_kraken = True
		
		if es_kraken:
			estadoeehn = EstadoEnvio.objects.get(pk=int(request.POST['estado_envio_guia']))
			estado_kraken = ''
			pk = 0
			if estadoeehn.pk == 1:
				estado_kraken = 'BODEGA EEUU'
				pk = 8
			elif estadoeehn.pk == 2:
				estado_kraken = 'PUERTO EEUU'
				pk = 9
			elif estadoeehn.pk == 3:
				estado_kraken = 'TRANSITO MARITIMO'
				pk = 10
			elif estadoeehn.pk == 4:
				estado_kraken = 'PUERTO HONDURAS'
				pk = 11
			elif estadoeehn.pk == 5:
				estado_kraken = 'BODEGA HONDURAS'
				pk = 12
			elif estadoeehn.pk == 6:
				estado_kraken = 'EN TRANSITO PARA ENTREGA'
				pk = 13
			elif estadoeehn.pk == 7:
				estado_kraken = 'ENTREGADO'
				pk = 14
			estadokraken = SistemaEmpresaestadoenvio.objects.using('kraken_cargo').get(pk=pk)
			seguimiento_kraken_cargo = SistemaSeguimientoenvio.objects.db_manager('kraken_cargo').filter(codigo_envio=SistemaEmpresaenvio.objects.using('kraken_cargo').get(codigo =  envio.guia_revendedor)).update(estado=SistemaEmpresaestadoenvio.objects.using('kraken_cargo').get(pk=pk))
			historial_kraken_cargo = SistemaHistorialenvio.objects.db_manager('kraken_cargo').create(codigo_envio=SistemaEmpresaenvio.objects.using('kraken_cargo').get(codigo =  envio.guia_revendedor),estado=SistemaEmpresaestadoenvio.objects.using('kraken_cargo').get(pk=pk),usuario_registro=AuthUser.objects.using('kraken_cargo').get(pk=14),fechahora = timezone.now())
			envio_kraken = SistemaEmpresaenvio.objects.using('kraken_cargo').get(codigo =  envio.guia_revendedor)
			hoy = timezone.now()
		#return JsonResponse({'option': 'save','url':reverse('ver_contenedor_enviar',kwargs={ 'id': contenedor })})
		return HttpResponseRedirect(reverse('ver_contenedor_enviar',args = (int(contenedor), ))+"?reversion")	
		
	elif request.method == 'GET':
		return render(request, '404.html')


#@minified_response
@login_required()
def ver_contenedor_enviar(request,id):
	contenedor = Contenedor.objects.get(pk=id)
	detalle = DetalleEnvio.objects.filter(envio__contenedor=contenedor,fue_enviada=True)
	espacio_ocupado = DetalleEnvio.objects.filter(envio__contenedor=contenedor).aggregate(total = Sum("tipo_caja__tipo_caja__espacio_cubico"))
	historial_contenedor = HistorialContenedor.objects.filter(contenedor=contenedor)
	url = reverse('ver_contenedor_enviar',args = (contenedor.pk, ))
	disponible = 0
	estados = EstadoEnvio.objects.filter(pk__in=(1,2,3,4,5))
	if espacio_ocupado['total'] == None:
		disponible = contenedor.pies_cubicos
	else:
		disponible = contenedor.pies_cubicos - float(espacio_ocupado['total'])
	error = ''
	if request.method == 'POST':
		try:
			with transaction.atomic():

				guia_hija = request.POST['guia']
				if disponible <=0:
					error = 'Ya no hay espacio disponible'
					data = {'contenedor':contenedor,
							'detalle_contenedor':detalle,
							'disponible':round(disponible,2),
							'error':error,
							'estados':estados,
							'historial_contenedor':historial_contenedor,
							'url':url}
					return render(request, 'ver_contenedor_enviar.html',data)
				elif request.POST['guia'] == '':
					error = 'Debe ingresar la guía'
					data = {'contenedor':contenedor,
							'detalle_contenedor':detalle,
							'disponible':round(disponible,2),
							'error':error,
							'estados':estados,
							'historial_contenedor':historial_contenedor,
							'url':url}
					return render(request, 'ver_contenedor_enviar.html',data)
				else:
					detalle_save = DetalleEnvio.objects.get(codigo=guia_hija)
					detalle_save.fue_enviada = True 
					detalle_save.save()
					modificar_envio = Envio.objects.filter(pk=detalle_save.envio.pk).update(contenedor=contenedor)
					
		except Exception as e:
			transaction.rollback()
			print('Error------------------> ',e)
			error = 'No existe la guía u ocurrio un error'
			data = {'contenedor':contenedor,
					'detalle_contenedor':detalle,
					'disponible':round(disponible,2),
					'error':error,
					'estados':estados,
					'historial_contenedor':historial_contenedor,
					'url':url}
			return render(request, 'ver_contenedor_enviar.html',data)
		else: 
			transaction.commit()
			return HttpResponseRedirect(reverse('ver_contenedor_enviar',args = (id, ))+"?ok")
	elif request.method == 'GET':
		data = {'contenedor':contenedor,
				'detalle_contenedor':detalle,
				'disponible':round(disponible,2),
				'estados':estados,
				'historial_contenedor':historial_contenedor,
				'url':url}
		return render(request, 'ver_contenedor_enviar.html',data)

#@minified_response
@login_required()
def ver_camion_enviar(request,id):
	camion_asg = Camion.objects.get(pk=id)
	detalle = DetalleEnvio.objects.filter(envio__camion=camion_asg,fue_subida_camion=True,envio__estado_envio__in=(5,6))
	estados = EstadoEnvio.objects.filter(pk__in=(5,6))
	error = ''
	#print detalle
	
	if request.method == 'POST':
		guia_hija = request.POST['guia']
		accion = request.POST['accion']
		if accion == 'consulta':
			guia_hg = request.GET.get('guia')
			caja = DetalleEnvio.objects.filter(codigo=guia_hg)
			return JsonResponse(caja, safe=False)
		elif accion == 'trasladar':
			try:
				with transaction.atomic():
					detalle_save = DetalleEnvio.objects.get(codigo=guia_hija)
					detalle_save.fue_subida_camion = True 
					detalle_save.save()
					print('este es el envio', detalle_save.envio.pk)
					Envio.objects.filter(pk=detalle_save.envio.pk).update(camion=camion_asg)
					return HttpResponseRedirect(reverse('ver_camion_enviar',args = (id, ))+"?ok")
			except Exception as e:
				error = 'No existe la guía u ocurrio un error'
				data = {'camion':camion_asg,
						'disponible':round(0,2),
						'error':error,
						'estados':estados,
						}
				print('este es el error ->', e)
				transaction.rollback()
			return render(request, 'ver_camion_enviar.html',data)
	# elif request.method == 'GET':
	# 	guia_hg = request.GET.get('guia')
	# 	caja = DetalleEnvio.objects.filter(codigo=guia_hg)
	# 	return JsonResponse(caja, safe=False)
	else:
		data = {'camion':camion_asg,
				'detalle_camion':detalle,
				'estados':estados,
				}
		return render(request, 'ver_camion_enviar.html',data)

@login_required
def trasladar_contenedor(request):
	if request.method == 'POST':
		try:
			with transaction.atomic():
				estado = request.POST.get('estado_envio')
				pk_contenedor = request.POST.get('contenedor_pk')
				
				contenedor = Contenedor.objects.get(pk=pk_contenedor)
				acum = 0
				if int(estado) == 2:
					#esta validacion solamente se hace cuando se va trasladar de la bodega de houston a puerto
					#para evitar enviar guias sin todas sus cajas
					for envios in Envio.objects.filter(contenedor=contenedor):
						for detalle_envio in DetalleEnvio.objects.filter(envio=envios):
							if detalle_envio.fue_enviada == False:
								acum += 1
								#se valida que todo el detalle de cada envio este subida en el contenedor:
				tipo_error = 0
				detalle_error = ''
				if acum>0:
					tipo_error = 1
					detalle_error = 'HAY GUIAS CON CAJAS FALTANTES'
					#return JsonResponse({'option': 'error','detalle_error':'HAY GUIAS CON CAJAS FALTANTES','tipo_error':1})
				#else:
				estado_db = EstadoEnvio.objects.get(pk=request.POST.get('estado_envio'))
				error_ = 0
				try:
					historial_contenedor = HistorialContenedor.objects.create(contenedor=contenedor,
																				estado=estado_db,
																				comentario='Traslador a '+str(estado_db.estado),
																				usuario_registro=request.user)
					contenedor.estado =estado_db
					contenedor.save()
				except Exception as e:
					#print e, "error al crear historial de contenedor"
					return JsonResponse({'option': 'error','detalle_error':'ERROR DE SISTEMA','tipo_error':0}) 

				for e in Envio.objects.filter(contenedor=contenedor):
					envio_cliente = Envio.objects.get(pk=e.pk)
					detalle_envio = DetalleEnvio.objects.filter(envio=envio_cliente)

					empresa = Envio.objects.get(pk=e.pk)
					estado_final = EstadoEnvio.objects.filter(empresa=empresa.empresa).last()
					###UPDATE EEHN
					envio = SeguimientoEnvio.objects.filter(codigo_envio=e.pk).update(estado=estado_db)
					try:
						historial = HistorialEnvio.objects.create(codigo_envio=envio_cliente,
																	estado=estado_db,
																	usuario_registro=request.user)
						
					except Exception as e:
						#print e,"ERROR AL CREAR SEGUIMIENTO DE ENVIO"
						return 0
					envio_cliente.estado_envio = EstadoEnvio.objects.get(pk=estado)
					for de in detalle_envio:
						de.estado_hija = EstadoEnvio.objects.get(pk=estado)
						de.save()
					envio_cliente.save()
					
					###UPDATE KRAKEN
					#OBTENER ENVIO EN kraken_cargo
					es_kraken = False
					kraken_envio = SistemaEmpresaenvio.objects.using('kraken_cargo').filter(codigo = empresa.guia_revendedor)
					#print kraken_envio.count(),'kraken_envio'
					if kraken_envio.count() >= 1:
						es_kraken = True
					
					if es_kraken:
						estadoeehn = EstadoEnvio.objects.get(pk=estado)
						#print estadoeehn.estado, 'estadoeehn'
						estado_kraken = ''
						pk = 0
						if estadoeehn.pk == 1:
							estado_kraken = 'BODEGA EEUU'
							pk = 8
						elif estadoeehn.pk == 2:
							estado_kraken = 'PUERTO EEUU'
							pk = 9
						elif estadoeehn.pk == 3:
							estado_kraken = 'TRANSITO MARITIMO'
							pk = 10
						elif estadoeehn.pk == 4:
							estado_kraken = 'PUERTO HONDURAS'
							pk = 11
						elif estadoeehn.pk == 5:
							estado_kraken = 'BODEGA HONDURAS'
							pk = 12
						elif estadoeehn.pk == 6:
							estado_kraken = 'EN TRANSITO PARA ENTREGA'
							pk = 13
						elif estadoeehn.pk == 7:
							estado_kraken = 'ENTREGADO'
							pk = 14
						estadokraken = SistemaEmpresaestadoenvio.objects.using('kraken_cargo').get(pk=pk)
						#print estadokraken.estado, 'estadokraken'
						#print envio_cliente.guia_revendedor,'guia revendedor'
						seguimiento_kraken_cargo = SistemaSeguimientoenvio.objects.db_manager('kraken_cargo').filter(codigo_envio=SistemaEmpresaenvio.objects.using('kraken_cargo').get(codigo =  envio_cliente.guia_revendedor)).update(estado=SistemaEmpresaestadoenvio.objects.using('kraken_cargo').get(pk=pk))
						historial_kraken_cargo = SistemaHistorialenvio.objects.db_manager('kraken_cargo').create(codigo_envio=SistemaEmpresaenvio.objects.using('kraken_cargo').get(codigo =  envio_cliente.guia_revendedor),estado=SistemaEmpresaestadoenvio.objects.using('kraken_cargo').get(pk=pk),usuario_registro=AuthUser.objects.using('kraken_cargo').get(pk=14),fechahora = timezone.now())
						envio_kraken = SistemaEmpresaenvio.objects.using('kraken_cargo').get(codigo =  envio_cliente.guia_revendedor)
						hoy = timezone.now()
		except Exception as e:
			print('Error en contenedor --> ',e)
			transaction.rollback()
		else:
			transaction.commit()
			return JsonResponse({'option': 'save','url': reverse('ver_contenedor_enviar',args = (pk_contenedor, )),'tipo_error':tipo_error,'detalle_error':detalle_error})

	elif request.method == 'GET':
		return render(request, '404.html')

@login_required
def trasladar_camion(request):
	if request.method == 'POST':
		
		estado = request.POST.get('estado_envio')
		pk_camion = request.POST.get('camion_pk')
		camion = Camion.objects.get(pk=pk_camion)
		acum = 0
		if int(estado) == 6:
			#esta validacion solamente se hace cuando se va trasladar de la bodega A TRANSITO
			#para evitar enviar guias sin todas sus cajas

			for envios in Envio.objects.filter(camion=camion):
				for detalle_envio in DetalleEnvio.objects.filter(envio=envios):
					if detalle_envio.fue_subida_camion == False:
						acum += 1
						#se valida que todo el detalle de cada envio este subida en el camion:
		tipo_error = 0
		detalle_error = ''
		if acum>0:
			tipo_error = 1
			detalle_error = 'HAY GUIAS CON CAJAS FALTANTES'
			#return JsonResponse({'option': 'error','detalle_error':'HAY GUIAS CON CAJAS FALTANTES','tipo_error':1})
		#else:
		estado_db = EstadoEnvio.objects.get(pk=request.POST.get('estado_envio'))
		error_ = 0
		for e in Envio.objects.filter(camion=camion):
			envio_cliente = Envio.objects.get(pk=e.pk)
			empresa = Envio.objects.get(pk=e.pk)
			estado_final = EstadoEnvio.objects.filter(empresa=empresa.empresa).last()
			###UPDATE EEHN
			envio = SeguimientoEnvio.objects.filter(codigo_envio=e.pk).update(estado=estado_db)
			try:
				historial = HistorialEnvio.objects.create(codigo_envio=envio_cliente,
														estado=estado_db,
														usuario_registro=request.user)

			except Exception as e:
				#print e,"ERROR AL CREAR SEGUIMIENTO DE ENVIO"
				return 0
			envio_cliente.estado_envio = EstadoEnvio.objects.get(pk=estado)
			envio_cliente.save()
			
			###UPDATE KRAKEN
			#OBTENER ENVIO EN kraken_cargo
			es_kraken = False
			kraken_envio = SistemaEmpresaenvio.objects.using('kraken_cargo').filter(codigo = empresa.guia_revendedor)
			#print kraken_envio.count(),'kraken_envio'
			if kraken_envio.count() >= 1:
				es_kraken = True
			
			if es_kraken:
				estadoeehn = EstadoEnvio.objects.get(pk=estado)
				#print estadoeehn.estado, 'estadoeehn'
				estado_kraken = ''
				pk = 0
				if estadoeehn.pk == 1:
					estado_kraken = 'BODEGA EEUU'
					pk = 8
				elif estadoeehn.pk == 2:
					estado_kraken = 'PUERTO EEUU'
					pk = 9
				elif estadoeehn.pk == 3:
					estado_kraken = 'TRANSITO MARITIMO'
					pk = 10
				elif estadoeehn.pk == 4:
					estado_kraken = 'PUERTO HONDURAS'
					pk = 11
				elif estadoeehn.pk == 5:
					estado_kraken = 'BODEGA HONDURAS'
					pk = 12
				elif estadoeehn.pk == 6:
					estado_kraken = 'EN TRANSITO PARA ENTREGA'
					pk = 13
				elif estadoeehn.pk == 7:
					estado_kraken = 'ENTREGADO'
					pk = 14
				estadokraken = SistemaEmpresaestadoenvio.objects.using('kraken_cargo').get(pk=pk)
				#print estadokraken.estado, 'estadokraken'
				#print envio_cliente.guia_revendedor,'guia revendedor'
				seguimiento_kraken_cargo = SistemaSeguimientoenvio.objects.db_manager('kraken_cargo').filter(codigo_envio=SistemaEmpresaenvio.objects.using('kraken_cargo').get(codigo =  envio_cliente.guia_revendedor)).update(estado=SistemaEmpresaestadoenvio.objects.using('kraken_cargo').get(pk=pk))
				historial_kraken_cargo = SistemaHistorialenvio.objects.db_manager('kraken_cargo').create(codigo_envio=SistemaEmpresaenvio.objects.using('kraken_cargo').get(codigo =  envio_cliente.guia_revendedor),estado=SistemaEmpresaestadoenvio.objects.using('kraken_cargo').get(pk=pk),usuario_registro=AuthUser.objects.using('kraken_cargo').get(pk=14),fechahora = timezone.now())
				envio_kraken = SistemaEmpresaenvio.objects.using('kraken_cargo').get(codigo =  envio_cliente.guia_revendedor)
				hoy = timezone.now()	
		return JsonResponse({'option': 'save','url': reverse('ver_camion_enviar',args = (pk_camion, )),'tipo_error':tipo_error,'detalle_error':detalle_error})
	elif request.method == 'GET':
		return render(request, '404.html')
@login_required
def lista_guias_faltante(request,id_contenedor):
	contenedor = Contenedor.objects.get(pk=id_contenedor)
	detalle_contenedor = DetalleEnvio.objects.filter(envio__contenedor=id_contenedor,fue_enviada=False)
	data = {'detalle_contenedor':detalle_contenedor,'contenedor':contenedor}
	return render(request, 'lista_guias_faltante.html',data)

@login_required
def lista_guias_faltante_camion(request,id_camion):
	camion = Camion.objects.get(pk=id_camion)
	#print camion
	detalle_camion = DetalleEnvio.objects.filter(envio__camion=camion,fue_subida_camion=False)
	#print detalle_camion
	data = {'detalle_camion':detalle_camion,'camion':camion}
	return render(request, 'lista_guias_faltante_camion.html',data)

@login_required
def buscar_guia(request):
	tiene_datos = 0
	resultado = {}
	data = []
	if request.is_ajax():
		seguimiento = SeguimientoEnvio.objects.filter(codigo_envio__codigo= request.GET.get('guia')).order_by('-id')[0]
		#print seguimiento
		detalle = DetalleEnvio.objects.filter(envio__codigo=request.GET.get('guia'))
		lista_detalle = []
		if detalle:
			tiene_datos = 1
			for d in detalle:
				dic = {}
				dic['cantidad'] = d.cantidad
				dic['caja'] = d.tipo_caja.tipo_caja.descripcion 
				dic['precio'] = d.precio
				dic['total'] = d.total
				if d.fue_enviada:
					dic['enviada'] = 'SI'
				else:
					dic['enviada'] = 'NO'

				if d.fue_subida_camion:
					dic['camion'] = 'SI'
				else:
					dic['camion'] = 'NO'
				lista_detalle.append(dic) 
		if seguimiento.codigo_envio.contenedor:
			contendor = str(seguimiento.codigo_envio.contenedor.codigo_original)+'|'+str(seguimiento.codigo_envio.contenedor.codigo_express)
		else:
			contendor ='NO ESTA ASIGNADA A CONTENEDOR'

		if seguimiento.codigo_envio.camion:
			camion = str(seguimiento.codigo_envio.camion.descripcion)+'|'+str(seguimiento.codigo_envio.camion.placa)
		else:
			camion ='NO ESTA ASIGNADA A CONTENEDOR'
		data = {'pk':seguimiento.pk,
				'codigo':seguimiento.codigo_envio.codigo,
				'envia':seguimiento.codigo_envio.quien_envia.nombre_completo,
				'recibe':seguimiento.codigo_envio.quien_recibe.nombre_completo,
				'estado':seguimiento.estado.estado,
				'fechahora':str(seguimiento.fechahora),
				'usuario_registro':seguimiento.usuario_registro.username,
				'contenedor':contendor,
				'tiene_datos':tiene_datos,
				'lista_detalle':lista_detalle,
				'camion':camion}
		return HttpResponse(simplejson.dumps(data), content_type='application/javascript')
	return render(request, 'buscar_guia.html')

@login_required
def seleccionar_camion(request):
	camiones = Camion.objects.filter(estado=True)
	data = {'camiones':camiones}
	return render(request, 'seleccionar_camion.html',data)

@login_required
def enviar_transito(request):
	tiene_datos = 0
	resultado = {}
	data = []
	if request.is_ajax():
		try:
			seguimiento = SeguimientoEnvio.objects.get(codigo_envio__codigo= request.GET.get('guia'),estado=5)

			detalle = DetalleEnvio.objects.filter(envio__codigo=seguimiento.codigo_envio.codigo)
			lista_detalle = []
			if detalle:
				tiene_datos = 1
				for d in detalle:
					dic = {}
					dic['cantidad'] = d.cantidad
					dic['caja'] = d.tipo_caja.tipo_caja.descripcion 
					dic['precio'] = d.precio
					dic['total'] = d.total
					if d.fue_enviada:
						dic['enviada'] = 'SI'
					else:
						dic['enviada'] = 'NO'
					lista_detalle.append(dic) 
			if seguimiento.codigo_envio.contenedor:
				contendor = str(seguimiento.codigo_envio.contenedor.codigo_original)+'|'+str(seguimiento.codigo_envio.contenedor.codigo_express)
			else:
				contendor ='NO ESTA ASIGNADA A CONTENEDOR'
			data = {'pk':seguimiento.pk,
					'codigo':seguimiento.codigo_envio.codigo,
					'envia':seguimiento.codigo_envio.quien_envia.nombre_completo,
					'recibe':seguimiento.codigo_envio.quien_recibe.nombre_completo,
					'estado':seguimiento.estado.estado,
					'fechahora':str(seguimiento.fechahora),
					'usuario_registro':seguimiento.usuario_registro.username,
					'contenedor':contendor,
					'tiene_datos':tiene_datos,
					'lista_detalle':lista_detalle,
					'url' : reverse('trasladar_a_transito', kwargs={'envio':request.GET.get('guia')})}
		except Exception as e:
			data = {'tiene_datos':0}
			#print 'errores', e
		
		return HttpResponse(simplejson.dumps(data), content_type='application/javascript')
	return render(request, 'enviar_transito.html')

@login_required
def trasladar_a_transito(request,envio):
	e = Envio.objects.get(codigo=envio)
	estado = 6
	envio_cliente = Envio.objects.get(pk=e.pk)
	empresa = Envio.objects.get(pk=e.pk)
	estado_final = EstadoEnvio.objects.filter(empresa=empresa.empresa).last()
	

	###UPDATE EEHN
	envio = SeguimientoEnvio.objects.filter(codigo_envio=Envio.objects.get(pk=e.pk)).update(estado=estado)
	historial = HistorialEnvio.objects.create(codigo_envio=Envio.objects.get(pk=e.pk),estado=EstadoEnvio.objects.get(pk=estado),usuario_registro=request.user)
	###UPDATE KRAKEN
	#OBTENER ENVIO EN kraken_cargo
	es_kraken = False
	kraken_envio = SistemaEmpresaenvio.objects.using('kraken_cargo').filter(codigo = empresa.guia_revendedor)
	#print kraken_envio.count(),'kraken_envio'
	if kraken_envio.count() >= 1:
		es_kraken = True
	#print es_kraken,'es_kraken'
	if es_kraken:
		estadoeehn = EstadoEnvio.objects.get(pk=estado)
		#print estadoeehn.estado, 'estadoeehn'
		estadokraken = SistemaEmpresaestadoenvio.objects.using('kraken_cargo').get(estado=estadoeehn.estado)
		#print estadokraken.estado, 'estadokraken'
		seguimiento_kraken_cargo = SistemaSeguimientoenvio.objects.db_manager('kraken_cargo').filter(codigo_envio=SistemaEmpresaenvio.objects.using('kraken_cargo').get(codigo =  empresa.guia_revendedor)).update(estado=SistemaEmpresaestadoenvio.objects.using('kraken_cargo').get(estado=estadoeehn.estado))
		historial_kraken_cargo = SistemaHistorialenvio.objects.db_manager('kraken_cargo').create(codigo_envio=SistemaEmpresaenvio.objects.using('kraken_cargo').get(codigo =  empresa.guia_revendedor),estado=SistemaEmpresaestadoenvio.objects.using('kraken_cargo').get(estado=estadoeehn.estado),usuario_registro=AuthUser.objects.using('kraken_cargo').get(pk=14),fechahora = timezone.now())
		

		envio_kraken = SistemaEmpresaenvio.objects.using('kraken_cargo').get(codigo =  empresa.guia_revendedor)
		hoy = timezone.now()
		

		###PUERTO HONDURAS CREDITO
		if envio_kraken.credito:
			if int(estado) == 4:
				account_sid = 'AC17aec886df904438ae784475a29acd60'
				auth_token = 'd36a6c11614c87a982ff21406c2e8bb3'
				twilio_client= Client(account_sid, auth_token)
				message = twilio_client.messages.create(
				body='ESTE ES UN MENSAJE GENERADO AUTOMATICAMENTE FAVOR NO RESPONDER. ' + '\n' + 'Saludos ' + envio_kraken.quien_envia.nombre_completo + envio_kraken.quien_envia.apellidos +'\n\n' +envio_kraken.empresa.nombre_empresa+' le informa que su envio '+envio_kraken.codigo + ' para ' + envio_kraken.quien_recibe.nombre_completo+ ', se encuentra en el Puerto de Honduras esperando el desembarque. Para evitar atrasos con su entrega favor realizar el pago pendiente de $'+ str(envio_kraken.saldo_pendiente) + '. Para más información contáctarse al número ' + envio_kraken.empresa.telefono_empresa ,
				from_='+13473345592',
				to='+1' + envio_kraken.quien_envia.celular
				)
				mensaje = SistemaEmpresamensaje.objects.db_manager('kraken_cargo').create(tipo_mensaje=SistemaTipomensajes.objects.using('kraken_cargo').get(pk=2),envio=envio_kraken,texto = message.body,fecha= hoy.date(),precio=0.25,pagado=False,fecha_pago = hoy.date())
				#print(message.sid)
		#### PUERTO HONDURAS INFO
		if int(estado) == 4:
			account_sid = 'AC17aec886df904438ae784475a29acd60'
			auth_token = 'd36a6c11614c87a982ff21406c2e8bb3'
			twilio_client= Client(account_sid, auth_token)
			message = twilio_client.messages.create(
			body='ESTE ES UN MENSAJE GENERADO AUTOMATICAMENTE FAVOR NO RESPONDER. ' + '\n' + 'Saludos ' + envio_kraken.quien_envia.nombre_completo + envio_kraken.quien_envia.apellidos +'\n\n' +envio_kraken.empresa.nombre_empresa+' le informa que su envio '+envio_kraken.codigo + ' para ' + envio_kraken.quien_recibe.nombre_completo+ ', se encuentra en el Puerto de Honduras esperando el desembarque. Muchas gracias por preferirnos. Para más información contáctarse al número ' + envio_kraken.empresa.telefono_empresa ,
			from_='+13473345592',
			to='+1' + envio_kraken.quien_envia.celular
			)
			mensaje = SistemaEmpresamensaje.objects.db_manager('kraken_cargo').create(tipo_mensaje=SistemaTipomensajes.objects.using('kraken_cargo').get(pk=2),envio=envio_kraken,texto = message.body,fecha= hoy.date(),precio=0.25,pagado=False,fecha_pago = hoy.date())
			#print(message.sid)
		#### MENSAJE DE ENTREGADO
		if int(estado) == int(estado_final.pk):
			account_sid = 'AC17aec886df904438ae784475a29acd60'
			auth_token = 'd36a6c11614c87a982ff21406c2e8bb3'
			twilio_client= Client(account_sid, auth_token)
			message = twilio_client.messages.create(
			body='ESTE ES UN MENSAJE GENERADO AUTOMATICAMENTE FAVOR NO RESPONDER. ' + '\n' +'Saludos ' + envio_kraken.quien_envia.nombre_completo + envio_kraken.quien_envia.apellidos +'\n\n' +envio_kraken.empresa.nombre_empresa+' le informa que su envio '+envio_kraken.codigo + ' para ' + envio_kraken.quien_recibe.nombre_completo+ ', ha sido entregado satisfactoriamente.' + '\n' + 'Muchas gracias por confiar en nuestros servicios esperamos pronto seguirle brindando el mejor servicio.',
			from_='+13473345592',
			to='+1' + envio_kraken.quien_envia.celular
			)
			#print(message.sid)
			mensaje = SistemaEmpresamensaje.objects.db_manager('kraken_cargo').create(tipo_mensaje=SistemaTipomensajes.objects.using('kraken_cargo').get(pk=1),envio=envio_kraken,texto = message.body,fecha = hoy.date(),precio=0.25,pagado=False,fecha_pago = hoy.date())
	## MENSAJES EXPRESS ENVIOS
	if envio_cliente.credito:
		if int(estado) == 4:
			account_sid = 'AC17aec886df904438ae784475a29acd60'
			auth_token = 'd36a6c11614c87a982ff21406c2e8bb3'
			twilio_client= Client(account_sid, auth_token)
			message = twilio_client.messages.create(
			body='ESTE ES UN MENSAJE GENERADO AUTOMATICAMENTE FAVOR NO RESPONDER. ' + '\n' + 'Saludos ' + envio_cliente.quien_envia.nombre_completo + '\n\n' +envio_cliente.empresa.nombre_empresa+' le informa que su envio '+envio_cliente.codigo + ' para ' + envio_cliente.quien_recibe.nombre_completo+ ', se encuentra en el Puerto de Honduras esperando el desembarque. Para evitar atrasos con su entrega favor realizar el pago pendiente de $'+ str(envio_cliente.saldo_pendiente) + '. Para más información contáctarse al número ' + envio_cliente.empresa.telefono_empresa ,
			from_='+13473345592',
			to='+1' + envio_cliente.quien_envia.celular
			)
			mensaje = EmpresaMensaje.objects.create(tipo_mensaje=TipoMensajes.objects.get(pk=2),envio=envio_cliente,texto = message.body)
			#print(message.sid)
	
	if int(estado) == int(estado_final.pk):
		account_sid = 'AC17aec886df904438ae784475a29acd60'
		auth_token = 'd36a6c11614c87a982ff21406c2e8bb3'
		twilio_client= Client(account_sid, auth_token)
		message = twilio_client.messages.create(
		body='ESTE ES UN MENSAJE GENERADO AUTOMATICAMENTE FAVOR NO RESPONDER. ' + '\n' +'Saludos ' + envio_cliente.quien_envia.nombre_completo + '\n\n' +envio_cliente.empresa.nombre_empresa+' le informa que su envio '+envio_cliente.codigo + ' para ' + envio_cliente.quien_recibe.nombre_completo+ ', ha sido entregado satisfactoriamente.' + '\n' + 'Muchas gracias por confiar en nuestros servicios esperamos pronto seguirle brindando el mejor servicio.',
		from_='+13473345592',
		to='+1' + envio_cliente.quien_envia.celular
		)
		#print(message.sid)
		mensaje = EmpresaMensaje.objects.create(tipo_mensaje=TipoMensajes.objects.get(pk=1),envio=envio_cliente,texto = message.body)
	return HttpResponseRedirect(reverse('enviar_transito'))


# @login_required
# def buscar_caja_transito(request):
# 	tiene_datos = 0
# 	resultado = {}
# 	data = []
# 	if request.is_ajax():
# 		try:
# 			#seguimiento = SeguimientoEnvio.objects.get(codigo_envio__codigo= request.GET.get('guia'),estado=6)
# 			seguimiento = SeguimientoEnvio.objects.filter(codigo_envio__codigo= request.GET.get('guia'),estado=6).order_by('-id')[0]
# 			detalle = DetalleEnvio.objects.filter(envio__codigo=seguimiento.codigo_envio.codigo)
# 			lista_detalle = []
# 			if detalle:
# 				tiene_datos = 1
# 				for d in detalle:
# 					dic = {}
# 					dic['cantidad'] = d.cantidad
# 					dic['caja'] = d.tipo_caja.tipo_caja.descripcion 
# 					dic['precio'] = d.precio
# 					dic['total'] = d.total
# 					if d.fue_enviada:
# 						dic['enviada'] = 'SI'
# 					else:
# 						dic['enviada'] = 'NO'
# 					lista_detalle.append(dic) 
# 			if seguimiento.codigo_envio.contenedor:
# 				contendor = str(seguimiento.codigo_envio.contenedor.codigo_original)+'|'+str(seguimiento.codigo_envio.contenedor.codigo_express)
# 			else:
# 				contendor ='NO ESTA ASIGNADA A CONTENEDOR'
# 			data = {'pk':seguimiento.pk,
# 					'codigo':seguimiento.codigo_envio.codigo,
# 					'envia':seguimiento.codigo_envio.quien_envia.nombre_completo,
# 					'recibe':seguimiento.codigo_envio.quien_recibe.nombre_completo,
# 					'estado':seguimiento.estado.estado,
# 					'fechahora':str(seguimiento.fechahora),
# 					'usuario_registro':seguimiento.usuario_registro.username,
# 					'contenedor':contendor,
# 					'tiene_datos':tiene_datos,
# 					'lista_detalle':lista_detalle,
# 					'url' : reverse('entregar_caja', kwargs={'envio':request.GET.get('guia')})}
# 		except Exception as e:
# 			data = {'tiene_datos':0}
# 			#print 'errores', e
		
# 		return HttpResponse(simplejson.dumps(data), content_type='application/javascript')
# 	return render(request, 'buscar_caja_transito.html')

# @login_required
# def entregar_caja(request,envio):
# 	e = Envio.objects.get(codigo=envio)
# 	estado = 7
# 	envio_cliente = Envio.objects.get(pk=e.pk)
# 	empresa = Envio.objects.get(pk=e.pk)
# 	estado_final = EstadoEnvio.objects.filter(empresa=empresa.empresa).last()
	

# 	###UPDATE EEHN
# 	envio = SeguimientoEnvio.objects.filter(codigo_envio=Envio.objects.get(pk=e.pk)).update(estado=estado)
# 	historial = HistorialEnvio.objects.create(codigo_envio=Envio.objects.get(pk=e.pk),estado=EstadoEnvio.objects.get(pk=estado),usuario_registro=request.user)
# 	envio_cliente.estado_envio = estado_final
# 	envio_cliente.save()
# 	###UPDATE KRAKEN
# 	#OBTENER ENVIO EN kraken_cargo
# 	es_kraken = False
# 	kraken_envio = SistemaEmpresaenvio.objects.using('kraken_cargo').filter(codigo = envio_cliente.guia_revendedor)
# 	if kraken_envio.count() >= 1:
# 		es_kraken = True
	
# 	if es_kraken:
# 		estado_kraken = ''
# 		pk = 0
# 		if estado_final.pk == 1:
# 			estado_kraken = 'BODEGA EEUU'
# 			pk = 8
# 		elif estado_final.pk == 2:
# 			estado_kraken = 'PUERTO EEUU'
# 			pk = 9
# 		elif estado_final.pk == 3:
# 			estado_kraken = 'TRANSITO MARITIMO'
# 			pk = 10
# 		elif estado_final.pk == 4:
# 			estado_kraken = 'PUERTO HONDURAS'
# 			pk = 11
# 		elif estado_final.pk == 5:
# 			estado_kraken = 'BODEGA HONDURAS'
# 			pk = 12
# 		elif estado_final.pk == 6:
# 			estado_kraken = 'EN TRANSITO PARA ENTREGA'
# 			pk = 13
# 		elif estado_final.pk == 7:
# 			estado_kraken = 'ENTREGADO'
# 			pk = 14
# 		estadokraken = SistemaEmpresaestadoenvio.objects.using('kraken_cargo').get(pk=pk)
# 		seguimiento_kraken_cargo = SistemaSeguimientoenvio.objects.db_manager('kraken_cargo').filter(codigo_envio=SistemaEmpresaenvio.objects.using('kraken_cargo').get(codigo =  envio_cliente.guia_revendedor)).update(estado=SistemaEmpresaestadoenvio.objects.using('kraken_cargo').get(pk=pk))
# 		historial_kraken_cargo = SistemaHistorialenvio.objects.db_manager('kraken_cargo').create(codigo_envio=SistemaEmpresaenvio.objects.using('kraken_cargo').get(codigo =  envio_cliente.guia_revendedor),estado=SistemaEmpresaestadoenvio.objects.using('kraken_cargo').get(pk=pk),usuario_registro=AuthUser.objects.using('kraken_cargo').get(pk=14),fechahora = timezone.now())
# 		envio_kraken = SistemaEmpresaenvio.objects.using('kraken_cargo').get(codigo =  envio_cliente.guia_revendedor)
# 		hoy = timezone.now()
# 	return HttpResponseRedirect(reverse('enviar_transito'))


@login_required
def reporte_ventas(request):
	empleados = Empleado.objects.all()
	if request.method == 'POST':
		
		if request.POST.get('fecha_inicio') !='' and request.POST.get('fecha_fin') != '':
			inicio = datetime.strptime(request.POST.get('fecha_inicio'), '%d/%m/%Y')
			fin = datetime.strptime(request.POST.get('fecha_fin'), '%d/%m/%Y')
			empleado = Empleado.objects.get(pk=int(request.POST.get('vendedor')))
			user = User.objects.get(pk=empleado.usuario.pk)
			envios = Envio.objects.filter(usuario_registro = user, aprobado = True,fecha_cierre__range=(inicio,fin))
			cajas = ReciboCaja.objects.filter(usuario_registro = user,fecha__range=(inicio,fin))
			vehiculos = ReciboVehiculos.objects.filter(usuario_registro = user,fecha_cierre__range=(inicio,fin))
			contenedores = ReciboContenedor.objects.filter(usuario_registro = user,fecha_cierre__range=(inicio,fin))
			pagos = PagosCredito.objects.filter(usuario_registro = user,fecha_cierre__range=(inicio,fin))
			saldo = 0
			saldo_cajas = 0
			saldo_vehiculos = 0
			saldo_contenedores = 0
			saldo_pagos = 0
			saldo_pagos_efectivo = 0
			dic_datos = []
			dic_cajas = []
			dic_vehiculos = []
			dic_contenedores = []
			dic_pagos = []

			for e in envios:
				datos = {}
				datos['pk'] = e.pk,
				datos['codigo'] = e.codigo
				datos['cliente'] = e.quien_envia.nombre_completo
				datos['pais'] = e.pais_destino
				datos['depto'] = e.departamento_destino
				datos['embarque']  = e.descripcion_embarque
				datos['comentario']  = e.comentario
				datos['celular'] = e.quien_envia.celular
				datos['fecha'] = e.fecha_recoleccion.strftime("%H:%M"),
				datos['saldo'] = e.saldo_pendiente
				datos['pago'] = e.pago_recibido
				saldo += e.pago_recibido
				dic_datos.append(datos)
			
			for i in cajas:
				datos = {}
				datos['codigo'] = i.recibo_no
				datos['cliente'] = i.cliente
				datos['size_caja'] = i.size_caja
				datos['valor_caja'] = i.valor_caja
				saldo_cajas += i.valor_caja
				dic_cajas.append(datos)
			
			for i in vehiculos:
				datos = {}
				datos['codigo'] = i.recibo_no
				datos['cliente'] = i.cliente
				datos['vehiculo'] = i.marca_vehiculo
				datos['modelo_vehiculo'] = i.modelo_vehiculo
				datos['valor_vehiculo'] = i.valor_vehiculo
				saldo_vehiculos += i.valor_vehiculo
				dic_vehiculos.append(datos)
			
			for i in contenedores:
				datos = {}
				datos['codigo'] = i.recibo_no
				datos['cliente'] = i.cliente
				datos['pais_destino'] = i.pais_destino
				datos['puerto_destino'] = i.puerto_destino
				datos['tamano_contenedor'] = i.tamano_contenedor
				datos['valor_contenedor'] = i.valor_contenedor
				saldo_contenedores += i.valor_contenedor
				dic_contenedores.append(datos)

			for p in pagos:
				datos = {}
				datos['factura'] = p.envio.quien_envia
				datos['pago'] = p.pago
				datos['tipo'] = p.tipo_pago
				saldo_pagos += p.pago
				if p.tipo_pago == True:
					saldo_pagos_efectivo += p.pago
				dic_pagos.append(datos)


			ctx = {'datos':dic_datos,
					'dic_cajas':dic_cajas,
					'dic_vehiculos':dic_vehiculos,
					'dic_contenedores':dic_contenedores,
					'saldo':saldo,
					'saldo_cajas':saldo_cajas,
					'saldo_vehiculos':saldo_vehiculos,
					'saldo_contenedores':saldo_contenedores, 
					'pagos':dic_pagos, 
					'saldo_pagos':saldo_pagos, 
					'saldo_pagos_efectivo':saldo_pagos_efectivo,
					'inicio':inicio,
					'fin':fin,
					'empleado':empleado
					}
			return render(request, 'resultado_reporte_venta.html', ctx)
	data = {'empleados':empleados}
	return render(request, 'reporte_ventas.html',data)

@login_required
def camiones(request):
	camiones = Camion.objects.all()
	if request.method == 'POST':
		descripcion = request.POST.get('descripcion')
		placa = request.POST.get('placa')
		try:
			camion = Camion.objects.create(descripcion=descripcion,placa=placa)
			return HttpResponseRedirect(reverse('camiones')+"?ok")
		except Exception as e:
			return HttpResponseRedirect(reverse('camiones')+"?error")
	elif request.method == 'GET':
		ctx = {'camiones':camiones}
		return render(request,'camiones.html', ctx)

@login_required
def desactivar_camion(request,id):
	camion = Camion.objects.filter(pk = id).update(estado = False)
	return HttpResponseRedirect(reverse('camiones')+"?desactivar")

@login_required
def activar_camion(request,id):
	camion = Camion.objects.filter(pk = id).update(estado  = True)
	return HttpResponseRedirect(reverse('camiones')+"?activar")

@login_required
def reporte_abonos(request):
	
	if request.method == 'POST':
		
		if request.POST.get('fecha_inicio') !='' and request.POST.get('fecha_fin') != '':
			dic_datos = []
			dic_pagos = []
			inicio = datetime.strptime(request.POST.get('fecha_inicio'), '%d/%m/%Y')
			fin = datetime.strptime(request.POST.get('fecha_fin'), '%d/%m/%Y')
			
			cliente_q = Cliente.objects.get(pk=request.POST.get('cliente'))
			envios = Envio.objects.filter(quien_envia=cliente_q,aprobado = True,fecha_envio__range=(inicio,fin))
			pagos = PagosCredito.objects.filter(envio__quien_envia=cliente_q,fecha__range=(inicio,fin))
			for e in envios:
				datos = {}
				datos['cliente'] = e.quien_envia
				datos['fecha'] = e.fecha_envio
				datos['valor'] = e.pago_recibido
				datos['guia'] = e.codigo
				datos['registro'] = e.usuario_registro
				datos['descripcion'] = 'PAGO INICIAL'
				dic_datos.append(datos)

			for p in pagos:
				datos = {}
				datos['cliente'] = p.envio.quien_envia
				datos['fecha'] = p.fecha
				datos['valor'] = p.pago
				datos['guia'] = p.envio.codigo
				datos['registro'] = p.usuario_registro
				datos['descripcion'] = 'ABONO'
				dic_pagos.append(datos)

			
			ctx = {'abonos':dic_datos,'datos':dic_pagos,'comentario':'ABONOS DEL '+str(inicio)+'AL '+str(fin)}
			return render(request,'resultados_abonos.html', ctx)
		else:
			ctx = {'clientes':Cliente.objects.all(),'error_fecha':'DEBE SELECCIONAR FECHAS'}
			return render(request,'reporte_abonos.html', ctx)
	else:
		
		ctx = {'clientes':Cliente.objects.all()}
		return render(request,'reporte_abonos.html', ctx)

#--------------------------------------------------------------------------------------------------------

@login_required
def distribuir_cajas(request):
	if request.method == 'GET' :
		camiones = Camion.objects.all()
		return render(request, 'distribuir_cajas.html',{'camiones':camiones, })

@login_required
def cargar_caja_camion(request):
	if request.method == 'POST':
		try:
			with transaction.atomic():
				
				codigo = request.POST['codigo']
				id_camion = request.POST['codigo_camion']
				print('codigo ', codigo)
				caja_hija = DetalleEnvio.objects.get(codigo=codigo)
				envio_cliente = Envio.objects.get(pk=caja_hija.envio.pk)
				caja = Envio.objects.get(pk=caja_hija.envio.pk)
				if caja_hija.estado_hija_id == 6 or caja_hija.fue_subida_camion == True:
					mensaje = 'La caja ' + caja_hija.codigo + ' ya ha sido cargada un camion'
					return JsonResponse({'error': mensaje})
				caja.estado_envio_id = 6
				caja.camion_id = id_camion
				caja.save()
				caja_hija.estado_hija_id = 6
				caja_hija.fue_subida_camion = True
				caja_hija.save()
				
				seguimiento = SeguimientoEnvio.objects.filter(codigo_envio=envio_cliente).update(estado=EstadoEnvio.objects.get(pk=6))
				historico = HistorialEnvio.objects.get_or_create(codigo_envio=envio_cliente,
															estado_id=6,
															usuario_registro=request.user)
		except Exception as e:
			print ('Error -> ', e)
			transaction.rollback()
		else:
			transaction.commit()
			
			return JsonResponse({'mensaje': 'Caja agregada a la bodega', 'caja': {'codigo': caja.codigo,
									'envia': caja.quien_envia.nombre_completo,
									'recibe': caja.quien_recibe.nombre_completo,
									'tamano': caja_hija.tipo_caja.tipo_caja.descripcion,
									'direccion': caja.direccion_registrar,
									'hija':caja_hija.codigo}})

@login_required
def ver_cajas_camion(request,id):
	camion = Camion.objects.get(pk=id)
	envios = Envio.objects.filter(camion_id=camion.pk, estado_envio_id=6)
	envios_detalle = []
	hoy = timezone.now()
	fecha = hoy.strftime('%B %d, %Y')

	for envio in envios:
		detalle={}
		detalle['guia'] = envio.codigo
		detalle['cliente'] = envio.quien_recibe.nombre_completo
		telefonos = envio.celular_registrar.split('/')
		detalle['telefono'] = telefonos
		detalle['departamento'] = envio.departamento_destino.nombre 
		detalle['direccion'] = envio.direccion_registrar
		detalle['cantidad'] = DetalleEnvio.objects.filter(envio_id=envio.pk).count()
		detalle['comentario'] = envio.comentario
		debe = False
		if (envio.saldo_pendiente > 0):
			debe = True
		detalle['saldo'] = debe

		cajas = DetalleEnvio.objects.filter(envio_id=envio.pk)
		cajas_agrupadas = {}  # diccionario para almacenar el resultado
		for caja in cajas:
			if caja.estado_hija_id == 6:
				descripcion = caja.tipo_caja.tipo_caja.descripcion
				if descripcion not in cajas_agrupadas:
					cajas_agrupadas[descripcion] = 1
				else:
					cajas_agrupadas[descripcion] += 1

		# construimos el string resultante a partir del diccionario
		cajas_string = ''
		for descripcion, cantidad in cajas_agrupadas.items():
			cajas_string += f'{descripcion}:({cantidad}), '
		cajas_string = cajas_string[:-2]  # eliminamos la última coma y espacio

		detalle['cajas'] = cajas_string
		
		envios_detalle.append(detalle)
		
	if request.method == 'POST':
		envios_detalle_ordenado = sorted(envios_detalle, key=lambda x: x['guia'])
		motorista = request.POST['motorista']
		ayuda = request.POST['ayuda']
		return generar_pdf('ver_cajas_camion_pdf.html',{
				'camion':camion,
				'envios':envios_detalle_ordenado, 
				'hoy':fecha, 
				'motorista':motorista, 
				'ayuda':ayuda,
			})
			
	return render(request, 'ver_cajas_camion.html',{'camion':camion, 'hoy':fecha})

@login_required
def entregar_caja(request,envio):
	try:
		with transaction.atomic():
			e = Envio.objects.get(codigo=envio)
			estado = 7
			envio_cliente = Envio.objects.get(pk=e.pk)
			empresa = Envio.objects.get(pk=e.pk)
			estado_final = EstadoEnvio.objects.filter(empresa=empresa.empresa).last()
			#detalle de cajas
			detalle_envio = DetalleEnvio.objects.filter(envio = envio_cliente)
			for de in detalle_envio:
						de.estado_hija = estado_final
						de.save()
			###UPDATE EEHN
			# Actualiza el seguimiento
			SeguimientoEnvio.objects.filter(codigo_envio=Envio.objects.get(pk=e.pk)).update(estado=estado)
			#crea el historico
			HistorialEnvio.objects.create(codigo_envio=Envio.objects.get(pk=e.pk),estado=EstadoEnvio.objects.get(pk=estado),usuario_registro=request.user)
			envio_cliente.estado_envio = estado_final

			envio_cliente.save()
			hoy = timezone.now()
			return HttpResponseRedirect(reverse('enviar_transito'))
	except Exception as e:
		print('Error ---> ',e)
		transaction.rollback()
	return HttpResponseRedirect(reverse('enviar_transito'))

@login_required
def buscar_caja_transito(request):
	tiene_datos = 0
	resultado = {}
	data = []
	if request.is_ajax():
		try:
			with transaction.atomic():
				#seguimiento = SeguimientoEnvio.objects.get(codigo_envio__codigo= request.GET.get('guia'),estado=6)
				guia = request.GET.get('guia')
				guia_envio=''
				try:
					if Envio.objects.get(codigo=guia):
						guia_envio = guia
				except Exception as e:
					# si el envio no existe salta a la siguiente comparacion
					#buscamos la guia hija
					guia_detalle = DetalleEnvio.objects.filter(codigo=guia)
					id_envio_detalle = ''
					for d in guia_detalle:
						id_envio_detalle = d.envio.codigo
					guia_envio = id_envio_detalle
				seguimiento = SeguimientoEnvio.objects.filter(codigo_envio= Envio.objects.get(codigo=guia_envio),estado=6).order_by('-id')[0]
				detalle = DetalleEnvio.objects.filter(envio__codigo=seguimiento.codigo_envio.codigo)
				lista_detalle = []
				if detalle:
					tiene_datos = 1
					for d in detalle:
						dic = {}
						dic['cantidad'] = d.cantidad
						dic['caja'] = d.tipo_caja.tipo_caja.descripcion 
						dic['precio'] = d.precio
						dic['total'] = d.total
						if d.fue_enviada:
							dic['enviada'] = 'SI'
						else:
							dic['enviada'] = 'NO'
						lista_detalle.append(dic)
				if seguimiento.codigo_envio.contenedor:
					contendor = str(seguimiento.codigo_envio.contenedor.codigo_original)+'|'+str(seguimiento.codigo_envio.contenedor.codigo_express)
				else:
					contendor ='NO ESTA ASIGNADA A CONTENEDOR'
				data = {'pk':seguimiento.pk,
						'codigo':seguimiento.codigo_envio.codigo,
						'envia':seguimiento.codigo_envio.quien_envia.nombre_completo,
						'recibe':seguimiento.codigo_envio.quien_recibe.nombre_completo,
						'estado':seguimiento.estado.estado,
						'fechahora':str(seguimiento.fechahora),
						'usuario_registro':seguimiento.usuario_registro.username,
						'contenedor':contendor,
						'tiene_datos':tiene_datos,
						'lista_detalle':lista_detalle,
						'url' : reverse('entregar_caja', kwargs={'envio':guia_envio})}
		except Exception as e:
			print('Error ---> ',e)
			transaction.rollback()
			data = {'tiene_datos':0}
			#print 'errores', e
		
		return HttpResponse(simplejson.dumps(data), content_type='application/javascript')
	return render(request, 'buscar_caja_transito.html')
# ------- de puerto a bodega
@login_required
def recibir_bodega(request):
	if request.method == 'GET' :

		cajas = DetalleEnvio.objects.filter(fue_subida_camion=False,envio__estado_envio=4)
		return render(request, 'recibir_bodega_hn.html',{'cajas':cajas})

@login_required
def busqueda_caja_bodega(request):
	if request.method == 'POST':
		try:
			with transaction.atomic():
				
				codigo = request.POST['codigo']
				caja_hija = DetalleEnvio.objects.get(codigo=codigo)
				envio_cliente = Envio.objects.get(pk=caja_hija.envio.pk)
				caja = Envio.objects.get(pk=caja_hija.envio.pk)
				if caja_hija.estado_hija_id == 5:
					mensaje = 'La caja ' + caja_hija.codigo + ' ya ha sido recibida en bodega Hn'
					return JsonResponse({'error': mensaje})
				caja.estado_envio_id = 5
				caja.save()
				caja_hija.estado_hija_id = 5
				caja_hija.save()

				seguimiento = SeguimientoEnvio.objects.filter(codigo_envio=envio_cliente).update(estado=EstadoEnvio.objects.get(pk=5))
				historico = HistorialEnvio.objects.get_or_create(codigo_envio=envio_cliente,
															estado_id=5,
															usuario_registro=request.user)
		except Exception as e:
			print ('Error -> ', e)
			transaction.rollback()
		else:
			transaction.commit()
			
			return JsonResponse({'mensaje': 'Caja agregada a la bodega', 'caja': {'codigo': caja.codigo,
									'envia': caja.quien_envia.nombre_completo,
									'recibe': caja.quien_recibe.nombre_completo,
									'tamano': caja_hija.tipo_caja.tipo_caja.descripcion,
									'direccion': caja.direccion_registrar,
									'hija':caja_hija.codigo}})

@login_required
def recibido_bodega_hn(request):
	user_reg = request.user
	envios_detalle = []
	hoy = timezone.now()
	fecha = hoy.strftime('%B %d, %Y')
	envios = []
	if request.method == 'GET':
		try:
			envs = request.GET.get('codigos')
			envios = envs.split(",")
			for envio in envios:
				detalle={}
				envio_d = Envio.objects.get(codigo=envio)
				detalle['guia'] = envio_d.codigo
				detalle['cliente'] = envio_d.quien_recibe.nombre_completo
				detalle['telefono'] = envio_d.celular_registrar
				detalle['departamento'] = envio_d.departamento_destino.nombre 
				detalle['direccion'] = envio_d.direccion_registrar
				detalle['cantidad'] = DetalleEnvio.objects.filter(envio_id=envio_d.pk).count()
				detalle['comentario'] = envio_d.comentario
				debe = False
				if (envio_d.saldo_pendiente > 0):
					debe = True
				detalle['saldo'] = debe

				cajas = DetalleEnvio.objects.filter(envio_id=envio_d.pk)
				cajas_agrupadas = {}  # diccionario para almacenar el resultado
				for caja in cajas:
					if caja.estado_hija_id == 5:
						descripcion = caja.tipo_caja.tipo_caja.descripcion
						if descripcion not in cajas_agrupadas:
							cajas_agrupadas[descripcion] = 1
						else:
							cajas_agrupadas[descripcion] += 1

				# construimos el string resultante a partir del diccionario
				cajas_string = ''
				for descripcion, cantidad in cajas_agrupadas.items():
					cajas_string += f'{descripcion}:({cantidad}), '
				cajas_string = cajas_string[:-2]  # eliminamos la última coma y espacio

				detalle['cajas'] = cajas_string
				
				envios_detalle.append(detalle)
				envios_detalle_ordenado = sorted(envios_detalle, key=lambda x: x['guia'])
		
			return generar_pdf('ver_recibido_bodega_hn_pdf.html',{
									'envios':envios_detalle_ordenado, 
									'hoy':fecha, 
									'registra': user_reg,
									})
		except Exception as e:
			print('ERROR-----------------> ',e)
	else:
		return HttpResponseRedirect('Metodo no permitido')

@login_required
def cierre_anual(request):
	return render(request, 'cierre_anual.html')

@login_required
def cierre_anual_print(request):
	empleado = Empleado.objects.get(usuario=request.user)
	empresa_e = EmpresaEmpleado.objects.get(empleado = empleado)
	empresa = Empresa.objects.get(id=empresa_e.empresa_id)
	if request.method == 'POST':
		year = request.POST['year']
		try:
			envios = Envio.objects.filter(fecha_cierre__year = year, aprobado=True, cierre = True)
			cajas = ReciboCaja.objects.filter(fecha_cierre__year = year)
			vehiculos = ReciboVehiculos.objects.filter(fecha_cierre__year = year)
			contenedores = ReciboContenedor.objects.filter(fecha_cierre__year = year)
			
			resultados_envios = {}
			resultados_cajas = {}
			resultados_vehiculos = {}
			resultados_contenedores = {}

			total_anual_envios=0.0
			total_anual_cajas=0.0
			total_anual_vehiculos=0.0
			total_anual_contenedores=0.0

			for month in range(1,13):
				envios_mes = envios.filter(fecha_cierre__month = month)
				envios_sin_NaN = envios_mes.exclude(total='NaN')
				total_mes_envios = envios_sin_NaN.aggregate(Sum('total'))['total__sum'] or 0
				resultados_envios[month] = total_mes_envios
				total_anual_envios += total_mes_envios

				cajas_mes = cajas.filter(fecha_cierre__month = month)
				total_mes_cajas = cajas_mes.aggregate(Sum('valor_caja'))['valor_caja__sum'] or 0
				resultados_cajas[month] = total_mes_cajas
				total_anual_cajas += total_mes_cajas

				vehiculos_mes = vehiculos.filter(fecha_cierre__month = month)
				total_mes_vehiculos = vehiculos_mes.aggregate(Sum('valor_vehiculo'))['valor_vehiculo__sum'] or 0
				resultados_vehiculos[month] = total_mes_vehiculos
				total_anual_vehiculos += total_mes_vehiculos

				contenedores_mes = contenedores.filter(fecha_cierre__month = month)
				total_mes_contenedores = contenedores_mes.aggregate(Sum('valor_contenedor'))['valor_contenedor__sum'] or 0
				resultados_contenedores[month] = total_mes_contenedores
				total_anual_contenedores += total_mes_contenedores

			total_general = total_anual_envios + total_anual_cajas + total_anual_vehiculos + total_anual_contenedores
			return generar_pdf('cierre_anual_print.html',
						{'pagesize':'A4',
						'orientation':'landscape',
						'empresa':empresa, 'resultados_envios': resultados_envios, 'resultados_cajas': resultados_cajas, 'resultados_vehiculos': resultados_vehiculos,
                        'resultados_contenedores': resultados_contenedores, 'total_anual_envios': total_anual_envios,'total_anual_cajas': total_anual_cajas,
                        'total_anual_vehiculos': total_anual_vehiculos,'total_anual_contenedores': total_anual_contenedores, 'total_general': total_general,
						'year':year,'iterador':range(1,13),
						}
					)
		except Exception as e:
			print('Error--> ',str(e))

@login_required
def registrar_envio_rv(request):
	empleado = Empleado.objects.get(usuario=request.user)
	empresa = EmpresaEmpleado.objects.get(empleado = empleado)
	r = Revendedor.objects.filter(usuario=request.user)
	pais = Pais.objects.all()
	ver_clientes = Cliente.objects.all()
	departamento = Departamento.objects.all()
	quien_recibe = ClienteRecibe.objects.all()
	es_revendedor = False
	correlativo = Envio.objects.filter(empresa=empresa.empresa)
	#print correlativo.count(), 'correlativo'
	tipo_contenido = TipoContenido.objects.all()
	tipo_envio = TipoEnvio.objects.all()
	if r.count() >= 1:
		revendedor_creo = Revendedor.objects.get(usuario = request.user)
		es_revendedor = True
		revendedor_creo = Revendedor.objects.get(usuario = request.user)
		quien_envia = Cliente.objects.filter(empresa=empresa.empresa,revendedor = True,revenedor_creo=revendedor_creo.pk)
	else:
		quien_envia = Cliente.objects.filter(empresa=empresa.empresa)

	#imagen = request.FILES.get('imagen')
	if request.method == 'POST':
		valor_adicional = 0
		valor_emplasticado = 0
		valor_seguro = 0
		codigo_inicial = 2006
		correlativo = Envio.objects.filter(empresa=empresa.empresa)
		codigo_envio = int(codigo_inicial) + int(correlativo.count())
		codigo_empresa = empresa.empresa.codigo_empresa
		codigo_final = codigo_empresa + '00' + str(codigo_envio)
		ret_data,errores,query_envio,recibo,ingreso_correcto = {},{},{},{},{}
		recibo = validad_envio(request,quien_envia,quien_recibe,pais,departamento)
		ret_data.update(recibo['ret_data'])
		errores.update(recibo['errores'])
		query_envio.update({'empresa':empresa.empresa})
		query_envio.update({'codigo':codigo_final})
		query_envio.update(recibo['query_envio'])
		#//////
		ret_data.update({'precioadicional':valor_adicional})
		query_envio.update({'valor_adicional':valor_adicional})
		ret_data.update({'valor_emplasticado':valor_emplasticado})
		query_envio.update({'valor_emplasticado':valor_emplasticado})
		ret_data.update({'valor_seguro':valor_seguro})
		query_envio.update({'valor_seguro':valor_seguro})
		guia_revend = empleado.nombres_empleado +' '+ empleado.apellidos_empleado
		if request.POST.get('valor_seguro') != '':
			valor_seguro = request.POST.get('valor_seguro')
			ret_data.update({'valor_seguro':valor_seguro})
			query_envio.update({'valor_seguro':valor_seguro})

		if request.POST.get('valor_envio') == '':
			errores.update({'valor_envio':'DEBE REALIZAR EL CALCULO DEL ENVIO'})
		else:
			if float(request.POST.get('valor_envio')) <=0:
				errores.update({'valor_envio':'DEBE REALIZAR EL CALCULO DEL ENVIO O AGREGAR CAJAS'})
		if not errores:
			valor = float(request.POST.get('valor_envio'))
			total = valor + float(valor_adicional) + float(valor_emplasticado) + float(valor_seguro)
			pago_recibido = total
			
			if pago_recibido < total:
				credito = True
				saldo_pendiente = (total - pago_recibido)
			else:
				credito = False
				saldo_pendiente = 0.00
				ret_data.update({'pago_recibido':pago_recibido})
				query_envio.update({'pago_recibido':pago_recibido})
			if es_revendedor:
				
				if request.POST.get('guia_revendedor') != '':
					ret_data.update({'guia_revendedor':guia_revend})
					query_envio.update({'guia_revendedor':guia_revend.upper()})
					query_envio.update({'revendedor':True})
					query_envio.update({'aprobado':False})
				else:
					errores['guia_revendedor'] = u'Falta ingresar la guia de revendedor'

			query_envio.update({'guia_revendedor':guia_revend.upper()})

			if request.POST.get('tipo_contenido', '') and request.POST.get('tipo_contenido', '') != '0':
				ret_data['id_tipo_contenido'] = int(request.POST.get('tipo_contenido',''))
				tipo_contenido_query = tipo_contenido.get(pk=request.POST.get('tipo_contenido',''))
				query_envio['tipo_contenido'] = tipo_contenido_query
			else:
				tipo_contenido_query = TipoContenido.objects.get(pk=1)
				query_envio['tipo_contenido'] = tipo_contenido_query
			
			if request.POST.get('tipo_envio', '') and request.POST.get('tipo_envio', '') != '0':
				ret_data['id_tipo_envio'] = int(request.POST.get('tipo_envio',''))
				tipo_envio_query = tipo_envio.get(pk=request.POST.get('tipo_envio',''))
				query_envio['tipo_envio'] = tipo_envio_query
			else:
				tipo_envio_query = TipoEnvio.objects.get(pk=1)
				query_envio['tipo_envio'] = tipo_envio_query
				
			comentario = "Valor declarado: " + str(request.POST.get('valor_declarado')) + " | "+ request.POST.get('comentario').upper()

			query_envio.update({'total':total,'usuario_registro':request.user,
								'usuario_aprobo':request.user,
								'credito':credito,
								'valor_emplasticado':valor_emplasticado,
								'saldo_pendiente':saldo_pendiente,
								'comentario': comentario,
								'estado_envio':EstadoEnvio.objects.get(pk=1)})
			try:
				with transaction.atomic():
					envio = Envio(**query_envio)
					envio.save()
					acum = 0
					for detalle in request.POST.getlist('cajas'):
						acum+=1
						dato = detalle.split('|')
						pk_caja = dato[0]
						cantidad = dato[1]
						pk = dato[2]
						caja_pais = CajaPais.objects.get(pk=int(pk))
						precio = caja_pais.precio
						acum_prueba = 0
						for x in range(0, int(cantidad)):
							acum_prueba += 1
							codigo_detalle = codigo_final + str(caja_pais.pk) +str(acum_prueba)
							total = float(precio) * int(1)
							try:
								detalle = DetalleEnvio.objects.create(envio=envio,
																	tipo_caja=caja_pais,
																	precio=precio,
																	cantidad=1,
																	codigo_orden=acum_prueba,
																	codigo=codigo_detalle,
																	total=total)
							except Exception as e:
								print("Error registro caja detalle revendedor: --> ", e)
								return 0
						seguimiento = SeguimientoEnvio.objects.create(codigo_envio=Envio.objects.get(pk=envio.pk),estado=EstadoEnvio.objects.get(pk=1),empresa=empresa.empresa,usuario_registro=request.user)
						historial = HistorialEnvio.objects.create(codigo_envio=Envio.objects.get(pk=envio.pk),estado=EstadoEnvio.objects.get(pk=1),usuario_registro=request.user)

			except Exception as e:
				print("Error registro revendedor: --> ", e)
				errores['extra'] = e
				transaction.rollback()
				ctx = {'es_revendedor':es_revendedor,'pais':pais,'ret_data':ret_data,'errores':errores}
				return render(request,'registrar_envio_rv.html',ctx)
			else:
				transaction.commit()
				
				ingreso_correcto['mensaje'] = u"Datos guardados correctamente."
				ctx = {'es_revendedor':es_revendedor,'ingreso_correcto':ingreso_correcto, 'envio':envio,'tipo_contenido':tipo_contenido,'tipo_envio':tipo_envio}
				#return HttpResponseRedirect(reverse('registrar_envio')+"?ok" +"&envio="+ str(envio.pk))
				return render(request,'registrar_envio_rv.html',ctx)
		else:
			#print errores,'erroresdentro de errores'
			ctx = {'es_revendedor':es_revendedor,'pais':pais,'ret_data':ret_data,'errores':errores,'quien_envia':quien_envia,'tipo_contenido':tipo_contenido,'tipo_envio':tipo_envio}
			return render(request,'registrar_envio_rv.html',ctx)
	elif request.method == 'GET':
		ctx = {'es_revendedor':es_revendedor,
				'pais':pais,
				'quien_envia':quien_envia,
				'quien_recibe':quien_recibe,
				'tipo_contenido':tipo_contenido,
				'tipo_envio':tipo_envio,
				'empleado':empleado}
		return render(request,'registrar_envio_rv.html',ctx)


@login_required
def cajas_contenedor_pdf(request,id):
	empleado = Empleado.objects.get(usuario=request.user)
	empresa_e = EmpresaEmpleado.objects.get(empleado = empleado)
	empresa = Empresa.objects.get(id=empresa_e.empresa_id)
	
	envios = Envio.objects.filter(contenedor_id = id)
	
	resultados_envios = []
	contenedor = Contenedor.objects.get(pk = id)
	for e in envios:
		detalle = {}
		detalle['guia']= e.codigo
		detalle['persona_envia'] = e.quien_envia.nombre_completo
		detalle['persona_recibe'] = e.quien_recibe.nombre_completo
		detalle['persona_envia_telefono'] = e.quien_envia.celular
		detalle['persona_recibe_telefono'] = e.quien_recibe.celular
		detalle['pais_destino'] = e.pais_destino.nombre

		cajas = DetalleEnvio.objects.filter(envio_id=e.pk)
		cajas_agrupadas = {}  # diccionario para almacenar el resultado
		for caja in cajas:
			descripcion = caja.tipo_caja.tipo_caja.descripcion
			if descripcion not in cajas_agrupadas:
				cajas_agrupadas[descripcion] = 1
			else:
				cajas_agrupadas[descripcion] += 1

		# construimos el string resultante a partir del diccionario
		cajas_string = ''
		for descripcion, cantidad in cajas_agrupadas.items():
			cajas_string += f'{descripcion}:({cantidad}), '
		cajas_string = cajas_string[:-2]  # eliminamos la última coma y espacio

		detalle['cajas'] = cajas_string
		resultados_envios.append(detalle)
		envios_detalle_ordenado = sorted(resultados_envios, key=lambda x: x['guia'])
	if request.method == 'GET':
		return generar_pdf('cajas_contenedor_pdf.html',{'envios':envios_detalle_ordenado,'empresa':empresa, 'contenedor':contenedor})
	
	return render(request, 'cajas_contenedor_pdf.html',{'envios':envios_detalle_ordenado,'empresa':empresa, 'contenedor':contenedor})

@login_required
def cajas_contenedor_xls(request,id):
	# Crear un nuevo archivo de Excel
	workbook = Workbook()
	
	envios = Envio.objects.filter(contenedor_id = id).order_by('codigo')
	
	contenedor = Contenedor.objects.get(pk = id)
	# Seleccionar la hoja activa
	sheet = workbook.active
	#cabeceras
	sheet['A1'] = 'Guia'
	sheet['B1'] = 'Envia'
	sheet['C1'] = 'Recibe'
	sheet['D1'] = 'Pais Destino'
	sheet['E1'] = 'Departamento'
	sheet['F1'] = 'Direccion'
	sheet['G1'] = 'Cajas'

	for index,e in enumerate(envios):
		row = index + 2 #para empezar a escribir en la segunda fila
		sheet[f'A{row}'] = e.codigo
		sheet[f'B{row}'] = e.quien_envia.nombre_completo + "|" + e.quien_envia.celular
		sheet[f'C{row}'] = e.quien_recibe.nombre_completo + "|" + e.quien_recibe.celular
		sheet[f'D{row}'] = e.pais_destino.nombre
		sheet[f'E{row}'] = e.departamento_destino.nombre
		sheet[f'F{row}'] = e.quien_recibe.direccion

		cajas = DetalleEnvio.objects.filter(envio_id=e.pk)
		cajas_agrupadas = {}  # diccionario para almacenar el resultado
		for caja in cajas:
			descripcion = caja.tipo_caja.tipo_caja.descripcion
			if descripcion not in cajas_agrupadas:
				cajas_agrupadas[descripcion] = 1
			else:
				cajas_agrupadas[descripcion] += 1

		# construimos el string resultante a partir del diccionario
		cajas_string = ''
		for descripcion, cantidad in cajas_agrupadas.items():
			cajas_string += f'{descripcion}:({cantidad}), '
		cajas_string = cajas_string[:-2]  # eliminamos la última coma y espacio

		sheet[f'G{row}'] = cajas_string

	for column_cells in sheet.columns:
		length = max(len(str(cell.value)) for cell in column_cells)
		sheet.column_dimensions[column_cells[0].column_letter].width = length
	# Generar un nombre de archivo único
	hoy = timezone.now()
	fecha = hoy.strftime('%Y_%B%d_%H%M%S')
	filename = f'envios_{contenedor.codigo_express}_{fecha}.xlsx'
	# Guardar el archivo de Excel
	workbook.save(filename)

	# Crear una respuesta HTTP con el archivo adjunto
	with open(filename, 'rb') as excel_file:
		response = HttpResponse(excel_file.read(),
                                content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
		response['Content-Disposition'] = f'attachment; filename="{filename}"'
	os.remove(filename)
	return response

@login_required
def estado_cajas_contenedor_pdf(request,id):
	empleado = Empleado.objects.get(usuario=request.user)
	empresa_e = EmpresaEmpleado.objects.get(empleado = empleado)
	empresa = Empresa.objects.get(id=empresa_e.empresa_id)
	
	envios = Envio.objects.filter(contenedor_id = id)
	
	resultados_envios = []
	contenedor = Contenedor.objects.get(pk = id)

	
	for e in envios:
		detalle = {}
		detalle_envio = DetalleEnvio.objects.filter(envio_id = e.pk)
		for de in detalle_envio:
			detalle['guia'] = e.codigo
			detalle['persona_envia'] = e.quien_envia.nombre_completo
			detalle['persona_recibe'] = e.quien_recibe.nombre_completo
			detalle['persona_envia_telefono'] = e.quien_envia.celular
			detalle['persona_recibe_telefono'] = e.quien_recibe.celular
			detalle['pais_destino'] = e.pais_destino.nombre
			detalle['estado_padre'] = e.estado_envio.estado
			detalle['hija'] = de.codigo
			detalle['tamano'] = de.tipo_caja.tipo_caja.descripcion
			detalle['cantidad'] = de.cantidad
			detalle['estado_hija'] = de.estado_hija.estado
			resultados_envios.append(detalle)
		
		
		envios_detalle_ordenado = sorted(resultados_envios, key=lambda x: x['guia'])
	if request.method == 'GET':
		return generar_pdf('estado_cajas_contenedor_pdf.html',{'envios':envios_detalle_ordenado,'empresa':empresa, 'contenedor':contenedor})
	
	return render(request, 'estado_cajas_contenedor_pdf.html',{'envios':envios_detalle_ordenado,'empresa':empresa, 'contenedor':contenedor})
# Reversiones de estado en honduras
@login_required
def reversion_estado_honduras(request):
	tiene_datos = 0
	estados = EstadoEnvio.objects.filter(pk__in=(4,5,6))
	data = []
	if request.is_ajax():
		try:
			with transaction.atomic():
				guia = request.GET.get('guia')
				
				if Envio.objects.get(codigo=guia):
					guia_envio = guia
				
				envio_detalle = Envio.objects.get(codigo=guia_envio) #guardo el envio
				seguimiento = SeguimientoEnvio.objects.filter(codigo_envio=envio_detalle.pk).order_by('-id')[0] #busca el seguimiento para actualizarlo
				detalle = DetalleEnvio.objects.filter(envio__codigo=seguimiento.codigo_envio.codigo) #para buscas las cajas del envio
				lista_detalle = []
				if detalle:
					tiene_datos = 1
					for d in detalle:
						dic = {}
						dic['cantidad'] = d.cantidad
						dic['caja'] = d.tipo_caja.tipo_caja.descripcion 
						dic['precio'] = d.precio
						dic['total'] = d.total
						lista_detalle.append(dic)
				
				data = {'pk':seguimiento.pk,
						'codigo':seguimiento.codigo_envio.codigo,
						'envia':seguimiento.codigo_envio.quien_envia.nombre_completo,
						'recibe':seguimiento.codigo_envio.quien_recibe.nombre_completo,
						'estado':seguimiento.estado.estado,
						'fechahora':str(seguimiento.fechahora),
						'tiene_datos':tiene_datos,
						'lista_detalle':lista_detalle,
						'url' : reverse('entregar_caja', kwargs={'envio':guia_envio})}
		except Exception as e:
			print('Error ---> ',e)
			transaction.rollback()
			data = {'tiene_datos':0}
		
		return HttpResponse(simplejson.dumps(data), content_type='application/javascript')
	return render(request, 'revercion_estados_honduras.html',{'estados':estados})

@login_required()
def trasladar_guia_hn(request):
	if request.method== 'POST':
		empleado = str(request.user)
		id_envio = str(request.POST['codigo_guia'])
		estado_nuevo = int(request.POST['estado_envio_guia'])
		hoy = timezone.now()
		fecha = hoy.strftime('%B %d, %Y')
		try:
			with transaction.atomic():
				envio = Envio.objects.get(codigo=id_envio)
				empresa = Empresa.objects.get(pk=envio.empresa.pk)
				
				estado = EstadoEnvio.objects.get(pk=estado_nuevo)
				#actualziar el envio
				if estado_nuevo > envio.estado_envio_id:
					return HttpResponse(status=400)
				else:
					if envio.estado_envio_id == 6:
						envio.camion = None
					envio.estado_envio = estado
					envio.save()
					coment = 'Se reverso a '+estado.estado+' a la fecha '+fecha+' por '+empleado
					seguimiento_update = SeguimientoEnvio.objects.filter(codigo_envio=envio.pk).update(estado=estado,comentario=coment)
					historial = HistorialEnvio.objects.filter(codigo_envio=envio, estado_id__gt=estado_nuevo)
					historial.delete()
				#####
				#actualizacion de detalle del envio
				if estado_nuevo < 6:
					detalle = DetalleEnvio.objects.filter(envio=envio).update(estado_hija=estado, fue_subida_camion = False)
				else:
					detalle = DetalleEnvio.objects.filter(envio=envio).update( estado_hija=estado)
				
		except Exception as e:
			print('Reversion estado Hn error ---> ',e)
			transaction.rollback()
			return 0
		else:
			transaction.commit()
			return HttpResponseRedirect(reverse('reversion_estado_honduras'))	
		
	elif request.method == 'GET':
		return HttpResponseRedirect(reverse('reversion_estado_honduras'))
	
@login_required
def lista_guias_faltante_bodega_hn(request):
	try:
		with transaction.atomic():
			envios = Envio.objects.filter(estado_envio_id = 5)

			lista_faltante = []
			#por cada envio buscar las cajas
			for env in envios:
				faltantes = DetalleEnvio.objects.filter(envio=env,estado_hija_id = 4)
				for faltante in faltantes:
					dic={}
					dic['guia_padre'] = env.codigo
					dic['guia_hija'] = faltante.codigo
					dic['envia'] = env.quien_envia.nombre_completo
					dic['recibe'] = env.quien_recibe.nombre_completo
					dic['tamano'] = faltante.tipo_caja.tipo_caja.descripcion
					dic['direccion'] = env.direccion_registrar
					lista_faltante.append(dic)
			data = {'lista_faltante':lista_faltante}
	except Exception as e:
		print("Error al mostrar faltantes----> ", e)
		transaction.rollback()
	else:
		transaction.commit()
		return render(request, 'lista_faltante_bodega_hn.html',data)
	
@login_required
def editar_dato_cliente(request):
	if request.method == 'POST':
		try:
			with transaction.atomic():
				if request.POST['id_cliente'] == '':
					return JsonResponse({'option': 'error','detalle_error':'Recarga la pagina o contacta con alguien de soporte'})
				else:
					if request.POST['nombre_cliente'] == '' or request.POST['telefono_cliente'] == '' or request.POST['direcion_cliente'] == '':
						return JsonResponse({'option': 'error','detalle_error':'Revisa que los datos esten completos'})
					else:
						#print request.POST['recibe_pk_modal_Edit'],'recibe_pk_modal_Edit'
						cliente = Cliente.objects.filter(pk=request.POST['id_cliente']).update(
																nombre_completo = request.POST['nombre_cliente'],
																celular = request.POST['telefono_cliente'],
																direccion=request.POST['direcion_cliente'],
																usuario_registro=request.user
																)
						recibe_ = Cliente.objects.get(pk=request.POST['id_cliente'])
						option = '''
							<option value="{}" selected>{}</option>
						'''.format(recibe_.pk, recibe_.nombre_completo+'|'+recibe_.celular)
						# return JsonResponse({'option': option})
		except Exception as e:
			print('Error al editar cliente ----> ',e)
			transaction.rollback()
		else:
			return JsonResponse({'option': option})
	else:
		return render(request, '404.html')

@login_required
def obtener_clientes_ajax(request):
	empleado = Empleado.objects.get(usuario=request.user)
	empresa = EmpresaEmpleado.objects.get(empleado = empleado)
	clientes = Cliente.objects.filter(empresa=empresa.empresa).order_by('-pk')
	
	data = {
        'clientes': list(clientes.values()),
    }
	return JsonResponse(data)